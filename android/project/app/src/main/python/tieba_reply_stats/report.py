"""单用户报告（FR-011）：Excel + Markdown。"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from .version import __version__ as VERSION

FILENAME_BAD = re.compile(r'[\/:*?"<>|\r\n\t]+')


def safe_name(text: str) -> str:
    value = FILENAME_BAD.sub("_", text or "").strip(" ._")
    return value[:60] or "user"


def report_base(user: Dict, profile: Dict) -> str:
    name = str(user.get("name") or profile.get("nickname") or "user")
    uid = str(user.get("uid") or profile.get("user_id") or "")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"用户报告_{safe_name(name)}_{safe_name(uid)}_{stamp}"


def posts_base(user: Dict, profile: Dict) -> str:
    name = str(user.get("name") or profile.get("nickname") or "user")
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return f"贴吧发言_{safe_name(name)}_{stamp}"


def export_user_posts(
    path_without_suffix: str,
    user: Dict,
    posts: Sequence[Dict],
    start_date: str = "",
    end_date: str = "",
    coverage: Optional[Dict] = None,
) -> str:
    """导出发言：汇总 + 主题帖 + 回复，链接和标题可点击。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    base = Path(path_without_suffix)
    xlsx_path = base.with_suffix(".xlsx")
    book = Workbook()
    summary = book.active
    summary.title = "汇总"
    bold = Font(bold=True)
    link_font = Font(color="0563C1", underline="single")
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    rows = list(posts or [])
    topics = [post for post in rows if post.get("type") == "主题帖"]
    replies = [post for post in rows if post.get("type") == "回复"]

    info = [
        ("昵称", user.get("name") or ""),
        ("用户名", user.get("username") or ""),
        ("用户ID", user.get("uid") or ""),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("日期范围", f"{start_date or '不限'} ~ {end_date or '不限'}"),
        ("主题帖", len(topics)),
        ("回复", len(replies)),
        ("总条数", len(rows)),
        ("回复覆盖", (coverage or {}).get("reply", "未记录")),
        ("主题覆盖", (coverage or {}).get("topic", "未记录")),
    ]
    for index, (key, value) in enumerate(info, start=1):
        summary.cell(index, 1, key).font = bold
        summary.cell(index, 2, value)
    summary.column_dimensions["A"].width = 16
    summary.column_dimensions["B"].width = 58

    by_month: Dict[str, int] = {}
    by_forum: Dict[str, int] = {}
    for post in rows:
        month = str(post.get("time") or "")[:7] or "未知"
        by_month[month] = by_month.get(month, 0) + 1
        forum = str(post.get("bar") or "") or "未知"
        by_forum[forum] = by_forum.get(forum, 0) + 1

    row = len(info) + 3
    summary.cell(row, 1, "按月份").font = bold
    summary.cell(row, 2, "条数").font = bold
    for month, count in sorted(by_month.items()):
        row += 1
        summary.cell(row, 1, month)
        summary.cell(row, 2, count)
    row += 2
    summary.cell(row, 1, "按贴吧").font = bold
    summary.cell(row, 2, "条数").font = bold
    for forum, count in sorted(by_forum.items(), key=lambda item: (-item[1], item[0])):
        row += 1
        summary.cell(row, 1, forum)
        summary.cell(row, 2, count)

    headers = ["日期", "类型", "贴吧", "标题", "内容", "链接", "用户名", "用户ID"]
    widths = [12, 9, 14, 34, 60, 46, 16, 14]

    def add_sheet(title: str, items: Sequence[Dict]) -> None:
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = bold
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for post in items:
            url = str(post.get("url") or "")
            sheet.append(
                [
                    post.get("time") or "",
                    post.get("type") or "",
                    post.get("bar") or "",
                    post.get("title") or "",
                    post.get("content") or "",
                    url,
                    user.get("username") or "",
                    user.get("uid") or "",
                ]
            )
            line = sheet.max_row
            if url:
                for column in (4, 6):
                    cell = sheet.cell(line, column)
                    cell.hyperlink = url
                    cell.font = link_font
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"

    add_sheet("主题帖", topics)
    add_sheet("回复", replies)
    summary.freeze_panes = "A1"
    summary["A1"].comment = None
    footer_row = max(20, summary.max_row + 3)
    summary.cell(footer_row, 1, "说明").font = bold
    summary.cell(footer_row, 2, "内容为接口返回片段，最多 120 字；链接和标题可点击。")
    book.save(xlsx_path)
    return str(xlsx_path)


def _merged_forums(profile: Dict) -> List[Dict]:
    result: Dict[str, Dict] = {}
    for key, source in (
        ("followed_forums", "公开"),
        ("hidden_forums", "隐藏恢复"),
        ("like_forum", "精选"),
        ("probed_forums", "候选确认"),
    ):
        for forum in profile.get(key) or []:
            name = forum.get("name")
            if not name:
                continue
            result[name] = {
                "name": name,
                "level": forum.get("level") or 0,
                "level_name": forum.get("level_name") or "",
                "source": source,
            }
    return sorted(result.values(), key=lambda item: int(item.get("level") or 0), reverse=True)


def _participation(user: Dict, threads: Sequence) -> List[Dict]:
    rows = []
    uid = str(user.get("uid") or "")
    username = str(user.get("username") or "")
    name = str(user.get("name") or "")
    for thread in threads:
        for item in thread.user_stats(True):
            if (
                (uid and str(item.get("uid") or "") == uid)
                or (username and str(item.get("username") or "") == username)
                or (name and str(item.get("name") or "") == name)
            ):
                rows.append(
                    {
                        "thread": thread.meta.title or thread.meta.tid,
                        "forum": thread.meta.forum or "",
                        "floor": item.get("floor_count") or 0,
                        "comment": item.get("comment_count") or 0,
                        "total": item.get("total") or 0,
                        "first": item.get("first_floor") or "",
                    }
                )
                break
    return rows


def export_user_report(
    path_without_suffix: str,
    user: Dict,
    profile: Dict,
    threads: Sequence,
    mode: str = "",
    mark: Optional[Dict] = None,
) -> tuple:
    """返回 (xlsx_path, md_path)。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    base = Path(path_without_suffix)
    xlsx_path = base.with_suffix(".xlsx")
    md_path = base.with_suffix(".md")
    forums = _merged_forums(profile)
    posts = profile.get("recent_posts") or []
    participation = _participation(user, threads)

    book = Workbook()
    sheet = book.active
    sheet.title = "基本信息"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EAEFF7")
    info = [
        ("昵称", profile.get("nickname") or user.get("name") or ""),
        ("用户名", profile.get("username") or user.get("username") or ""),
        ("用户 ID", profile.get("user_id") or user.get("uid") or ""),
        ("等级", user.get("level") or ""),
        ("粉丝", profile.get("fans") if profile.get("fans") is not None else ""),
        ("关注数", profile.get("concern") if profile.get("concern") is not None else ""),
        ("发帖数", profile.get("post_num") or ""),
        ("抓取模式", mode),
        ("软件版本", VERSION),
        ("帖子范围", f"{len(threads)} 个帖子"),
        ("关注增强", "开" if mode == "登录" else "关"),
        ("关注数据状态", "已加载" if profile.get("follow_loaded", True) else "未加载"),
        ("数据完整性", f"有警告 {sum(len(t.warnings) for t in threads)} 条" if any(t.warnings for t in threads) else "完整"),
        ("红绿灯状态", "已完成" if mark else "未启用"),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if mark:
        info.append(("本地标记", f"{mark.get('level', '')} {mark.get('bar', '')}".strip()))
    for index, (key, value) in enumerate(info, start=1):
        sheet.cell(index, 1, key).font = bold
        sheet.cell(index, 1).fill = fill
        sheet.cell(index, 2, value)
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 50

    forum_sheet = book.create_sheet("关注贴吧")
    forum_sheet.append(["贴吧", "等级", "来源"])
    for forum in forums:
        forum_sheet.append([forum["name"], forum["level"] or "-", forum["source"]])
    for cell in forum_sheet[1]:
        cell.font = bold

    post_sheet = book.create_sheet("发言")
    post_sheet.append(["日期", "类型", "贴吧", "标题", "内容", "链接"])
    for post in posts:
        post_sheet.append(
            [
                post.get("time") or "",
                post.get("type") or "",
                post.get("bar") or "",
                post.get("title") or "",
                post.get("content") or "",
                post.get("url") or "",
            ]
        )
    for cell in post_sheet[1]:
        cell.font = bold

    part_sheet = book.create_sheet("本帖参与")
    part_sheet.append(["帖子", "贴吧", "楼层", "楼中楼", "合计", "首次楼层"])
    for row in participation:
        part_sheet.append(
            [row["thread"], row["forum"], row["floor"], row["comment"], row["total"], row["first"]]
        )
    for cell in part_sheet[1]:
        cell.font = bold

    if mark:
        mark_sheet = book.create_sheet("本地标记")
        mark_sheet.append(["颜色", "命中贴吧"])
        mark_sheet.append([mark.get("level") or "", mark.get("bar") or ""])
        for cell in mark_sheet[1]:
            cell.font = bold

    book.save(xlsx_path)

    lines = [
        f"# 用户报告：{profile.get('nickname') or user.get('name') or ''}",
        "",
        "## 基本信息",
        "",
    ]
    for key, value in info:
        lines.append(f"- {key}：{value}")
    lines += ["", "## 关注贴吧", "", "| 贴吧 | 等级 | 来源 |", "|---|---|---|"]
    for forum in forums:
        lines.append(f"| {forum['name']} | {forum['level'] or '-'} | {forum['source']} |")
    lines += ["", "## 发言", "", "| 日期 | 类型 | 贴吧 | 标题 | 内容 |", "|---|---|---|---|---|"]
    for post in posts:
        content = str(post.get("content") or "").replace("|", "\\|").replace("\n", " ")
        lines.append(
            f"| {post.get('time') or ''} | {post.get('type') or ''} | {post.get('bar') or ''} | "
            f"{post.get('title') or ''} | {content} |"
        )
    lines += ["", "## 本帖参与", "", "| 帖子 | 贴吧 | 楼层 | 楼中楼 | 合计 |", "|---|---|---|---|---|"]
    for row in participation:
        lines.append(
            f"| {row['thread']} | {row['forum']} | {row['floor']} | {row['comment']} | {row['total']} |"
        )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return str(xlsx_path), str(md_path)
