"""Excel 导出：把 ThreadData 写成带格式的工作簿。"""

from __future__ import annotations

import math
from dataclasses import dataclass
from datetime import datetime
from typing import Dict, Iterable, List, Optional, Sequence
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import ThreadData

MAX_CELL_CHARS = 32000
MAX_EXCEL_ROWS = 1_048_575

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
ZEBRA_FILL = PatternFill("solid", fgColor="F4F7FB")
TOTAL_FILL = PatternFill("solid", fgColor="EAEFF7")
HEADER_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10, color="262626")
BOLD_FONT = Font(name="微软雅黑", size=10, bold=True, color="262626")
LINK_FONT = Font(name="微软雅黑", size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
FILL_NAMES = ("回复统计", "楼层明细", "楼中楼明细", "帖子信息")
TAB_COLORS = ("1F4E79", "2E75B6", "548235", "7F7F7F")


@dataclass
class Col:
    title: str
    width: float
    align: str = "left"
    wrap: bool = False
    fmt: Optional[str] = None
    text: bool = False


def display_width(text: str) -> int:
    width = 0
    for char in text:
        width += 2 if ord(char) > 0x2E80 else 1
    return width


def clean_text(value: str) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", value or "")
    if len(text) > MAX_CELL_CHARS:
        text = text[: MAX_CELL_CHARS - 12] + "…（已截断）"
    return text


def set_text(cell, value: str) -> None:
    """写入纯文本；避免以 = + - @ 开头的内容被 Excel 当成公式。"""
    cell.value = clean_text(value)
    if cell.data_type == "f":
        cell.data_type = "s"


def set_widths(ws: Worksheet, cols: Sequence[Col]) -> None:
    for index, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(index)].width = col.width


def write_table(
    ws: Worksheet,
    cols: Sequence[Col],
    rows: Iterable[Sequence],
    total_row: Optional[Sequence] = None,
    links: Optional[List[Dict[int, str]]] = None,
) -> int:
    """写表头 + 数据 + 可选合计行，返回最后一行行号。

    links: 与数据行一一对应，每项 {列号(1 起): 超链接}。
    """
    set_widths(ws, cols)
    for index, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=index, value=col.title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 26

    data_last_row = 1
    for row_index, row in enumerate(rows, start=2):
        if row_index > MAX_EXCEL_ROWS:
            break
        data_last_row = row_index
        row_links = links[row_index - 2] if links and row_index - 2 < len(links) else {}
        for col_index, col in enumerate(cols, start=1):
            value = row[col_index - 1] if col_index - 1 < len(row) else None
            cell = ws.cell(row=row_index, column=col_index)
            _fill_cell(cell, col, value)
            if row_index % 2 == 0:
                cell.fill = ZEBRA_FILL
            url = row_links.get(col_index)
            if url and cell.value not in (None, "", "-"):
                cell.hyperlink = url
                cell.font = LINK_FONT
        _fit_row(ws, row_index, cols, row)

    last_column = get_column_letter(len(cols))
    ws.auto_filter.ref = f"A1:{last_column}{data_last_row}"

    row_index = data_last_row
    if total_row:
        row_index += 1
        for col_index, col in enumerate(cols, start=1):
            value = total_row[col_index - 1] if col_index - 1 < len(total_row) else None
            cell = ws.cell(row=row_index, column=col_index)
            _fill_cell(cell, col, value)
            cell.font = BOLD_FONT
            cell.fill = TOTAL_FILL
            cell.border = Border(left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color="B4C6E7"))

    ws.freeze_panes = "A2"
    return row_index


def _fill_cell(cell, col: Col, value) -> None:
    cell.font = BODY_FONT
    cell.border = BORDER
    alignment = Alignment(
        horizontal=col.align,
        vertical="top" if col.wrap else "center",
        wrap_text=col.wrap,
    )
    cell.alignment = alignment

    if value is None or value == "":
        cell.value = None
        return

    if col.text or isinstance(value, str):
        set_text(cell, str(value))
        if col.text and col.title in ("贴吧ID", "用户ID", "帖子ID"):
            cell.number_format = "@"
        return

    cell.value = value
    if col.fmt:
        cell.number_format = col.fmt


def _fit_row(ws: Worksheet, row_index: int, cols: Sequence[Col], row: Sequence) -> None:
    lines = 1
    for col_index, col in enumerate(cols):
        if not col.wrap or col_index >= len(row):
            continue
        value = row[col_index]
        if not isinstance(value, str) or not value:
            continue
        width = max(8.0, col.width - 1)
        estimated = 0
        for segment in value.split("\n"):
            estimated += max(1, math.ceil(display_width(segment) / width))
        lines = max(lines, min(estimated, 12))
    if lines > 1:
        ws.row_dimensions[row_index].height = min(15 * lines + 4, 190)


def _fmt_level(level: str) -> str:
    level = str(level or "").strip()
    if level in ("", "0", "None"):
        return "-"
    return f"Lv.{level}"


def _tieba_cell(item) -> str:
    if isinstance(item, dict):
        value = str(item.get("tieba_uid") or "")
        has_identity = bool(item.get("uid") or item.get("user_id") or item.get("username") or item.get("portrait"))
    else:
        value = str(getattr(item, "tieba_uid", "") or "")
        has_identity = bool(
            getattr(item, "author_id", "") or getattr(item, "author_name", "") or getattr(item, "portrait", "")
        )
    if value:
        return value
    failed = bool(item.get("tieba_uid_failed")) if isinstance(item, dict) else bool(
        getattr(item, "tieba_uid_failed", False)
    )
    if failed:
        return "-"
    return "待补充" if has_identity else "-"


def profile_url(username: str, portrait: str = "") -> str:
    if username:
        return "https://tieba.baidu.com/home/main?un=" + quote(str(username))
    if portrait:
        return "https://tieba.baidu.com/home/main?id=" + quote(str(portrait).split("?")[0])
    return ""


def post_url(tid: str, pid: str) -> str:
    if not tid or not pid:
        return ""
    return f"https://tieba.baidu.com/p/{tid}?pid={pid}"


def _user_rows(thread: ThreadData, include_op: bool = True) -> List[list]:
    stats = thread.user_stats(include_op)
    total = sum(item["total"] for item in stats) or 1
    rows = []
    for rank, item in enumerate(stats, start=1):
        rows.append(
            [
                rank,
                item["name"],
                item["username"] or "-",
                _tieba_cell(item),
                item["uid"] or "-",
                _fmt_level(item["level"]),
                item["badge"] or "-",
                item["floor_count"],
                item["comment_count"],
                item["total"],
                item["total"] / total,
                f"{item['first_floor']}楼" if item["first_floor"] else "-",
                "；".join(item.get("aliases") or []),
            ]
        )
    return rows


USER_COLS = [
    Col("排名", 7, "center"),
    Col("昵称", 20),
    Col("用户名", 20),
    Col("贴吧ID", 22, "center", text=True),
    Col("用户ID", 22, "center", text=True),
    Col("等级", 10, "center"),
    Col("头衔", 15),
    Col("主题楼层", 13, "center", fmt="0"),
    Col("楼中楼回复", 15, "center", fmt="0"),
    Col("回复合计", 13, "center", fmt="0"),
    Col("占比", 11, "center", fmt="0.0%"),
    Col("首次出现", 13, "center"),
    Col("曾用昵称", 28),
]

FLOOR_COLS = [
    Col("楼层", 9, "center"),
    Col("昵称", 20),
    Col("用户名", 20),
    Col("贴吧ID", 22, "center", text=True),
    Col("用户ID", 22, "center", text=True),
    Col("等级", 10, "center"),
    Col("头衔", 15),
    Col("发布时间", 19, "center", fmt="yyyy-mm-dd hh:mm"),
    Col("楼中楼数", 13, "center", fmt="0"),
    Col("内容", 56, wrap=True),
]

COMMENT_COLS = [
    Col("所在楼层", 11, "center"),
    Col("楼层作者", 18),
    Col("回复者", 20),
    Col("用户名", 20),
    Col("贴吧ID", 22, "center", text=True),
    Col("用户ID", 22, "center", text=True),
    Col("等级", 10, "center"),
    Col("回复时间", 19, "center", fmt="yyyy-mm-dd hh:mm"),
    Col("被回复", 16),
    Col("内容", 56, wrap=True),
]


def export_workbook(
    path: str,
    threads: List[ThreadData],
    generated_at: Optional[datetime] = None,
    include_op: bool = True,
    meta: Optional[Dict[str, str]] = None,
) -> str:
    """写出 xlsx，返回文件路径。"""
    if not threads:
        raise ValueError("没有可导出的帖子数据")

    generated_at = generated_at or datetime.now()
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "OK"
    workbook.properties.title = "OK的贴吧玩具箱"
    workbook.properties.description = "由 OK的贴吧玩具箱 自动生成"

    if len(threads) > 1:
        _write_overview(workbook, threads, include_op)

    for index, thread in enumerate(threads, start=1):
        prefix = f"{index}·" if len(threads) > 1 else ""
        _write_thread(workbook, thread, prefix, generated_at, include_op)

    if meta:
        _write_export_notes(workbook, meta, generated_at)
    workbook.save(path)
    return path


def _write_export_notes(workbook: Workbook, meta: Dict[str, str], generated_at: datetime) -> None:
    sheet = workbook.create_sheet("导出说明")
    sheet.sheet_properties.tabColor = "7F7F7F"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 80
    rows = [("生成时间", generated_at.strftime("%Y-%m-%d %H:%M:%S"))]
    rows.extend((str(key), str(value)) for key, value in meta.items())
    for index, (key, value) in enumerate(rows, start=1):
        cell_key = sheet.cell(row=index, column=1, value=key)
        cell_key.font = BOLD_FONT
        cell_key.fill = PatternFill("solid", fgColor="EAEFF7")
        cell_key.border = BORDER
        cell = sheet.cell(row=index, column=2)
        set_text(cell, value)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    sheet.freeze_panes = "A1"


def _write_thread(
    workbook: Workbook,
    thread: ThreadData,
    prefix: str,
    generated_at: datetime,
    include_op: bool = True,
) -> None:
    stats = _user_rows(thread, include_op)
    user_items = thread.user_stats(include_op)
    user_links = []
    for item in user_items:
        url = profile_url(item.get("username") or "", item.get("portrait") or "")
        user_links.append({2: url, 3: url, 4: url} if url else {})
    sheet = workbook.create_sheet(f"{prefix}回复统计")
    sheet.sheet_properties.tabColor = TAB_COLORS[0]
    total_count = sum(item["total"] for item in thread.user_stats(include_op))
    total_row = [
        "合计",
        f"{len(stats)} 位用户",
        "",
        "",
        "",
        "",
        sum(item["floor_count"] for item in thread.user_stats(include_op)),
        len(thread.comments),
        total_count,
        1 if total_count else 0,
        "",
        "",
    ]
    write_table(sheet, USER_COLS, stats, total_row=total_row, links=user_links)

    floor_rows = [
        [
            floor.post_no,
            floor.author,
            floor.author_name or "-",
            _tieba_cell(floor),
            floor.author_id or "-",
            _fmt_level(floor.level),
            floor.badge or "-",
            floor.posted_at,
            floor.reply_num,
            floor.content or "（空）",
        ]
        for floor in thread.floors
    ]
    floor_links = []
    for floor in thread.floors:
        profile = profile_url(floor.author_name, floor.portrait)
        floor_links.append({2: profile, 3: profile, 10: post_url(thread.meta.tid, floor.pid)})
    sheet = workbook.create_sheet(f"{prefix}楼层明细")
    sheet.sheet_properties.tabColor = TAB_COLORS[1]
    write_table(sheet, FLOOR_COLS, floor_rows, links=floor_links)

    comment_rows = [
        [
            f"{comment.floor_no}楼",
            comment.floor_author,
            comment.author,
            comment.author_name or "-",
            _tieba_cell(comment),
            comment.author_id or "-",
            _fmt_level(comment.level),
            comment.posted_at,
            comment.reply_to or "",
            comment.content or "（空）",
        ]
        for comment in thread.comments
    ]
    comment_links = []
    for comment in thread.comments:
        profile = profile_url(comment.author_name, comment.portrait)
        comment_links.append({3: profile, 4: profile, 10: post_url(thread.meta.tid, comment.pid)})
    sheet = workbook.create_sheet(f"{prefix}楼中楼明细")
    sheet.sheet_properties.tabColor = TAB_COLORS[2]
    write_table(sheet, COMMENT_COLS, comment_rows, links=comment_links)

    _write_info(workbook, thread, prefix, generated_at)


def _write_info(workbook: Workbook, thread: ThreadData, prefix: str, generated_at: datetime) -> None:
    sheet = workbook.create_sheet(f"{prefix}帖子信息")
    sheet.sheet_properties.tabColor = TAB_COLORS[3]
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 80

    rows = [
        ("帖子标题", thread.meta.title or "（未获取）"),
        ("所属贴吧", thread.meta.forum or "（未获取）"),
        ("帖子链接", thread.meta.url),
        ("帖子 ID", thread.meta.tid),
        ("楼层数", len(thread.floors)),
        ("楼中楼回复数", len(thread.comments)),
        ("回复总数", len(thread.floors) + len(thread.comments)),
        ("参与用户数", len(thread.user_stats())),
        ("抓取页数", thread.pages_fetched),
        ("抓取模式", {"guest": "免登录", "login": "登录"}.get(thread.meta.mode, "未知")),
        ("抓取时间", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("数据完整性", "完整" if thread.complete else "部分缺失（见备注）"),
        ("备注", "\n".join(thread.warnings) if thread.warnings else "无"),
        (
            "说明",
            "1. 主题楼层与楼中楼回复合并统计，按用户汇总。\n"
            "2. 正文/回复为纯文本：图片记为 [图片]，表情仅保留文字。\n"
            "3. 被折叠、被删除、需要登录才能查看的回复可能缺失。\n"
            "4. 时间均为北京时间（贴吧服务器时间）。",
        ),
    ]

    for row_index, (key, value) in enumerate(rows, start=1):
        cell_key = sheet.cell(row=row_index, column=1, value=key)
        cell_key.font = BOLD_FONT
        cell_key.alignment = Alignment(vertical="top", horizontal="left")
        cell_key.fill = PatternFill("solid", fgColor="EAEFF7")
        cell_key.border = BORDER

        cell_value = sheet.cell(row=row_index, column=2)
        cell_value.font = BODY_FONT
        cell_value.border = BORDER
        if isinstance(value, int):
            cell_value.value = value
            cell_value.number_format = "0"
            cell_value.alignment = Alignment(horizontal="left", vertical="center")
        else:
            set_text(cell_value, str(value))
            cell_value.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

        if key == "帖子链接":
            cell_value.hyperlink = thread.meta.url
            cell_value.font = LINK_FONT
        if key == "备注" and thread.warnings:
            cell_value.fill = PatternFill("solid", fgColor="FFF2CC")
        if key in ("备注", "说明") and str(value):
            sheet.row_dimensions[row_index].height = max(18, min(15 * str(value).count("\n") + 18, 150))

    sheet.freeze_panes = "A1"


def _write_overview(workbook: Workbook, threads: List[ThreadData], include_op: bool = True) -> None:
    sheet = workbook.create_sheet("总览")
    sheet.sheet_properties.tabColor = "C00000"
    cols = [
        Col("#", 6, "center"),
        Col("帖子标题", 46),
        Col("贴吧", 18),
        Col("楼层", 9, "center", fmt="0"),
        Col("楼中楼", 11, "center", fmt="0"),
        Col("回复总数", 12, "center", fmt="0"),
        Col("用户数", 10, "center", fmt="0"),
        Col("链接", 46),
        Col("状态", 12, "center"),
    ]
    rows = []
    for index, thread in enumerate(threads, start=1):
        rows.append(
            [
                index,
                thread.meta.title or thread.meta.tid,
                thread.meta.forum or "-",
                len(thread.floors),
                len(thread.comments),
                len(thread.floors) + len(thread.comments),
                len(thread.user_stats()),
                thread.meta.url,
                "完整" if thread.complete else "部分缺失",
            ]
        )
    last_row = write_table(sheet, cols, rows)
    for row in range(2, last_row + 1):
        cell = sheet.cell(row=row, column=8)
        if cell.value:
            cell.hyperlink = str(cell.value)
            cell.font = LINK_FONT

    # 全吧用户总榜
    totals: dict[str, dict] = {}
    for thread in threads:
        for item in thread.user_stats(include_op):
            key = item["uid"] or f"name:{item['name']}"
            entry = totals.setdefault(
                key,
                {
                    "name": item["name"],
                    "uid": item["uid"],
                    "tieba_uid": item.get("tieba_uid") or "",
                    "threads": 0,
                    "floors": 0,
                    "comments": 0,
                },
            )
            if not entry.get("tieba_uid"):
                entry["tieba_uid"] = item.get("tieba_uid") or ""
            entry["threads"] += 1
            entry["floors"] += item["floor_count"]
            entry["comments"] += item["comment_count"]
    ranked = sorted(totals.values(), key=lambda item: (-(item["floors"] + item["comments"]), item["name"]))
    sheet = workbook.create_sheet("用户总榜")
    sheet.sheet_properties.tabColor = "BF8F00"
    cols = [
        Col("排名", 7, "center"),
        Col("用户", 24),
        Col("贴吧ID", 22, "center", text=True),
        Col("用户ID", 22, "center", text=True),
        Col("参与帖子", 12, "center", fmt="0"),
        Col("主题楼层", 13, "center", fmt="0"),
        Col("楼中楼", 11, "center", fmt="0"),
        Col("回复合计", 13, "center", fmt="0"),
    ]
    rows = [
        [
            rank,
            item["name"],
            item.get("tieba_uid") or "-",
            item["uid"] or "-",
            item["threads"],
            item["floors"],
            item["comments"],
            item["floors"] + item["comments"],
        ]
        for rank, item in enumerate(ranked, start=1)
    ]
    write_table(sheet, cols, rows)
