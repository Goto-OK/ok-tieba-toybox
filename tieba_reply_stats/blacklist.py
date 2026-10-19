"""贴吧黑白名单：把贴吧分成「安全 / 警告 / 禁止」，用来给用户打标记。

配置保存在 %APPDATA%\\TiebaReplyStats\\blacklist.json，格式：
{"safe": ["贴吧名", ...], "warn": [...], "ban": [...]}
"""

from __future__ import annotations

import json
import os
import unicodedata
from typing import Dict, List, Tuple

APP_DIR = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "TiebaReplyStats")
CONFIG_PATH = os.path.join(APP_DIR, "blacklist.json")

LEVEL_LABEL = {
    "ban": "禁止",
    "warn": "警告",
    "safe": "安全",
}


def load_blacklist() -> Dict[str, List[str]]:
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as handle:
            data = json.load(handle)
    except (OSError, json.JSONDecodeError):
        data = {}
    result = {}
    for key in ("safe", "warn", "ban"):
        value = data.get(key) if isinstance(data, dict) else None
        result[key] = [str(item).strip() for item in (value or []) if str(item).strip()]
    return result


def save_blacklist(data: Dict[str, List[str]]) -> bool:
    payload = {
        key: [str(item).strip() for item in (data.get(key) or []) if str(item).strip()]
        for key in ("safe", "warn", "ban")
    }
    try:
        os.makedirs(APP_DIR, exist_ok=True)
        with open(CONFIG_PATH, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
        return True
    except OSError:
        return False


def _norm(name: str) -> str:
    """统一吧名：全角转半角、去空格、转小写、去掉末尾一个“吧”。"""
    text = unicodedata.normalize("NFKC", str(name or "")).strip().lower()
    text = text.replace(" ", "").replace("\u3000", "")
    if text.endswith("吧"):
        text = text[:-1]
    return text


def classify(bars: List[str], data: Dict[str, List[str]]) -> Tuple[str, str]:
    """根据用户发过言的贴吧列表，返回 (级别, 命中的吧名)。

    优先级：禁止(红) > 警告(黄) > 安全(绿)；都没命中返回 ("", "")。
    """
    normalized = [(bar, _norm(bar)) for bar in bars if bar]
    for level in ("ban", "warn", "safe"):
        wanted = {_norm(name) for name in data.get(level) or []}
        if not wanted:
            continue
        for original, name in normalized:
            if name and name in wanted:
                return level, original
    return "", ""


def merge_blacklists(old: Dict[str, List[str]], new: Dict[str, List[str]]) -> Dict[str, List[str]]:
    """合并两份名单；同吧按 禁止 > 警告 > 安全 保留最高级。"""
    rank = {"safe": 0, "warn": 1, "ban": 2}
    chosen: Dict[str, str] = {}
    for data in (old, new):
        for level in ("safe", "warn", "ban"):
            for name in data.get(level) or []:
                key = _norm(name)
                if not key:
                    continue
                if key not in chosen or rank[level] > rank[chosen[key]]:
                    chosen[key] = level
    result: Dict[str, List[str]] = {"safe": [], "warn": [], "ban": []}
    for key, level in chosen.items():
        result[level].append(key)
    for level in result:
        result[level].sort()
    return result


def marker(level: str) -> str:
    return {"ban": "🔴", "warn": "🟡", "safe": "🟢"}.get(level, "⚪")


CATEGORY_LABELS = {"safe": "安全", "warn": "警告", "ban": "禁止"}
LABEL_TO_KEY = {
    "安全": "safe",
    "警告": "warn",
    "禁止": "ban",
    "safe": "safe",
    "warn": "warn",
    "ban": "ban",
    "绿": "safe",
    "黄": "warn",
    "红": "ban",
}


def export_xlsx(path: str, data: Dict[str, List[str]]) -> bool:
    """导出为 Excel：类别 | 贴吧名，带下拉选项和颜色。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return False
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "黑名单"
    sheet.append(["类别", "贴吧名"])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    fills = {
        "safe": PatternFill("solid", fgColor="DCFCE7"),
        "warn": PatternFill("solid", fgColor="FEF3C7"),
        "ban": PatternFill("solid", fgColor="FEE2E2"),
    }
    for key in ("safe", "warn", "ban"):
        for name in data.get(key) or []:
            clean = str(name).strip()
            if not clean:
                continue
            sheet.append([CATEGORY_LABELS[key], clean])
            for cell in sheet[sheet.max_row]:
                cell.fill = fills[key]
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 34
    sheet.freeze_panes = "A2"
    last_row = max(2, sheet.max_row)
    validation = DataValidation(type="list", formula1='"安全,警告,禁止"', allow_blank=False)
    sheet.add_data_validation(validation)
    validation.add(f"A2:A{last_row}")
    try:
        workbook.save(path)
        return True
    except OSError:
        return False


def import_xlsx(path: str) -> Dict[str, List[str]]:
    """读取 export_xlsx 生成的表格（也兼容 类别,贴吧名 两列的手工表格）。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"safe": [], "warn": [], "ban": []}
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    result: Dict[str, List[str]] = {"safe": [], "warn": [], "ban": []}
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        category = str(row[0] or "").strip()
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        if not name:
            continue
        key = LABEL_TO_KEY.get(category) or LABEL_TO_KEY.get(category.lower())
        if not key:
            continue
        clean = name.strip().strip("，,")
        if clean and clean not in result[key]:
            result[key].append(clean)
    workbook.close()
    return result
