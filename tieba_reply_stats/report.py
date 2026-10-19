"""单用户报告（FR-011）：Excel + Markdown。"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from .version import __version__ as VERSION

FILENAME_BAD = re.compile(r'[\/:*?"<>|\r\n\t]+')


def safe_name(text: str) -> str:
    value = FILENAME_BAD.sub("_", text or "").strip(" ._")
    return value[:60] or "user"


def report_base(user: Dict, profile: Dict) -> str:
    name = str(user.get("name") or profile.get("nickname") or "user")
    uid = str(user.get("uid") or profile.get("user_id") or "")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"用户报告_{safe_name(name)}_{safe_name(uid)}_{stamp}"


def _merged_forums(profile: Dict) -> List[Dict]:
    result: Dict[str, Dict] = {}
    for key, source in (
        ("followed_forums", "公开"),
        ("hidden_forums", "隐藏恢复"),
        ("like_forum", "精选"),
        ("probed_forums", "候选确认"),
    ):
        for forum in profile.get(key) or []:
            name = forum.get("name")
            if not name:
                continue
            result[name] = {
                "name": name,
                "level": forum.get("level") or 0,
                "level_name": forum.get("level_name") or "",
                "source": source,
            }
    return sorted(result.values(), key=lambda item: int(item.get("level") or 0), reverse=True)


def _participation(user: Dict, threads: Sequence) -> List[Dict]:
    rows = []
    uid = str(user.get("uid") or "")
    username = str(user.get("username") or "")
    name = str(user.get("name") or "")
    for thread in threads:
        for item in thread.user_stats(True):
            if (
                (uid and str(item.get("uid") or "") == uid)
                or (username and str(item.get("username") or "") == username)
                or (name and str(item.get("name") or "") == name)
            ):
                rows.append(
                    {
                        "thread": thread.meta.title or thread.meta.tid,
                        "forum": thread.meta.forum or "",
                        "floor": item.get("floor_count") or 0,
                        "comment": item.get("comment_count") or 0,
                        "total": item.get("total") or 0,
                        "first": item.get("first_floor") or "",
                    }
                )
                break
    return rows


def export_user_report(
    path_without_suffix: str,
    user: Dict,
    profile: Dict,
    threads: Sequence,
    mode: str = "",
    mark: Optional[Dict] = None,
) -> tuple:
    """返回 (xlsx_path, md_path)。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    base = Path(path_without_suffix)
    xlsx_path = base.with_suffix(".xlsx")
    md_path = base.with_suffix(".md")
    forums = _merged_forums(profile)
    posts = profile.get("recent_posts") or []
    participation = _participation(user, threads)

    book = Workbook()
    sheet = book.active
    sheet.title = "基本信息"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EAEFF7")
    info = [
        ("昵称", profile.get("nickname") or user.get("name") or ""),
        ("用户名", profile.get("username") or user.get("username") or ""),
        ("用户 ID", profile.get("user_id") or user.get("uid") or ""),
        ("等级", user.get("level") or ""),
        ("粉丝", profile.get("fans") if profile.get("fans") is not None else ""),
        ("关注数", profile.get("concern") if profile.get("concern") is not None else ""),
        ("发帖数", profile.get("post_num") or ""),
        ("抓取模式", mode),
        ("软件版本", VERSION),
        ("帖子范围", f"{len(threads)} 个帖子"),
        ("关注数据状态", "已加载" if profile.get("follow_loaded", True) else "未加载"),
        ("数据完整性", f"有警告 {sum(len(t.warnings) for t in threads)} 条" if any(t.warnings for t in threads) else "完整"),
        ("红绿灯状态", "已完成" if mark else "未启用"),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if mark:
        info.append(("本地标记", f"{mark.get('level', '')} {mark.get('bar', '')}".strip()))
    for index, (key, value) in enumerate(info, start=1):
        sheet.cell(index, 1, key).font = bold
        sheet.cell(index, 1).fill = fill
        sheet.cell(index, 2, value)
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 50

    forum_sheet = book.create_sheet("关注贴吧")
    forum_sheet.append(["贴吧", "等级", "来源"])
    for forum in forums:
        forum_sheet.append([forum["name"], forum["level"] or "-", forum["source"]])
    for cell in forum_sheet[1]:
        cell.font = bold

    post_sheet = book.create_sheet("最近发言")
    post_sheet.append(["日期", "类型", "贴吧", "标题", "内容", "链接"])
    for post in posts:
        post_sheet.append(
            [
                post.get("time") or "",
                post.get("type") or "",
                post.get("bar") or "",
                post.get("title") or "",
                post.get("content") or "",
                post.get("url") or "",
            ]
        )
    for cell in post_sheet[1]:
        cell.font = bold

    part_sheet = book.create_sheet("本帖参与")
    part_sheet.append(["帖子", "贴吧", "楼层", "楼中楼", "合计", "首次楼层"])
    for row in participation:
        part_sheet.append(
            [row["thread"], row["forum"], row["floor"], row["comment"], row["total"], row["first"]]
        )
    for cell in part_sheet[1]:
        cell.font = bold

    if mark:
        mark_sheet = book.create_sheet("本地标记")
        mark_sheet.append(["颜色", "命中贴吧"])
        mark_sheet.append([mark.get("level") or "", mark.get("bar") or ""])
        for cell in mark_sheet[1]:
            cell.font = bold

    book.save(xlsx_path)

    lines = [
        f"# 用户报告：{profile.get('nickname') or user.get('name') or ''}",
        "",
        "## 基本信息",
        "",
    ]
    for key, value in info:
        lines.append(f"- {key}：{value}")
    lines += ["", "## 关注贴吧", "", "| 贴吧 | 等级 | 来源 |", "|---|---|---|"]
    for forum in forums:
        lines.append(f"| {forum['name']} | {forum['level'] or '-'} | {forum['source']} |")
    lines += ["", "## 最近发言", "", "| 日期 | 类型 | 贴吧 | 标题 | 内容 |", "|---|---|---|---|---|"]
    for post in posts:
        content = str(post.get("content") or "").replace("|", "\\|").replace("\n", " ")
        lines.append(
            f"| {post.get('time') or ''} | {post.get('type') or ''} | {post.get('bar') or ''} | "
            f"{post.get('title') or ''} | {content} |"
        )
    lines += ["", "## 本帖参与", "", "| 帖子 | 贴吧 | 楼层 | 楼中楼 | 合计 |", "|---|---|---|---|---|"]
    for row in participation:
        lines.append(
            f"| {row['thread']} | {row['forum']} | {row['floor']} | {row['comment']} | {row['total']} |"
        )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return str(xlsx_path), str(md_path)
