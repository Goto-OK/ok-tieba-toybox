#!/usr/bin/env python3
"""OK的贴吧玩具箱 —— 图形界面（给不想用命令行的用户）。

直接运行：python tieba_gui.py
打包后双击 exe 即可。
"""

from __future__ import annotations

import ctypes
import json
import os
import queue
import re
import sys
import threading
import time
import tkinter as tk
import traceback
import webbrowser
from dataclasses import replace
from datetime import datetime, timedelta
from tkinter import filedialog, messagebox, simpledialog, ttk
from typing import Any, Callable, Dict, List, Optional
from urllib.parse import quote

from tieba_gui_grid import GridMixin
from tieba_gui_layout import APP_TITLE, GuiLayoutMixin
from tieba_gui_lookup import LookupMixin
from tieba_gui_messages import MessageMixin
from tieba_gui_settings import SettingsMixin
from tieba_reply_stats import profile_pb, third_party, webview2
from tieba_reply_stats.blacklist import (
    LEVEL_LABEL,
    classify,
    load_blacklist,
    marker,
    merge_blacklists,
    save_blacklist,
)
from tieba_reply_stats.blacklist import (
    export_xlsx as export_blacklist_xlsx,
)
from tieba_reply_stats.blacklist import (
    import_xlsx as import_blacklist_xlsx,
)
from tieba_reply_stats.collect import (
    CollectCancelled,
    CollectError,
    CollectOptions,
    collect_thread,
    enrich_identities,
    enrich_tieba_uids,
    enrich_user_levels,
    extract_tid,
)
from tieba_reply_stats.compare import MODES, compare_threads
from tieba_reply_stats.config_backup import preview as backup_preview
from tieba_reply_stats.config_backup import read_backup, write_backup
from tieba_reply_stats.cookie_store import clear_cookie, load_cookie, save_cookie
from tieba_reply_stats.errors import friendly_error
from tieba_reply_stats.export import export_workbook
from tieba_reply_stats.fetch import (
    CaptchaError,
    Fetcher,
    FetchError,
    RateLimitError,
    ThreadGoneError,
    configure_gate,
    parse_cookie_input,
    pause_requests,
    resume_requests,
    set_requests_paused,
    set_stop_checker,
)
from tieba_reply_stats.filters import filter_users
from tieba_reply_stats.follow_probe import candidate_bars, probe_forums
from tieba_reply_stats.link_parser import parse_links
from tieba_reply_stats.lookup_supplement import merge_current_threads
from tieba_reply_stats.pbapi import fill_thread_levels
from tieba_reply_stats.post_cache import DEFAULT_LIMIT_MB, MAX_LIMIT_MB, MIN_LIMIT_MB
from tieba_reply_stats.post_cache import cache as post_cache
from tieba_reply_stats.profile import (
    cached_user_posts,
    clear_post_cache,
    fetch_user_profile,
    resolve_user,
)
from tieba_reply_stats.report import (
    export_user_posts,
    export_user_report,
    posts_base,
    report_base,
)
from tieba_reply_stats.safe_delete import remove_file as safe_remove_file
from tieba_reply_stats.safe_delete import remove_tree as safe_remove_tree
from tieba_reply_stats.settings import APP_DIR, SETTINGS_PATH, load_settings, save_settings
from tieba_reply_stats.storage import (
    IncompatibleResult,
    archive_incompatible,
    load_threads,
    save_threads,
)
from tieba_reply_stats.storage import (
    clear as clear_stored_result,
)
from tieba_reply_stats.storage import (
    old_files as stored_old_files,
)
from tieba_reply_stats.theme import map_color, resolve_palette
from tieba_reply_stats.version import __version__ as VERSION

_EMOJI_CODE_RE = re.compile(r"\[U\+([0-9A-Fa-f]{4,6})\]", re.I)

EMOJI_PICKER = (
    "😀", "😁", "😂", "🤣", "😊", "😍", "🤔", "😅",
    "😭", "😡", "👍", "👎", "❤", "🔥", "🎉", "💩",
    "🤡", "👀", "😴", "🤮", "🤤", "🙄", "😳", "😱",
    "🤯", "🥰", "😎", "🤓", "😐", "😑", "🙃", "🙏",
)


def display_text(value) -> str:
    """界面直接显示 emoji；Tk 8.6 以单色字形渲染。"""
    return str(value or "")


def decode_display(value) -> str:
    """把 [U+XXXX] 还原成 emoji，用于回查。"""
    return _EMOJI_CODE_RE.sub(lambda match: chr(int(match.group(1), 16)), str(value or ""))


BDUSS_HELP = """只填 BDUSS 时，这样复制：

1. 用 Chrome 或 Edge 打开贴吧页面，确认已登录。
2. 按 F12，打开开发者工具。
3. 点「应用 / Application」→「存储 / Storage」→「Cookie / Cookies」→ https://tieba.baidu.com
4. 在列表里找到名为 BDUSS 的一项，复制它的「值 / Value」。
5. 粘贴到本窗口的 BDUSS 输入框，点「应用」。
6. 勾选「记住 BDUSS」后，下次打开自动填入。

只复制 BDUSS 的值，不要复制整行 Cookie。
BDUSS 等同于登录凭证，只在自己的电脑上使用，不要发给别人。
"""

COOKIE_HELP = """需要 Cookie 时，这样复制：

1. 用 Chrome 或 Edge 打开目标帖子，确认能正常看到楼层。
2. 按键盘上的 F12，打开开发者工具。
3. 点顶部的「网络 / Network」标签，然后按 F5 刷新页面。
4. 在左侧列表里点第一个以帖子数字开头的请求（类型是 document）。
5. 在右侧找到「标头 / Headers」→「请求标头 / Request Headers」，
   里面有一行 Cookie:，把它冒号后面的整段内容复制下来。
6. 粘贴到本窗口的「完整 Cookie」输入框（也可以复制后点「从剪贴板粘贴」），
   然后点「开始统计」。
7. 需要下次自动填入时，勾选「记住完整 Cookie」再点「应用」。

为什么不能全自动读取？
浏览器从 2024 年起对 Cookie 做了「应用绑定加密」，只有浏览器自己能解密，
而且还锁着 Cookie 文件，其他程序读不了。所以需要你手动复制一次，
之后由本程序用 Windows 账户加密保存在本机，下次自动使用。

提示：
- Cookie 相当于账号凭证，只在自己的电脑上使用，不要发给别人。
- Cookie 一般能用几个月；如果提示验证失败，就重新复制一次。
- 如果 Cookie 也一直提示验证失败，可以换个网络（比如手机热点）再试。
"""


def default_output_path() -> str:
    desktop = ""
    try:
        import winreg

        with winreg.OpenKey(
            winreg.HKEY_CURRENT_USER,
            r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders",
        ) as key:
            value, _kind = winreg.QueryValueEx(key, "Desktop")
        desktop = os.path.expandvars(value)
    except OSError:
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    if not os.path.isdir(desktop):
        desktop = os.path.expanduser("~")
    return os.path.join(desktop, f"OK的贴吧玩具箱_{datetime.now():%Y%m%d_%H%M}.xlsx")


def resource_dir() -> str:
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def _optional_int(text: str) -> Optional[int]:
    try:
        return int(str(text).strip())
    except ValueError:
        return None







class WorkerMixin:
    """抓取生命周期与 worker 方法，从 App 拆出。"""

    last_threads: Any
    worker: Any
    _webview_cookie_session: Any
    _identity_choice_queue: Any
    _rate_choice_queue: Any
    _manual_timing: Any
    _tieba_uid_failed: Any
    _user_id_failed: Any
    _emoji_window: Any
    _active_choice_window: Any

    def __getattr__(self, name: str) -> Any:
        raise AttributeError(name)

    def _start_from_key(self, _event=None):
        self._on_start()
        return "break"


    def _on_start(self) -> None:
        if self.running or getattr(self, "_supplement_running", False):
            return
        raw = self.url_text.get("1.0", "end").strip()
        tokens = [item for item in re.split(r"[\s,，;；]+", raw) if item]
        urls, bad = [], []
        for token in tokens:
            try:
                extract_tid(token)
                urls.append(token)
            except ValueError:
                bad.append(token)
        if not urls:
            self.status_var.set("先粘贴至少一个帖子链接")
            self.url_text.focus_set()
            return
        if bad:
            messagebox.showwarning(
                "有内容无法识别",
                "下面这些不是有效的帖子链接，将被忽略：\n\n" + "\n".join(bad[:6]),
                parent=self.root,
            )

        existing = {thread.meta.tid for thread in self.last_threads}
        current = {extract_tid(url) for url in urls}
        removed = existing - current
        removed_action = self.settings.get("removed_thread_action", "keep")
        if removed and removed_action == "ask":
            choice = self._ask_choice(
                "旧帖处理",
                f"有 {len(removed)} 个已抓过的帖子不在本次输入框中。\n"
                "要移除这些旧结果，还是保留？",
                [("移除旧结果", "remove"), ("保留旧结果", "keep"), ("取消", None)],
                default=0,
            )
            if choice is None:
                return
            removed_action = choice
        if removed and removed_action == "remove":
            self.last_threads = [thread for thread in self.last_threads if thread.meta.tid not in removed]
            existing = {thread.meta.tid for thread in self.last_threads}
            self._append_log(f"已移除 {len(removed)} 个旧帖结果。", "muted")
        if existing and not self.settings.get("auto_refresh_old", False):
            fresh = [url for url in urls if extract_tid(url) not in existing]
            if not fresh:
                messagebox.showinfo("没有新帖", "这些帖子都已抓过；可在设置里开启自动刷新。", parent=self.root)
                return
            if len(fresh) != len(urls):
                self._append_log(f"跳过已抓过的 {len(urls) - len(fresh)} 个帖子。", "muted")
            urls = fresh
        export = self.export_var.get()
        output = self.output_var.get().strip() or default_output_path()
        if export:
            if not output.lower().endswith(".xlsx"):
                output += ".xlsx"
            self._set_output(output)
            if os.path.exists(output) and not messagebox.askyesno(
                "文件已存在",
                f"这个文件已经存在：\n{output}\n\n要覆盖它吗？",
                parent=self.root,
                default=messagebox.NO,
            ):
                return
        else:
            output = ""

        try:
            delay = float(self.delay_var.get())
        except ValueError:
            delay = 0.8
        delay = min(5.0, max(0.3, delay))

        mode = self.mode_var.get()
        cookie = (
            self._request_cookie()
            if mode == "guest"
            else parse_cookie_input(self._current_cookie_text())
        )
        if mode == "login" and self.login_method_var.get() == "bduss" and not cookie.get("BDUSS"):
            messagebox.showwarning(
                "需要 BDUSS",
                "「仅 BDUSS」模式需要填写 BDUSS；也可以切换到「完整 Cookie」。",
                parent=self.root,
            )
            self.bduss_entry.focus_set()
            return
        if mode == "login" and not webview2.has_login_cookie(cookie):
            messagebox.showwarning(
                "需要登录 Cookie",
                "完整 Cookie 里没有 BDUSS/TIEBA_SID，无法确认登录态。\n"
                "请重新获取完整 Cookie，或改用游客模式。",
                parent=self.root,
            )
            self.cookie_entry.focus_set()
            return
        if mode == "guest":
            self._append_log("本次使用免登录模式：公开帖子，楼中楼按设置补全", "muted")
        elif getattr(self, "_webview_cookie_session", False):
            self._append_log("本次完整 Cookie 来自内置验证，不自动保存到本机。", "muted")
            self._webview_cookie_session = False
        elif self._remember_current_login():
            if save_cookie(self._current_cookie_text()):
                self._append_log("已记住登录信息（本机加密保存），下次会自动填入。", "muted")
            else:
                self._append_log("登录信息保存失败，本次仍可正常使用。", "muted")

        options = CollectOptions(
            mode=mode,
            include_comments=self.comments_var.get(),
            comment_depth=self._comment_depth_value(),
            comment_concurrency=max(1, min(3, int(self.comment_concurrency_var.get() or 1))),
            enrich_levels=True,
            enrich_identities=False,
            pb_levels=True,
            verbose=True,
            log=self._log_from_thread,
            should_stop=self.stop_event.is_set,
        )
        include_op = not self.exclude_op_var.get()
        self.last_include_op = include_op

        self._update_risk_summary()
        risk_items = self._high_risk_items()
        if risk_items:
            if not self._ask_choice(
                "高风险设置",
                "本次有以下高风险设置：\n- "
                + "\n- ".join(risk_items)
                + "\n\n可能触发验证码或临时限流，不保证不会受限。",
                [("返回调整", False), ("仍然开始", True)],
                default=0,
            ):
                return

        self._current_urls = list(urls)
        resume_requests()
        configure_gate(float(delay or 0))
        self.stop_event.clear()
        self._task_paused = False
        set_stop_checker(self.stop_event.is_set)
        self._identity_choice_queue = queue.Queue()
        self._rate_choice_queue = queue.Queue()
        self._manual_timing = {}
        self._tieba_uid_failed = set()
        self._user_id_failed = set()
        self._write_log("fetch_start", threads=len(urls))
        self._set_running(True)
        self._append_log(f"\n开始统计 {len(urls)} 个帖子，请耐心等待…", "title")
        base_threads = list(self.last_threads)
        self.worker = threading.Thread(
            target=self._run_worker,
            args=(
                urls,
                output,
                options,
                cookie,
                delay,
                include_op,
                export,
                base_threads,
                self._identity_strategy(),
                int(self.identity_concurrency_var.get() or 3),
            ),
            daemon=True,
        )
        self.worker.start()


    def _on_pause(self) -> None:
        if not self.running:
            return
        paused = not getattr(self, "_task_paused", False)
        self._task_paused = paused
        set_requests_paused(paused)
        self.pause_btn.configure(text="继续" if paused else "暂停")
        if paused:
            self.status_var.set("已暂停；已完成数据保留")
            self._append_log("已暂停，不再发新请求。", "muted")
        else:
            self.status_var.set("已继续抓取")
            self._append_log("已继续抓取。", "muted")


    def _on_stop(self) -> None:
        resume_requests()
        active_choice = getattr(self, "_active_choice_window", None)
        if active_choice is not None:
            try:
                active_choice.destroy()
            except tk.TclError:
                pass
            self._active_choice_window = None
        if getattr(self, "_supplement_running", False):
            self._supplement_stop.set()
            self.stop_btn.configure(state="disabled")
            self.status_var.set("正在停止补全…")
            return
        if not self.running:
            return
        self.stop_event.set()
        self.stop_btn.configure(state="disabled")
        self.status_var.set("正在停止…")


    def _run_worker(
        self,
        urls,
        output,
        options,
        cookie,
        delay,
        include_op,
        export=True,
        base_threads=None,
        identity_strategy="full",
        identity_concurrency=3,
    ) -> None:
        completed = []
        began = time.perf_counter()
        stage_at = began
        timings: Dict[str, float] = {}
        try:
            fetcher = Fetcher(cookie=cookie, delay=0, verbose=False)
            for index, url in enumerate(urls, start=1):
                self._post(("log", f"\n[{index}/{len(urls)}] {url}", "title"))
                try:
                    tid = extract_tid(url)
                except ValueError:
                    tid = ""
                self._write_log("thread_start", tid=tid)
                rate_attempt = 0
                thread = None
                while True:
                    try:
                        thread = collect_thread(url, fetcher, options)
                        break
                    except RateLimitError as error:
                        if rate_attempt >= 3:
                            raise
                        rate_attempt += 1
                        self._post(
                            ("rate_limit_ask", str(error), float(error.retry_after or 0), url)
                        )
                        choice = None
                        while choice is None:
                            try:
                                choice = self._rate_choice_queue.get(timeout=0.2)
                            except queue.Empty:
                                if self.stop_event.is_set():
                                    raise CollectCancelled("已停止") from error
                        if choice == "stop":
                            raise CollectCancelled("已停止") from error
                        pause_requests(float(choice))
                        continue
                    except CaptchaError:
                        raise
                    except CollectCancelled:
                        raise
                    except (ThreadGoneError, FetchError, CollectError, ValueError) as error:
                        self._write_log("thread_error", tid=tid, error=type(error).__name__)
                        self._post(("log", f"  ✗ 抓取失败：{error}", "err"))
                        thread = None
                        break
                if thread is None:
                    continue
                completed.append(thread)
                users = thread.user_stats()
                leveled = sum(1 for item in users if item.get("level"))
                self._write_log(
                    "thread_done",
                    tid=thread.meta.tid,
                    floors=len(thread.floors),
                    comments=len(thread.comments),
                    users=len(users),
                )
                self._write_log("level_done", tid=thread.meta.tid, leveled=leveled, users=len(users))
                title = thread.meta.title or thread.meta.tid
                self._post(
                    (
                        "log",
                        f"  ✓ {title}：{len(thread.floors)} 层 / {len(thread.comments)} 条楼中楼 / "
                        f"{len(thread.user_stats())} 位用户",
                        "ok",
                    )
                )
                for warning in thread.warnings[:3]:
                    self._post(("log", f"    ! {warning}", "muted"))

            timings["抓取"] = time.perf_counter() - stage_at
            stage_at = time.perf_counter()
            if not completed:
                self._post(
                    ("error", "没有抓到任何帖子。\n\n可能原因：帖子被删、链接不对，或者被安全验证拦截。")
                )
                return

            final_threads = []
            by_tid = {str(thread.meta.tid): thread for thread in (base_threads or [])}
            for thread in completed:
                by_tid[str(thread.meta.tid)] = thread
            final_threads = list(by_tid.values())
            timings["合并"] = time.perf_counter() - stage_at
            stage_at = time.perf_counter()

            identity_options = CollectOptions(
                verbose=True,
                log=self._log_from_thread,
                should_stop=self.stop_event.is_set,
            )
            missing_uid = 0
            missing_tieba = 0
            seen_users: Dict[str, Dict[str, object]] = {}
            for thread in final_threads:
                for item in list(thread.floors) + list(thread.comments):
                    uid = str(getattr(item, "author_id", "") or "")
                    key = (
                        f"uid:{uid}"
                        if uid
                        else (
                            f"un:{item.author_name}"
                            if item.author_name
                            else (f"pt:{item.portrait}" if item.portrait else f"nm:{item.author}")
                        )
                    )
                    info = seen_users.setdefault(key, {"uid": uid, "tieba": ""})
                    if getattr(item, "tieba_uid", ""):
                        info["tieba"] = item.tieba_uid
            for info in seen_users.values():
                if not info["uid"]:
                    missing_uid += 1
                elif not info["tieba"]:
                    missing_tieba += 1

            strategy = identity_strategy
            if strategy == "ask" and (missing_uid or missing_tieba):
                self._post(("identity_ask", missing_uid, missing_tieba))
                choice = None
                remember = False
                while choice is None:
                    try:
                        choice, remember = self._identity_choice_queue.get(timeout=0.2)
                    except queue.Empty:
                        if self.stop_event.is_set():
                            choice, remember = "skip", False
                strategy = choice
                if remember:
                    self._post(("identity_remember", strategy))
            elif strategy == "ask":
                strategy = "skip"

            if strategy in ("full", "user_id"):
                self._post(("stage", "抓取完成，正在补全身份…"))
                for thread in final_threads:
                    enrich_identities(thread, fetcher, identity_options, self._log_from_thread)
                self._identity_missing = sum(
                    1
                    for thread in final_threads
                    for item in list(thread.floors) + list(thread.comments)
                    if not item.author_id and (item.author_name or item.portrait)
                )
            else:
                self._identity_missing = missing_uid
            self._write_log(
                "identity_done",
                missing=self._identity_missing,
                users=sum(len(thread.user_stats()) for thread in final_threads),
            )
            if self._identity_missing:
                self._post(
                    ("log", f"  ⚠ 仍有 {self._identity_missing} 条身份待补充（不影响已抓内容）", "err")
                )

            if strategy in ("full", "user_id"):
                self._user_id_failed = {
                    self._identity_fail_key(item)
                    for thread in final_threads
                    for item in list(thread.floors) + list(thread.comments)
                    if not item.author_id and (item.author_name or item.portrait)
                }
            else:
                self._user_id_failed = set()
            self._apply_user_id_failed_states(final_threads)
            timings["用户ID"] = time.perf_counter() - stage_at
            stage_at = time.perf_counter()
            if strategy == "full":
                self._post(("stage", "身份完成，正在补全贴吧ID…"))
                uid_stats = enrich_tieba_uids(
                    final_threads,
                    fetcher,
                    identity_options,
                    self._log_from_thread,
                    progress=lambda done, total: self._post(
                        ("stage", f"正在补全贴吧ID {done}/{total}…")
                    ),
                    concurrency=max(1, int(identity_concurrency or 3)),
                )
                self._tieba_uid_missing = int(uid_stats.get("failed") or 0)
                self._tieba_uid_failed = set(uid_stats.get("failed_uids") or [])
            else:
                self._tieba_uid_missing = missing_tieba
                self._tieba_uid_failed = set()
            self._last_identity_strategy = strategy
            self._apply_tieba_failed_states(final_threads)
            self._write_log("tieba_uid_done", missing=self._tieba_uid_missing, users=len(final_threads))
            timings["贴吧ID"] = time.perf_counter() - stage_at
            stage_at = time.perf_counter()
            self._post(("stage", "正在保存结果…"))
            self._persist_result(final_threads)
            self._write_log("save", threads=len(final_threads), stage="identity")
            timings["保存"] = time.perf_counter() - stage_at
            stage_at = time.perf_counter()

            if export:
                self._post(("stage", "正在生成 Excel…"))
                export_workbook(output, final_threads, include_op=include_op)
                timings["导出"] = time.perf_counter() - stage_at
                stage_at = time.perf_counter()
                self._write_log("export", ok=True, threads=len(final_threads))
                self._write_log(
                    "done",
                    threads=len(final_threads),
                    users=sum(len(thread.user_stats()) for thread in final_threads),
                    exported=True,
                )
                timings["总耗时"] = time.perf_counter() - began
                self._post(("done", output, len(final_threads), list(final_threads), dict(timings)))
            else:
                timings["导出"] = 0.0
                timings["总耗时"] = time.perf_counter() - began
                self._write_log(
                    "done",
                    threads=len(final_threads),
                    users=sum(len(thread.user_stats()) for thread in final_threads),
                    exported=False,
                )
                self._post(("done", None, len(final_threads), list(final_threads), dict(timings)))
        except CaptchaError as error:
            self._write_log("error", error="CaptchaError")
            partial = []
            items = list(completed)
            if getattr(error, "partial", None) is not None:
                items.append(error.partial)
            if items:
                by_tid = {str(thread.meta.tid): thread for thread in (base_threads or [])}
                for thread in items:
                    by_tid[str(thread.meta.tid)] = thread
                partial = list(by_tid.values())
                self._persist_result(partial)
            self._post(("captcha", str(error), list(partial)))
        except RateLimitError as error:
            self._write_log("error", error="RateLimitError")
            self._post(("rate_limit", str(error), float(error.retry_after or 0)))
        except CollectCancelled:
            self._write_log("cancelled", threads=len(completed), exported=bool(export and completed))
            if completed:
                by_tid = {str(thread.meta.tid): thread for thread in (base_threads or [])}
                for thread in completed:
                    by_tid[str(thread.meta.tid)] = thread
                final_threads = list(by_tid.values())
                self._persist_result(final_threads)
                if export:
                    export_workbook(output, final_threads, include_op=include_op)
                    self._write_log("export", ok=True, threads=len(final_threads))
                self._post(
                    (
                        "cancelled_partial",
                        output if export else None,
                        len(final_threads),
                        list(final_threads),
                    )
                )
            else:
                self._post(("cancelled",))
        except Exception as error:  # noqa: BLE001 - 界面上要给出可读信息
            self._write_log("error", error=type(error).__name__)
            self._post(("log", traceback.format_exc(), "muted"))
            self._post(("error", friendly_error(error)))


class ExportMixin:
    """导出/打开结果方法，从 App 拆出。"""

    last_output: Any

    def __getattr__(self, name: str) -> Any:
        raise AttributeError(name)

    def _open_result(self) -> None:
        raw = self.last_outputs or ([self.last_output] if self.last_output else [])
        paths = [path for path in raw if isinstance(path, str) and os.path.isfile(path)]
        if not paths:
            return
        if len(paths) > 1:
            choice = self._ask_choice(
                "打开拆分文件",
                f"共 {len(paths)} 个文件，选择打开方式。",
                [("只打开第一个", False), ("全部打开", True), ("取消", None)],
                default=0,
            )
            if choice is None:
                return
            if not choice:
                paths = paths[:1]
        failed = []
        for path in paths:
            try:
                os.startfile(path)  # type: ignore[attr-defined]
            except OSError:
                failed.append(path)
        if failed:
            messagebox.showerror(
                "部分文件打不开", "打不开：" + chr(10) + chr(10).join(failed), parent=self.root
            )


    def _split_export_threads(self, threads, limit: int = 500_000):
        parts = []
        current: list = []
        rows = 0
        for thread in threads:
            size = len(thread.floors) + len(thread.comments)
            if current and rows + size > limit:
                parts.append(current)
                current = []
                rows = 0
            current.append(thread)
            rows += size
        if current:
            parts.append(current)
        return parts


    def _export_now(self) -> None:
        if self._network_task_active():
            messagebox.showinfo(
                "正在处理",
                "抓取、补全或重试仍在进行，完成后才能导出。",
                parent=self.root,
            )
            return
        self._export_now_actual()


    def _export_now_actual(self) -> None:
        if not self.last_threads:
            self.status_var.set("还没有可导出的数据")
            self._append_log("导出前需要先抓取帖子。", "muted")
            return
        threads = self._selected_grid_threads()
        if len(self.last_threads) > 1 and not self.grid_thread_var.get().startswith("全部"):
            scope = self._ask_choice(
                "导出帖子范围",
                "选择这次导出的帖子范围。",
                [("当前表格", True), ("全部帖子", False), ("取消", None)],
                default=0,
            )
            if scope is None:
                return
            threads = list(self.last_threads) if scope is False else threads
        if self._people_filter_active():
            scope = self._ask_choice(
                "导出用户范围",
                "当前有参与者筛选，选择导出哪些用户。",
                [("当前筛选用户", True), ("全部用户", False), ("取消", None)],
                default=0,
            )
            if scope is None:
                return
            if scope is True:
                filtered_ids = {thread.meta.tid for thread in self._filtered_threads()}
                threads = [thread for thread in threads if thread.meta.tid in filtered_ids]
                if not any(thread.floors or thread.comments for thread in threads):
                    messagebox.showinfo("没有结果", "当前筛选没有可导出的用户。", parent=self.root)
                    return
        if not threads:
            messagebox.showinfo("没有结果", "当前范围没有可导出的帖子。", parent=self.root)
            return
        initial = self.last_output or default_output_path()
        path = filedialog.asksaveasfilename(
            title="导出 Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile=os.path.basename(initial),
            initialdir=os.path.dirname(initial),
            parent=self.root,
        )
        if not path:
            return
        visible = self._visible_people()
        scope_user = "当前筛选用户" if self._people_filter_active() else "全部用户"
        warnings = [f"{thread.meta.title or thread.meta.tid}：{len(thread.warnings)} 条"
                    for thread in threads if thread.warnings]
        empty = [thread.meta.title or thread.meta.tid for thread in threads if not (thread.floors or thread.comments)]
        all_users = list(getattr(self, "people_users", []))
        pending_tieba = sum(
            1 for user in all_users if not user.get("tieba_uid") and (user.get("uid") or user.get("username"))
        )
        timing_enabled = bool(getattr(self, "identity_timing_var", tk.BooleanVar(value=True)).get())
        stage_times = getattr(self, "_last_stage_times", {}) if timing_enabled else {}
        manual_times = getattr(self, "_manual_timing", {}) if timing_enabled else {}
        meta = {
            "软件版本": VERSION,
            "补全策略": {
                "full": "补全用户ID + 贴吧ID",
                "ask": "每次询问",
                "skip": "跳过补全",
                "user_id": "只补用户ID",
                "detail": "详情再补",
            }.get(
                str(getattr(self, "_last_identity_strategy", self.settings.get("identity_strategy", "full"))),
                "补全用户ID + 贴吧ID",
            ),
            "补全状态": (
                "跳过"
                if str(
                    getattr(
                        self,
                        "_last_identity_strategy",
                        self.settings.get("identity_strategy", "full"),
                    )
                )
                in ("skip", "detail")
                else ("部分" if pending_tieba or getattr(self, "_tieba_uid_missing", 0) else "完成")
            ),
            "补全用户ID失败": str(int(getattr(self, "_identity_missing", 0) or 0)),
            "补全贴吧ID失败": str(int(getattr(self, "_tieba_uid_missing", 0) or 0)),
            "待补充贴吧ID用户": str(pending_tieba),
            "不完整帖子": str(
                sum(1 for thread in self.last_threads if not getattr(thread, "complete", True))
            ),
            "不完整帖子标题": "；".join(
                display_text(thread.meta.title or thread.meta.tid)
                for thread in self.last_threads
                if not getattr(thread, "complete", True)
            )
            or "无",
            "总耗时": (
                f"{float(stage_times.get('总耗时') or 0):.1f} 秒"
                if stage_times.get("总耗时")
                else "未开启"
            ),
            "手动补全耗时": (
                f"{float(manual_times.get('总耗时') or 0):.1f} 秒"
                if manual_times.get("总耗时")
                else "未开启"
            ),
            "手动补全用户ID耗时": (
                f"{float(manual_times.get('用户ID') or 0):.1f} 秒"
                if manual_times
                else "未开启"
            ),
            "手动补全贴吧ID耗时": (
                f"{float(manual_times.get('贴吧ID') or 0):.1f} 秒"
                if manual_times
                else "未开启"
            ),
            "帖子范围": "当前表格范围" if len(threads) < len(self.last_threads) else "全部帖子",
            "帖子数": str(len(threads)),
            "用户范围": scope_user,
            "筛选后人数": str(len(visible)),
            "实际导出用户数": str(len(visible) if scope_user == "当前筛选用户" else len(self.people_users)),
            "筛选条件": self._filter_summary(),
            "关注增强": "开" if self.settings.get("follow_enhance") else "关",
            "包含关注": "是" if self.include_follows_var.get() else "否",

            "数据警告": "；".join(warnings) if warnings else "无",
            "空帖子": "；".join(empty) if empty else "无",
        }
        parts = self._split_export_threads(self._expand_large_threads(threads))
        paths = []
        try:
            for index, part in enumerate(parts, start=1):
                if index == 1:
                    part_path = path
                else:
                    root, ext = os.path.splitext(path)
                    part_path = f"{root}_part{index}{ext or '.xlsx'}"
                part_meta = dict(meta)
                part_meta["分片"] = f"{index}/{len(parts)}"
                export_workbook(
                    part_path,
                    part,
                    include_op=getattr(self, "last_include_op", True),
                    meta=part_meta,
                )
                paths.append(part_path)
        except Exception as error:  # noqa: BLE001
            messagebox.showerror("导出失败", friendly_error(error), parent=self.root)
            return
        self.last_output = paths[0]
        self.open_btn.configure(state="normal")
        self._append_log("已导出 Excel：" + "；".join(paths), "ok")
        messagebox.showinfo(
            "导出完成",
            "已保存到：" + chr(10) + chr(10).join(paths) + chr(10) + chr(10) + "现在打开第一个吗？",
            parent=self.root,
        )
        self._open_result()


    def _export_posts_for(self, user, profile, start: str, end: str) -> None:

        sources = profile.get("post_sources") or {}
        if sources.get("topic_truncated"):
            if not messagebox.askokcancel(
                "主题覆盖未到开始日期",
                "主题流已达到本次页数上限，导出的主题可能不完整。\n"
                "仍按现有缓存覆盖导出？",
                parent=self.root,
            ):
                return
        uid = str(user.get("uid") or profile.get("user_id") or "")
        cached = post_cache().get_user(uid) if uid else None
        posts = list(profile.get("recent_posts") or [])
        if cached and cached.get("posts"):
            posts = list(cached["posts"])
        cache_start = cache_end = ""
        if cached:
            oldest = int(cached.get("oldest_ts") or 0)
            newest = int(cached.get("newest_ts") or 0)
            if oldest:
                cache_start = datetime.fromtimestamp(oldest).strftime("%Y-%m-%d")
            if newest:
                cache_end = datetime.fromtimestamp(newest).strftime("%Y-%m-%d")
        chosen = self._post_range_dialog(user, start, end, cache_start, cache_end)
        if chosen is None:
            return
        chosen_start, chosen_end = chosen
        filtered = []
        for post in posts:
            day = str(post.get("time") or "")
            if chosen_start and day and day < chosen_start:
                continue
            if chosen_end and day and day > chosen_end:
                continue
            filtered.append(post)
        if not filtered:
            messagebox.showinfo("没有发言", "选定范围内没有缓存发言。", parent=self.root)
            return
        topics = [post for post in filtered if post.get("type") == "主题帖"]
        replies = [post for post in filtered if post.get("type") == "回复"]
        initial = posts_base(user, profile) + ".xlsx"
        path = filedialog.asksaveasfilename(
            title="导出发言",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile=initial,
            parent=self.root,
        )
        if not path:
            return
        base = os.path.splitext(path)[0]

        def coverage_text(items, label: str) -> str:
            dates = []
            unknown = 0
            for item in items:
                day = str(item.get("time") or "").strip()
                if day:
                    dates.append(day)
                else:
                    unknown += 1
            if not dates:
                return f"无{label}"
            text = f"{min(dates)} ~ {max(dates)}"
            if unknown:
                text += f"（含 {unknown} 条未知时间）"
            return text

        try:
            coverage = {
                "reply": coverage_text(replies, "回复"),
                "topic": coverage_text(topics, "主题"),
            }
            output = export_user_posts(base, user, filtered, chosen_start, chosen_end, coverage)
        except Exception as error:  # noqa: BLE001
            messagebox.showerror("导出失败", friendly_error(error), parent=self.root)
            return
        self._append_log(f"发言已导出：{output}", "ok")
        messagebox.showinfo("导出完成", "已生成：" + chr(10) + output, parent=self.root)


class BatchIdentityMixin:
    """身份补全与批量标记方法，从 App 拆出。"""

    _active_choice_window: Any
    _batch_pending: Any
    _tieba_uid_failed: Any

    def __getattr__(self, name: str) -> Any:
        raise AttributeError(name)

    def _supplement_identity_manual(self) -> None:
        if getattr(self, "_supplement_running", False) or not getattr(self, "last_threads", None):
            return
        users = getattr(self, "people_users", [])
        missing_uid = sum(
            1 for user in users if not user.get("uid") and (user.get("username") or user.get("portrait"))
        )
        missing_tieba = sum(1 for user in users if user.get("uid") and not user.get("tieba_uid"))
        if not missing_uid and not missing_tieba:
            messagebox.showinfo("无需补全", "当前没有可补充的身份。", parent=self.root)
            return
        strategy, remember = self._ask_identity_strategy(
            missing_uid, missing_tieba, allow_detail=False
        )
        if remember:
            self.identity_strategy_var.set(
                {
                    "full": "补全用户ID + 贴吧ID",
                    "skip": "跳过补全",
                    "user_id": "只补用户ID",
                    "detail": "详情再补",
                }.get(strategy, "每次询问")
            )
            self._save_identity_strategy()
        if strategy in ("skip", "detail"):
            return
        resume_requests()
        configure_gate(float(self.delay_var.get() or 0.8))
        self._set_supplement_running(True)
        self._last_identity_strategy = strategy
        self._tieba_uid_failed = set()
        self._supplement_stop.clear()
        set_stop_checker(self._supplement_stop.is_set)
        self.supplement_btn.configure(state="disabled")
        self.stop_btn.configure(state="normal")
        self.people_hint.configure(text="正在补全身份…")

        def worker():
            try:
                began = time.perf_counter()
                stage_at = began
                timings: Dict[str, float] = {}
                fetcher = Fetcher(cookie=self._request_cookie(), delay=0, verbose=False)
                options = CollectOptions(
                    verbose=True,
                    log=self._log_from_thread,
                    should_stop=self._supplement_stop.is_set,
                )
                for thread in self.last_threads:
                    enrich_identities(thread, fetcher, options, self._log_from_thread)
                self._user_id_failed = {
                    self._identity_fail_key(item)
                    for thread in self.last_threads
                    for item in list(thread.floors) + list(thread.comments)
                    if not item.author_id and (item.author_name or item.portrait)
                }
                self._apply_user_id_failed_states(self.last_threads)
                timings["用户ID"] = time.perf_counter() - stage_at
                stage_at = time.perf_counter()
                if strategy == "full":
                    uid_stats = enrich_tieba_uids(
                        self.last_threads,
                        fetcher,
                        options,
                        self._log_from_thread,
                        progress=lambda done, total: self._post(
                            ("identity_supplement_progress", done, total)
                        ),
                        concurrency=max(1, int(self.identity_concurrency_var.get() or 3)),
                    )
                    self._tieba_uid_failed = set(uid_stats.get("failed_uids") or [])
                self._apply_tieba_failed_states(self.last_threads)
                timings["贴吧ID"] = time.perf_counter() - stage_at
                timings["总耗时"] = time.perf_counter() - began
                self._manual_timing = dict(timings)
                self._persist_result(self.last_threads)
                self._post(("identity_supplement_done", strategy, dict(timings)))
            except CollectCancelled:
                self._persist_result(self.last_threads)
                self._post(("identity_supplement_stopped", dict(timings)))
            except Exception as error:  # noqa: BLE001
                self._post(("identity_supplement_error", str(error)))

        threading.Thread(target=worker, daemon=True).start()


    def _auto_fill_levels(self) -> None:
        """切到参与者时，如果还有缺等级的，自动在后台补一次（不用手动点）。"""
        if self.mode_var.get() == "guest":
            return
        if not parse_cookie_input(self.cookie_var.get()).get("BDUSS"):
            return
        if getattr(self, "_auto_filling", False) or getattr(self, "_auto_filled", False):
            return
        missing = [
            user
            for user in self.people_users
            if not user.get("level") and (user.get("uid") or user.get("username"))
        ]
        if not missing or not self.last_threads:
            return
        self._auto_filling = True
        self._auto_filled = True
        self.people_hint.configure(text=f"正在补全 {len(missing)} 人的等级…")

        def worker():
            try:
                fetcher = Fetcher(cookie=parse_cookie_input(self.cookie_var.get()), delay=0.3, verbose=False)
                options = CollectOptions(verbose=True, log=self._log_from_thread)
                for thread in self.last_threads:
                    fill_thread_levels(thread, fetcher, self._log_from_thread)
                    enrich_user_levels(thread, fetcher, options, self._log_from_thread)
                self._post(("levels_filled",))
            except Exception as error:  # noqa: BLE001
                self._post(("levels_error", str(error)))

        threading.Thread(target=worker, daemon=True).start()


    def _identity_fail_key(self, item) -> str:
        if isinstance(item, dict):
            uid = str(item.get("uid") or "")
            username = str(item.get("username") or "")
            portrait = str(item.get("portrait") or "")
            name = str(item.get("name") or "")
        else:
            uid = str(getattr(item, "author_id", "") or "")
            username = str(getattr(item, "author_name", "") or "")
            portrait = str(getattr(item, "portrait", "") or "")
            name = str(getattr(item, "author", "") or "")
        if uid:
            return f"uid:{uid}"
        if username:
            return f"un:{username}"
        if portrait:
            return f"pt:{portrait}"
        return f"nm:{name}"


    def _ask_identity_strategy(self, missing_uid: int, missing_tieba: int, allow_detail: bool = True):
        """返回 (strategy, remember)；关闭 = skip。"""
        result = {"strategy": "skip", "remember": False}
        win = tk.Toplevel(self.root)
        win.title("抓取完成，选择补全方式" if allow_detail else "补全身份")
        win.configure(bg="#ffffff")
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        self._active_choice_window = win
        frame = tk.Frame(win, bg="#ffffff", padx=20, pady=16)
        frame.pack(fill="both", expand=True)
        tk.Label(
            frame,
            text="抓取完成，选择补全方式" if allow_detail else "补全身份",
            bg="#ffffff",
            fg="#1f2937",
            font=("Microsoft YaHei UI", 12, "bold"),
        ).pack(anchor="w")
        tk.Label(
            frame,
            text=f"缺用户ID：{missing_uid} 人\n缺贴吧ID：{missing_tieba} 人",
            bg="#ffffff",
            fg="#475569",
            font=("Microsoft YaHei UI", 9),
            justify="left",
        ).pack(anchor="w", pady=(6, 10))
        remember_var = tk.BooleanVar(
            value=bool(self.settings.get("identity_remember_choice", True))
        )
        ttk.Checkbutton(frame, text="记住本次选择", variable=remember_var).pack(anchor="w", pady=(0, 8))

        def remember_state() -> bool:
            checked = bool(remember_var.get())
            self.settings["identity_remember_choice"] = checked
            save_settings(self.settings)
            return checked

        def choose(strategy: str) -> None:
            result["strategy"] = strategy
            result["remember"] = remember_state()
            win.destroy()

        choices = [
            ("跳过本次，直接完成" if allow_detail else "跳过本次", "skip"),
            ("只补用户ID", "user_id"),
            ("补全用户ID + 贴吧ID", "full"),
        ]
        if allow_detail:
            choices.append(("详情再补", "detail"))
        def make_command(value: str) -> Callable[[], None]:
            def command() -> None:
                choose(value)

            return command

        for text, strategy in choices:
            ttk.Button(frame, text=text, command=make_command(str(strategy))).pack(
                fill="x", pady=2
            )
        tk.Label(
            frame,
            text="关闭窗口或按 Esc = 跳过本次补全，会先二次确认，不记住选择。",
            bg="#ffffff",
            fg="#8a94a6",
            font=("Microsoft YaHei UI", 8),
            wraplength=420,
            justify="left",
        ).pack(anchor="w", pady=(8, 0))

        def on_close() -> None:
            if messagebox.askyesno(
                "跳过本次补全？",
                "关闭将跳过本次补全，并使用现有数据保存/导出。\n不记住本次选择。",
                parent=win,
                default=messagebox.NO,
            ):
                remember_state()
                result["strategy"] = "skip"
                result["remember"] = False
                win.destroy()

        win.protocol("WM_DELETE_WINDOW", on_close)
        win.bind("<Escape>", lambda _event: on_close())
        win.update_idletasks()
        width = 460
        height = max(320, win.winfo_reqheight() + 16)
        x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - width) // 2)
        y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - height) // 2)
        win.geometry(f"{width}x{height}+{x}+{y}")
        self.root.wait_window(win)
        if getattr(self, "_active_choice_window", None) is win:
            self._active_choice_window = None
        self._clear_identity_strategy_selection()
        try:
            self.root.focus_set()
        except tk.TclError:
            pass
        return result["strategy"], result["remember"]


    def _clear_identity_strategy_selection(self) -> None:
        box = getattr(self, "identity_strategy_box", None)
        if box is None:
            return
        try:
            box.selection_clear()
            box.icursor("end")
        except tk.TclError:
            pass


    def _save_identity_strategy(self) -> None:
        value = {
            "补全用户ID + 贴吧ID": "full",
            "每次询问": "ask",
            "跳过补全": "skip",
            "只补用户ID": "user_id",
            "详情再补": "detail",
        }.get(self.identity_strategy_var.get(), "full")
        self.settings["identity_strategy"] = value
        if hasattr(self, "identity_timing_var"):
            self.settings["identity_timing"] = bool(self.identity_timing_var.get())
        if hasattr(self, "identity_concurrency_var"):
            self.settings["identity_concurrency"] = int(self.identity_concurrency_var.get() or 3)
        save_settings(self.settings)
        self._clear_identity_strategy_selection()
        self._update_identity_risk_hint()
        self._update_risk_summary()


    def _update_identity_risk_hint(self) -> None:
        if not hasattr(self, "identity_risk_label"):
            return
        strategy = self._identity_strategy()
        value = str(getattr(self, "identity_concurrency_var", tk.StringVar(value="3")).get())
        unused = strategy in ("skip", "user_id", "detail")
        if hasattr(self, "identity_concurrency_box"):
            self.identity_concurrency_box.configure(state="disabled" if unused else "readonly")
        if unused:
            self.identity_risk_label.configure(text="未使用", fg="#8a94a6")
        elif value == "5":
            self.identity_risk_label.configure(
                text="高风险：可能触发验证码或临时限流", fg="#b91c1c"
            )
        else:
            self.identity_risk_label.configure(text="较稳（默认）", fg="#64748b")


    def _identity_strategy(self) -> str:
        return {
            "补全用户ID + 贴吧ID": "full",
            "每次询问": "ask",
            "跳过补全": "skip",
            "只补用户ID": "user_id",
            "详情再补": "detail",
        }.get(getattr(self, "identity_strategy_var", tk.StringVar(value="")).get(), "full")


    def _enrich_levels_now(self) -> None:
        if not self.last_threads:
            self.status_var.set("还没有可补全等级的数据")
            return
        if getattr(self, "_enriching", False):
            return
        self._enriching = True
        self.people_hint.configure(text="正在补全等级，请稍候…")

        def worker():
            try:
                fetcher = Fetcher(cookie=self._request_cookie(), delay=0, verbose=False)
                options = CollectOptions(verbose=True, log=self._log_from_thread)
                for thread in self.last_threads:
                    filled = fill_thread_levels(thread, fetcher, self._log_from_thread)
                    if filled:
                        self._log_from_thread(f"  新版网页接口补全等级：{filled} 处")
                    enrich_user_levels(thread, fetcher, options, self._log_from_thread)
                if self.last_output:
                    export_workbook(
                        self.last_output,
                        self.last_threads,
                        include_op=getattr(self, "last_include_op", True),
                    )
                self._post(("enrich_done",))
            except Exception as error:  # noqa: BLE001
                self._post(("enrich_error", str(error)))

        threading.Thread(target=worker, daemon=True).start()


    def _batch_refresh_marks(self, resume_users=None) -> None:
        if getattr(self, "_batch_running", False):
            self._batch_cancel = True
            self.batch_btn.configure(text="正在停止…")
            return
        if not self._risk_enabled():
            messagebox.showinfo(
                "本地标记未启用",
                "先在「设置 → 通用」里启用本地标记与红绿灯。",
                parent=self.root,
            )
            return
        if not self._has_mark_rules():
            messagebox.showinfo(
                "没有本地标记规则",
                "先在「设置 → 本地标记」里填写安全/可疑/警告贴吧，再刷红绿灯。",
                parent=self.root,
            )
            self._focus_local_marks()
            return
        if not getattr(self, "people_users", None):
            messagebox.showinfo("还没有数据", "先点「开始统计」抓一个帖子。", parent=self.root)
            return
        users = list(resume_users) if resume_users is not None else self._visible_people()
        if not users:
            messagebox.showinfo("没有可刷新的用户", "当前筛选没有用户。", parent=self.root)
            return
        self._batch_running = True
        self._batch_cancel = False
        self._batch_done = 0
        self._batch_total = len(users)
        self._batch_plan = list(users)
        self._batch_pending = []
        self._marks_stale = False
        self._batch_success = 0
        self._batch_skip = 0
        self._batch_fail = 0
        self.resume_batch_btn.pack_forget()
        self.resume_batch_btn.configure(state="disabled", text="继续未完成")
        self.remark_btn.state(["disabled"])
        self.refresh_people_btn.state(["disabled"])
        self.include_follows_check.state(["disabled"])
        self._set_people_controls_enabled(False)
        self.batch_btn.configure(text="停止查询")
        self.people_hint.configure(text=f"批量查询 0/{self._batch_total}…")
        count = self._post_count()
        start_date, end_date = self._date_params()
        prefer = (
            "start"
            if self.post_date_mode_var.get() == "指定日期" and "最早" in self.post_prefer_var.get()
            else "end"
        )

        include_follows = bool(self.include_follows_var.get() and self.settings.get("follow_enhance"))
        guest = self.mode_var.get() == "guest"
        cookie = {} if guest else parse_cookie_input(self.cookie_var.get())
        concurrency = 1 if include_follows else max(1, min(3, int(self.settings.get("batch_concurrency", 1) or 1)))

        def resolve(user):
            name = str(user.get("name") or "").strip()
            if not any((user.get("uid"), user.get("username"), user.get("portrait"), name)):
                return {"ok": False, "skipped": True, "error": "缺少用户身份"}
            key = (
                f"{user.get('uid') or user.get('username') or name}|{count}|{start_date}|{end_date}|"
                f"{prefer}|{int(include_follows)}|{int(guest)}"
            )
            hit = self._batch_cache.get(key)
            if hit and time.monotonic() - hit[0] < 60:
                return hit[1]
            try:
                fetcher = Fetcher(cookie=cookie, delay=0.3, verbose=False)
                uid = user.get("uid") or ""
                username = user.get("username") or ""
                if not uid:
                    info = {}
                    if username:
                        info = resolve_user(fetcher, username, "用户名")
                    elif name:
                        info = resolve_user(fetcher, name, "昵称")
                    if not info.get("uid") and not user.get("portrait") and name and username:
                        info = resolve_user(fetcher, name, "昵称")
                    if info.get("uid"):
                        uid = info.get("uid") or ""
                        username = info.get("username") or username
                        user["uid"] = uid
                        user["username"] = username
                        user["portrait"] = info.get("portrait") or user.get("portrait") or ""
                    elif not user.get("portrait"):
                        return {"ok": False, "skipped": True, "error": "用户名/昵称无法唯一定位"}
                    # 没有 uid/username 但保留 portrait 时，仍可用 portrait 查发言。
                cached_posts = cached_user_posts(
                    uid, count, start_date, end_date, prefer, promote=False
                )
                if cached_posts is not None and not include_follows:
                    return {
                        "ok": True,
                        "recent_posts": cached_posts,
                        "bar_level": str(user.get("level") or ""),
                        "follow_loaded": False,
                        "followed_forums": [],
                        "hidden_forums": [],
                        "post_sources": {"cached": True},
                    }
                profile = fetch_user_profile(
                    fetcher,
                    uid,
                    user.get("username") or "",
                    user.get("portrait") or "",
                    self._user_bar(user),
                    count,
                    start_date,
                    end_date,
                    prefer,
                    load_follow=include_follows,
                    topic_max_pages=1,
                    use_cache=False,
                    cache_priority=0,
                    cached_posts=cached_posts,
                )
            except Exception as error:  # noqa: BLE001
                return {"ok": False, "error": str(error)}
            if profile.get("ok"):
                if len(self._batch_cache) >= 500:
                    self._batch_cache.clear()
                self._batch_cache[key] = (time.monotonic(), profile)
            return profile

        def worker():
            if concurrency <= 1:
                for index, user in enumerate(users):
                    if self._batch_cancel:
                        break
                    self._post(("batch_profile", index, user, resolve(user)))
            else:
                from concurrent.futures import ThreadPoolExecutor

                with ThreadPoolExecutor(max_workers=concurrency) as pool:
                    futures = [pool.submit(resolve, user) for user in users]
                    for index, (user, future) in enumerate(zip(users, futures, strict=False)):
                        if self._batch_cancel:
                            break
                        try:
                            profile = future.result()
                        except Exception as error:  # noqa: BLE001
                            profile = {"ok": False, "error": str(error)}
                        self._post(("batch_profile", index, user, profile))
            self._post(("batch_done",))

        threading.Thread(target=worker, daemon=True).start()


    def _resume_batch(self) -> None:
        if self._batch_pending:
            pending = list(self._batch_pending)
            self._batch_pending = []
            self._batch_refresh_marks(resume_users=pending)


    def _apply_batch_mark(self, user, profile) -> None:
        if not self._risk_enabled():
            return
        iid = user.get("_iid")
        if iid is None:
            return
        level = str(profile.get("bar_level") or "") if profile.get("ok") else ""
        if level in ("0", "None"):
            level = ""
        if not level and profile.get("ok"):
            level = str(user.get("level") or "")
            if level in ("0", "None"):
                level = ""
        if level and not user.get("level"):
            user["level"] = level
        if profile.get("ok"):
            if profile.get("user_id") and not user.get("uid"):
                user["uid"] = str(profile.get("user_id") or "")
            if profile.get("username") and not user.get("username"):
                user["username"] = str(profile.get("username") or "")
            if profile.get("portrait") and not user.get("portrait"):
                user["portrait"] = str(profile.get("portrait") or "")
        hit_level = ""
        hit_bar = ""
        if profile.get("ok"):
            hit_level, hit_bar = self._profile_hit(profile, user.get("bars") or [self._user_bar(user)])
        hidden = self._hidden_follows(profile) if profile.get("ok") else False
        user["mark_level"] = hit_level
        user["hit"] = hit_level
        user["mark_bar"] = hit_bar or ""
        user["mark_hidden"] = bool(hidden)
        user["mark_at"] = time.time()
        user["mark_error"] = str(profile.get("error") or "")
        self._set_person_mark_row(user)


class ThemeMixin:
    """主题应用方法，从 App 拆出。"""

    palette: Any

    def __getattr__(self, name: str) -> Any:
        raise AttributeError(name)

    def _on_theme_changed(self) -> None:
        mode = {"浅色": "light", "深色": "dark", "跟随系统": "system"}.get(
            self.theme_mode_var.get(), "light"
        )
        self.settings["theme_mode"] = mode
        save_settings(self.settings)
        self._apply_theme()


    def _theme_widget_tree(self, widget, old_palette=None) -> None:
        palette = self.palette
        for option in (
            "bg",
            "fg",
            "activebackground",
            "activeforeground",
            "selectbackground",
            "selectforeground",
            "highlightbackground",
            "insertbackground",
            "disabledforeground",
        ):
            try:
                current = widget.cget(option)
            except tk.TclError:
                continue
            value = str(current)
            token_order = ("accent", "accent_hover", "selection_bg", "selection_text")
            old_token = ""
            current_token = ""
            for token in token_order:
                if old_palette and value.lower() == str(old_palette.get(token) or "").lower():
                    old_token = token
                    break
                if value.lower() == str(palette.get(token) or "").lower():
                    current_token = token
                    break
            try:
                text = str(widget.cget("text")) if "text" in widget.keys() else ""
            except tk.TclError:
                text = ""
            if widget.winfo_class() == "Label" and text in ("红", "黄", "绿", "白"):
                new = value
            elif old_token:
                new = palette[old_token]
            elif current_token:
                new = palette[current_token]
            else:
                new = map_color(option, value, palette)
            klass = widget.winfo_class()
            if value.startswith("System"):
                if option in ("bg", "background"):
                    new = palette["surface"] if klass in ("Text", "Entry", "Listbox") else palette["window_bg"]
                elif option in ("fg", "foreground"):
                    new = palette["text"]
                elif option == "insertbackground":
                    new = palette["text"]
            if new != current:
                try:
                    widget.configure(**{option: new})
                except tk.TclError:
                    pass
        try:
            if widget.winfo_class() in ("Text", "Entry", "Listbox"):
                widget.configure(
                    bg=palette["surface"],
                    fg=palette["text"],
                    insertbackground=palette["text"],
                    selectbackground=palette["selection_bg"],
                    selectforeground=palette["selection_text"],
                )
        except tk.TclError:
            pass
        for child in widget.winfo_children():
            self._theme_widget_tree(child, old_palette)


    def _apply_theme(self) -> None:
        old_palette = dict(self.palette) if hasattr(self, "palette") else None
        palette = resolve_palette(
            str(self.settings.get("theme_mode", "light")),
            str(self.settings.get("accent", "blue")),
        )
        self.palette = palette
        self.root.configure(bg=palette["window_bg"])
        style = ttk.Style(self.root)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure(
            "TFrame",
            background=palette["window_bg"],
        )
        style.configure("TLabel", background=palette["window_bg"], foreground=palette["text"])
        style.configure(
            "TLabelframe",
            background=palette["window_bg"],
            bordercolor=palette["border"],
            lightcolor=palette["border"],
            darkcolor=palette["border"],
        )
        style.configure(
            "TLabelframe.Label",
            background=palette["window_bg"],
            foreground=palette["text"],
        )
        style.configure(
            "TCheckbutton",
            background=palette["window_bg"],
            foreground=palette["text"],
            indicatorbackground=palette["surface"],
            indicatorforeground=palette["surface"],
            upperbordercolor=palette["border"],
            lowerbordercolor=palette["border"],
        )
        style.map(
            "TCheckbutton",
            indicatorbackground=[("selected", palette["accent"]), ("!selected", palette["surface"])],
            indicatorforeground=[("selected", palette["accent_text"]), ("!selected", palette["surface"])],
        )
        style.configure(
            "TRadiobutton",
            background=palette["window_bg"],
            foreground=palette["text"],
            indicatorbackground=palette["surface"],
            indicatorforeground=palette["surface"],
            upperbordercolor=palette["border"],
            lowerbordercolor=palette["border"],
        )
        style.map(
            "TRadiobutton",
            indicatorbackground=[("selected", palette["accent"]), ("!selected", palette["surface"])],
            indicatorforeground=[("selected", palette["accent_text"]), ("!selected", palette["surface"])],
        )
        style.configure(
            "TButton",
            background=palette["button_bg"],
            foreground=palette["button_text"],
            bordercolor=palette["border"],
            lightcolor=palette["button_bg"],
            darkcolor=palette["button_bg"],
        )
        style.map(
            "TButton",
            background=[("active", palette["button_hover"]), ("disabled", palette["disabled_bg"])],
            foreground=[("disabled", palette["disabled_text"])],
        )
        style.configure(
            "TEntry",
            fieldbackground=palette["surface"],
            foreground=palette["text"],
            insertcolor=palette["text"],
            bordercolor=palette["border"],
        )
        style.configure(
            "TCombobox",
            fieldbackground=palette["surface"],
            foreground=palette["text"],
            background=palette["button_bg"],
            arrowcolor=palette["text"],
            bordercolor=palette["border"],
        )
        style.map(
            "TCombobox",
            fieldbackground=[("readonly", palette["surface"]), ("!disabled", palette["surface"])],
            foreground=[("readonly", palette["text"])],
            selectbackground=[("readonly", palette["selection_bg"])],
            selectforeground=[("readonly", palette["selection_text"])],
        )
        style.configure(
            "TSpinbox",
            fieldbackground=palette["surface"],
            foreground=palette["text"],
            background=palette["button_bg"],
            arrowcolor=palette["text"],
            bordercolor=palette["border"],
        )
        style.map(
            "TSpinbox",
            fieldbackground=[("!disabled", palette["surface"])],
            foreground=[("!disabled", palette["text"])],
        )
        style.configure(
            "Treeview",
            background=palette["table_bg"],
            fieldbackground=palette["table_bg"],
            foreground=palette["text"],
            bordercolor=palette["border"],
        )
        style.configure(
            "Treeview.Heading",
            background=palette["table_header_bg"],
            foreground=palette["table_header_text"],
            bordercolor=palette["border"],
        )
        style.map(
            "Treeview",
            background=[("selected", palette["selection_bg"])],
            foreground=[("selected", palette["selection_text"])],
        )
        style.configure("TNotebook", background=palette["window_bg"], bordercolor=palette["border"])
        style.configure(
            "TNotebook.Tab",
            background=palette["button_bg"],
            foreground=palette["text_secondary"],
            padding=(9, 3),
        )
        style.configure(
            "TNotebook.Tab",
            background=palette["button_bg"],
            foreground=palette["text_secondary"],
        )
        style.map(
            "TNotebook.Tab",
            background=[("selected", palette["surface"])],
            foreground=[("selected", palette["accent"])],
        )
        style.configure("Main.TNotebook", background=palette["window_bg"])
        style.configure(
            "Main.TNotebook.Tab",
            background=palette["button_bg"],
            foreground=palette["text_secondary"],
            padding=(16, 7),
            font=("Microsoft YaHei UI", 10, "bold"),
        )
        style.map(
            "Main.TNotebook.Tab",
            background=[("selected", palette["surface"])],
            foreground=[("selected", palette["accent"])],
        )
        style.configure("Sub.TNotebook", background=palette["window_bg"])
        style.configure(
            "Sub.TNotebook.Tab",
            background=palette["button_bg"],
            foreground=palette["text_secondary"],
            padding=(9, 3),
            font=("Microsoft YaHei UI", 9),
        )
        style.map(
            "Sub.TNotebook.Tab",
            background=[("selected", palette["surface"])],
            foreground=[("selected", palette["accent"])],
        )
        style.configure(
            "TScrollbar",
            background=palette["button_bg"],
            troughcolor=palette["surface_alt"],
            bordercolor=palette["border"],
            lightcolor=palette["button_bg"],
            darkcolor=palette["button_bg"],
            arrowcolor=palette["text_secondary"],
        )
        style.map("TScrollbar", background=[("active", palette["button_hover"])])
        for element in ("Vertical.Scrollbar.trough", "Vertical.Scrollbar.thumb"):
            style.configure(
                element,
                background=palette["button_bg"],
                troughcolor=palette["surface_alt"],
                bordercolor=palette["border"],
                lightcolor=palette["button_bg"],
                darkcolor=palette["button_bg"],
                arrowcolor=palette["text_secondary"],
            )
        for element in ("Horizontal.Scrollbar.trough", "Horizontal.Scrollbar.thumb"):
            style.configure(
                element,
                background=palette["button_bg"],
                troughcolor=palette["surface_alt"],
                bordercolor=palette["border"],
                lightcolor=palette["button_bg"],
                darkcolor=palette["button_bg"],
                arrowcolor=palette["text_secondary"],
            )
        style.configure(
            "Vertical.TScrollbar",
            background=palette["button_bg"],
            troughcolor=palette["surface_alt"],
            arrowcolor=palette["text_secondary"],
            bordercolor=palette["border"],
            lightcolor=palette["button_bg"],
            darkcolor=palette["button_bg"],
        )
        style.configure(
            "Horizontal.TScrollbar",
            background=palette["button_bg"],
            troughcolor=palette["surface_alt"],
            arrowcolor=palette["text_secondary"],
            bordercolor=palette["border"],
            lightcolor=palette["button_bg"],
            darkcolor=palette["button_bg"],
        )
        style.configure(
            "TProgressbar",
            background=palette["accent"],
            troughcolor=palette["surface_alt"],
        )
        self.root.option_add("*TCombobox*Listbox.background", palette["surface"])
        self.root.option_add("*TCombobox*Listbox.foreground", palette["text"])
        self.root.option_add("*TCombobox*Listbox.selectBackground", palette["selection_bg"])
        self.root.option_add("*TCombobox*Listbox.selectForeground", palette["selection_text"])

        self._theme_widget_tree(self.root, old_palette)

        if hasattr(self, "log_text"):
            self.log_text.tag_configure("title", foreground=palette["accent"])
            self.log_text.tag_configure("ok", foreground=palette["safe"])
            self.log_text.tag_configure("err", foreground=palette["danger"])
            self.log_text.tag_configure("muted", foreground=palette["text_muted"])

        def all_trees(widget):
            found = []
            if isinstance(widget, ttk.Treeview):
                found.append(widget)
            for child in widget.winfo_children():
                found.extend(all_trees(child))
            return found

        for tree in all_trees(self.root):
            tree.tag_configure("ban", background=palette["danger_bg"])
            tree.tag_configure("warn", background=palette["warning_bg"])
            tree.tag_configure("safe", background=palette["safe_bg"])
            tree.tag_configure("odd", background=palette["table_alt"])
        for key, value in (("ban", "danger"), ("warn", "warning"), ("safe", "safe")):
            entry = getattr(self, "lookup_traffic_labels", {}).get(key)
            if entry:
                # 只更新备用颜色，不能因为切主题就把无结果的查用户红绿灯点亮。
                self.lookup_traffic_labels[key] = (entry[0], palette[value])
        if getattr(self, "_last_lookup_profile", None):
            self._render_lookup(self._last_lookup_profile)
        else:
            self._reset_lookup_traffic()


class ProfileMixin:
    """用户资料与关注页方法，从 App 拆出。"""

    love_native_edit: Any
    love_native_host: Any
    _profile_cache: Any
    _lookup_cache: Any

    def __getattr__(self, name: str) -> Any:
        raise AttributeError(name)

    def _create_love_native_edit(self) -> bool:
        try:
            import ctypes
            from ctypes import wintypes

            user32 = ctypes.WinDLL("user32", use_last_error=True)
            gdi32 = ctypes.WinDLL("gdi32", use_last_error=True)
            kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
            user32.CreateWindowExW.restype = wintypes.HWND
            user32.CreateWindowExW.argtypes = [
                wintypes.DWORD, wintypes.LPCWSTR, wintypes.LPCWSTR, wintypes.DWORD,
                ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int,
                wintypes.HWND, wintypes.HMENU, wintypes.HINSTANCE, wintypes.LPVOID,
            ]
            user32.SendMessageW.restype = ctypes.c_ssize_t
            user32.SendMessageW.argtypes = [
                wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM
            ]
            user32.GetWindowTextLengthW.argtypes = [wintypes.HWND]
            user32.GetWindowTextW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
            user32.SetWindowTextW.argtypes = [wintypes.HWND, wintypes.LPCWSTR]
            user32.MoveWindow.argtypes = [
                wintypes.HWND, ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int, wintypes.BOOL
            ]
            gdi32.CreateFontW.restype = wintypes.HFONT
            kernel32.GetModuleHandleW.restype = wintypes.HMODULE
            style = 0x40000000 | 0x10000000 | 0x00010000 | 0x00800000 | 0x00000080
            edit = user32.CreateWindowExW(
                0x00000200, "EDIT", "", style, 0, 0, 100, 28,
                wintypes.HWND(self.love_edit_host.winfo_id()), None, kernel32.GetModuleHandleW(None), None,
            )
            if not edit:
                return False
            self.love_native_edit = edit
            self.love_native_host = self.love_edit_host
            font = gdi32.CreateFontW(
                -16, 0, 0, 0, 400, 0, 0, 0, 1, 0, 0, 0, 0, "Microsoft YaHei UI"
            )
            self._love_native_font = font
            if font:
                user32.SendMessageW(edit, 0x0030, font, 1)

            def resize(event) -> None:
                user32.MoveWindow(edit, 0, 0, max(40, event.width), max(20, event.height), True)

            self.love_edit_host.bind("<Configure>", resize)
            resize(type("E", (), {"width": 260, "height": 28})())

            user32.CallWindowProcW.restype = ctypes.c_ssize_t
            user32.CallWindowProcW.argtypes = [
                ctypes.c_void_p, wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM
            ]

            def on_key(hwnd, msg, wparam, lparam):
                if msg == 0x0100 and wparam == 0x0D:
                    self.root.after_idle(self._start_love)
                    return 0
                return user32.CallWindowProcW(old_proc, hwnd, msg, wparam, lparam)

            proc_type = ctypes.WINFUNCTYPE(
                ctypes.c_ssize_t, wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM
            )
            self._love_native_proc = proc_type(on_key)
            set_ptr = getattr(user32, "SetWindowLongPtrW", user32.SetWindowLongW)
            set_ptr.restype = ctypes.c_void_p
            set_ptr.argtypes = [wintypes.HWND, ctypes.c_int, ctypes.c_void_p]
            old_proc = set_ptr(edit, -4, ctypes.cast(self._love_native_proc, ctypes.c_void_p))
            self._love_native_old_proc = old_proc
            self._love_native_user32 = user32
            self.root.bind_all("<Button-1>", self._release_lookup_native_focus, add="+")
            return True
        except Exception:  # noqa: BLE001
            self.love_native_edit = None
            self.love_native_host = None
            return False


    def _love_text(self) -> str:
        edit = getattr(self, "love_native_edit", None)
        user32 = getattr(self, "_love_native_user32", None)
        if edit and user32 is not None:
            length = user32.GetWindowTextLengthW(edit)
            buffer = ctypes.create_unicode_buffer(length + 1)
            user32.GetWindowTextW(edit, buffer, length + 1)
            return buffer.value
        return self.love_query_var.get()


    def _love_focus(self) -> None:
        edit = getattr(self, "love_native_edit", None)
        user32 = getattr(self, "_love_native_user32", None)
        if edit and user32 is not None:
            user32.SetFocus(edit)
            return
        if self.love_entry is not None:
            self.love_entry.focus_set()


    def _probe_profile_forums(self) -> None:
        if getattr(self, "_probe_running", False):
            self._probe_stop.set()
            return
        user = getattr(self, "_profile_user", None)
        profile = getattr(self, "_profile_data", None)
        if not user or not profile:
            return
        self._start_probe(user, profile, self.profile_probe_btn)


    def _set_profile_cache_note(self, text: str) -> None:
        self.profile_cache_label.configure(text=text)
        if text:
            if not self.profile_cache_label.winfo_manager():
                self.profile_cache_label.pack(
                    fill="x", padx=6, pady=(0, 1), before=self.profile_action_bar
                )
        else:
            self.profile_cache_label.pack_forget()


    def _attach_love_levels(self, profile) -> None:
        if not self.settings.get("love_level_enabled", False):
            return
        username = str(profile.get("username") or "")
        if not username:
            return
        cache = getattr(self, "_all_levels_cache", None)
        if cache is None:
            cache = {}
            self._all_levels_cache = cache
        if username in cache:
            levels = cache[username]
        else:
            try:
                levels = third_party.fetch_all_levels(username)
            except Exception:  # noqa: BLE001
                levels = []
            cache[username] = levels
        if levels:
            profile["level_forums"] = levels


    def _request_love_levels(self, profile, forums, tree, names, kind) -> None:
        if not self.settings.get("love_level_enabled", False):
            return
        uid = str(profile.get("tieba_uid") or "")
        username = str(profile.get("username") or "")
        if not uid and not username:
            return
        missing = [forum for forum in forums if not int(forum.get("level") or 0)]
        if not missing:
            return
        cache = getattr(self, "_level_cache", None)
        if cache is None:
            cache = {}
            self._level_cache = cache
        pending = [forum for forum in missing if (uid or username, forum.get("name")) not in cache]
        pending = pending[:50]
        if not pending:
            return

        def worker():
            results = {}
            for forum in pending:
                name = str(forum.get("name") or "")
                if not name:
                    continue
                try:
                    data = third_party.fetch_level(uid, name, username)
                except Exception:  # noqa: BLE001
                    data = None
                if data and data.get("level"):
                    results[name] = data
                    cache[(uid or username, name)] = data
            if results:
                self._post(("levels_ready", kind, profile, forums, results))

        threading.Thread(target=worker, daemon=True).start()


    def _open_current_profile(self) -> None:
        selection = self.people_tree.selection()
        if not selection:
            return
        try:
            user = self.people_users[int(selection[0])]
        except (IndexError, ValueError):
            return
        username = user.get("username") or ""
        portrait = str(user.get("portrait") or "").split("?")[0]
        if username:
            webbrowser.open("https://tieba.baidu.com/home/main?un=" + quote(username))
        elif portrait:
            webbrowser.open("https://tieba.baidu.com/home/main?id=" + quote(portrait))
        else:
            messagebox.showinfo("无法打开主页", "这个用户缺少用户名和头像标识。", parent=self.root)


    def _build_love_tab(self) -> None:
        page = self.tab_love
        top = tk.Frame(page, bg="#ffffff", padx=12, pady=8)
        top.pack(fill="x")
        row = tk.Frame(top, bg="#ffffff")
        row.pack(fill="x")
        tk.Label(row, text="查询方式：", bg="#ffffff", fg="#475569", font=("Microsoft YaHei UI", 9)).pack(side="left")
        self.love_mode_var = tk.StringVar(value="用户名")
        ttk.Combobox(
            row,
            values=("用户名", "贴吧ID", "昵称"),
            textvariable=self.love_mode_var,
            width=7,
            state="readonly",
        ).pack(side="left", padx=(0, 6))
        tk.Label(row, text="接口：", bg="#ffffff", fg="#475569", font=("Microsoft YaHei UI", 9)).pack(side="left")
        self.love_api_var = tk.StringVar(value="查发言（新）")
        ttk.Combobox(
            row,
            values=("查发言（新）", "查发言（旧）"),
            textvariable=self.love_api_var,
            width=10,
            state="readonly",
        ).pack(side="left", padx=(0, 6))
        tk.Label(row, text="内容：", bg="#ffffff", fg="#475569", font=("Microsoft YaHei UI", 9)).pack(side="left")
        self.love_query_var = tk.StringVar()
        self.love_entry = None
        self.love_edit_host = tk.Frame(row, bg="#ffffff", height=28)
        self.love_edit_host.pack(side="left", fill="x", expand=True, padx=(0, 4))
        self.love_edit_host.pack_propagate(False)
        if os.name != "nt" or not self._create_love_native_edit():
            self.love_entry = ttk.Entry(self.love_edit_host, textvariable=self.love_query_var, width=24)
            self.love_entry.pack(fill="x", expand=True)
        self.love_auto_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(row, text="自动连续翻页", variable=self.love_auto_var).pack(side="left", padx=(8, 0))
        ttk.Button(row, text="抓取", command=self._start_love).pack(side="left", padx=(8, 0))
        ttk.Button(row, text="下一页", command=self._next_love).pack(side="left", padx=(4, 0))
        ttk.Button(row, text="停止", command=self._stop_love).pack(side="left", padx=(4, 0))
        ttk.Button(row, text="打开主页", command=self._open_love_profile).pack(side="left", padx=(8, 0))
        ttk.Button(row, text="打开选中链接", command=self._open_love_link).pack(side="left", padx=(8, 0))
        ttk.Button(row, text="导出已加载", command=self._export_love).pack(side="left", padx=(8, 0))
        self.love_status = tk.Label(
            top,
            text="数据来自恋文（第三方）+ 官方主题补全；非百度官方。",
            bg="#ffffff",
            fg="#8a94a6",
            font=("Microsoft YaHei UI", 8),
        )
        self.love_status.pack(anchor="w", pady=(4, 0))

        wrap = tk.Frame(page, bg="#eef1f5")
        wrap.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        self.love_tree = ttk.Treeview(
            wrap,
            columns=("time", "type", "bar", "title", "content", "floor", "link"),
            show="headings",
            height=18,
        )
        for key, title, width in (
            ("time", "日期", 120),
            ("type", "类型", 60),
            ("bar", "贴吧", 100),
            ("title", "标题", 220),
            ("content", "内容", 300),
            ("floor", "楼层", 60),
            ("link", "链接", 260),
        ):
            self.love_tree.heading(key, text=title)
            self.love_tree.column(key, width=width, anchor="center" if key in ("time", "type", "floor") else "w")
        yscroll = ttk.Scrollbar(wrap, orient="vertical", command=self.love_tree.yview)
        xscroll = ttk.Scrollbar(wrap, orient="horizontal", command=self.love_tree.xview)
        self.love_tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.love_tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        wrap.rowconfigure(0, weight=1)
        wrap.columnconfigure(0, weight=1)
        self._love_urls: Dict[str, str] = {}
        self.love_tree.bind("<Double-1>", lambda _event: self._open_love_link())
        self.love_tree.bind("<Control-c>", lambda _event: self._copy_love_link())
        self.love_menu = tk.Menu(self.love_tree, tearoff=0)
        self.love_menu.add_command(label="打开链接", command=self._open_love_link)
        self.love_menu.add_command(label="复制链接", command=self._copy_love_link)
        self.love_tree.bind(
            "<Button-3>",
            lambda event: self._popup_tree_menu(self.love_tree, self.love_menu, event),
        )
        self._love_rows: List[Dict] = []
        self._love_seen: set = set()
        self._love_page = 0
        self._love_has_next = True
        self._love_running = False
        self._love_stop = threading.Event()


    def _selected_love_url(self) -> str:
        selection = self.love_tree.selection()
        return self._love_urls.get(selection[0], "") if selection else ""


    def _open_love_profile(self) -> None:
        query = self._love_text().strip()
        if not query:
            self.love_status.configure(text="先输入用户名 / 贴吧ID / 昵称")
            return
        mode = self.love_mode_var.get()
        username = ""
        portrait = ""
        if mode == "用户名":
            username = query
        else:
            try:
                fetcher = Fetcher(cookie={}, delay=0.2, verbose=False)
                if mode == "贴吧ID":
                    info = profile_pb.fetch_profile_by_tieba_uid(fetcher, query)
                    username = str(info.get("username") or "")
                    portrait = str(info.get("portrait") or "")
                else:
                    info = resolve_user(fetcher, query, "昵称")
                    username = str(info.get("username") or "")
                    portrait = str(info.get("portrait") or "")
            except Exception as error:  # noqa: BLE001
                self.love_status.configure(text=f"解析主页失败：{error}")
                return
        if not username and not portrait:
            self.love_status.configure(text="没有解析到可打开的主页")
            return
        if username:
            webbrowser.open("https://tieba.baidu.com/home/main?un=" + quote(username))
        else:
            webbrowser.open("https://tieba.baidu.com/home/main?id=" + quote(portrait))


    def _open_love_link(self) -> None:
        url = self._selected_love_url()
        if url:
            webbrowser.open(url)


    def _copy_love_link(self) -> None:
        url = self._selected_love_url()
        if url:
            self._copy_text(url)


    def _stop_love(self) -> None:
        self._love_stop.set()
        self.love_status.configure(text="正在停止…")


    def _start_love(self) -> None:
        query = self._love_text().strip()
        if not query:
            self.love_status.configure(text="请输入贴吧ID或用户名")
            return
        if not getattr(self, "_love_disclaimer_seen", False):
            if not messagebox.askyesno(
                "恋文增强（第三方）",
                "本功能使用恋文（第三方）的公开数据。\n"
                "恋文不是百度官方接口，与本软件无隶属关系。\n"
                "第三方数据可能不完整、滞后或不准确。\n"
                "本功能不发送 Cookie/BDUSS。\n\n继续？",
                parent=self.root,
            ):
                return
            self._love_disclaimer_seen = True
        self._love_rows = []
        self._love_seen = set()
        self._love_urls = {}
        self._love_page = 0
        self._love_has_next = True
        for item in self.love_tree.get_children():
            self.love_tree.delete(item)
        mode = self.love_mode_var.get()
        uid = query if mode == "贴吧ID" else ""
        username = query if mode == "用户名" else ""
        nickname = query if mode == "昵称" else ""
        self._love_identity = (uid, username, nickname)
        self._love_use_old = self.love_api_var.get() == "查发言（旧）"
        self._load_love_page()


    def _next_love(self) -> None:
        if self._love_running or not self._love_has_next:
            return
        self._load_love_page()


    def _load_love_page(self) -> None:
        if self._love_running:
            return
        uid, username, nickname = getattr(self, "_love_identity", ("", "", ""))
        if not uid and not username and not nickname:
            return
        self._love_running = True
        self._love_stop.clear()
        next_page = self._love_page + 1
        self.love_status.configure(text=f"正在抓取第 {next_page} 页…")

        def worker():
            try:
                fetcher = Fetcher(cookie={}, delay=0.3, verbose=False)
                data = third_party.fetch_posts_page(
                    uid,
                    username,
                    next_page,
                    use_old=bool(getattr(self, "_love_use_old", False)),
                    nickname=nickname,
                )
                items = data.get("items") or []
                hrefs = "\n".join(str(item.get("href") or "") for item in items)
                parsed = parse_links(hrefs, fetcher, should_stop=self._love_stop.is_set) if hrefs.strip() else {"rows": [], "errors": []}
                rows_by_pid = {str(row.get("pid") or ""): row for row in parsed.get("rows") or []}
                rows = []
                for item in items:
                    link = third_party.parse_post_link(str(item.get("href") or ""))
                    row = rows_by_pid.get(str(link.get("pid") or ""))
                    if row is None:
                        row = {
                            "pid": link.get("pid") or "",
                            "time": "",
                            "type": "",
                            "bar": "",
                            "title": item.get("title") or "",
                            "content": item.get("content") or "",
                            "floor": "",
                            "url": item.get("href") or "",
                            "source": "恋文（第三方）",
                        }
                    else:
                        row = dict(row)
                    row["source"] = "恋文（第三方）"
                    rows.append(row)
                self._post(("love_ready", {"page": next_page, "rows": rows, "has_next": bool(data.get("has_next")), "error": ""}, ""))
            except Exception as error:  # noqa: BLE001
                self._post(("love_ready", None, str(error)))

        threading.Thread(target=worker, daemon=True).start()


    def _render_love(self, result: Dict, error: str) -> None:
        self._love_running = False
        if result is None:
            self.love_status.configure(text=f"失败：{error}")
            return
        self._love_page = int(result.get("page") or self._love_page)
        self._love_has_next = bool(result.get("has_next"))
        added = 0
        for row in result.get("rows") or []:
            key = str(row.get("pid") or row.get("url") or "")
            if key and key in self._love_seen:
                continue
            if key:
                self._love_seen.add(key)
            self._love_rows.append(row)
            iid = str(len(self._love_rows) - 1)
            self.love_tree.insert(
                "",
                "end",
                iid=iid,
                values=(
                    row.get("time") or "",
                    row.get("type") or "",
                    row.get("bar") or "",
                    row.get("title") or "",
                    row.get("content") or "",
                    row.get("floor") or "",
                    row.get("url") or "",
                ),
            )
            self._love_urls[iid] = row.get("url") or ""
            added += 1
        self.love_status.configure(
            text=(
                f"第 {self._love_page} 页 / 已加载 {len(self._love_rows)} 条 / 本页新增 {added} / "
                f"{'还有更多' if self._love_has_next else '已到末页'}；数据来自恋文（第三方）"
            )
        )
        if self.love_auto_var.get() and self._love_has_next and self._love_page < 100 and not self._love_stop.is_set():
            self.root.after(120, self._next_love)
        elif self._love_page >= 100 and self._love_has_next:
            self.love_status.configure(
                text=f"第 {self._love_page} 页 / 自动补全已到 100 页上限，可手动继续；数据来自恋文（第三方）"
            )


    def _export_love(self) -> None:
        if not self._love_rows:
            self.love_status.configure(text="没有可导出的结果")
            return
        path = filedialog.asksaveasfilename(
            parent=self.root,
            title="导出恋文增强",
            defaultextension=".xlsx",
            initialfile="恋文增强.xlsx",
            filetypes=[("Excel 表格", "*.xlsx")],
        )
        if not path:
            return
        try:
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "恋文增强"
            ws.append(["日期", "类型", "贴吧", "标题", "内容", "楼层", "链接", "数据来源"])
            for row in self._love_rows:
                ws.append(
                    [
                        row.get("time") or "",
                        row.get("type") or "",
                        row.get("bar") or "",
                        row.get("title") or "",
                        row.get("content") or "",
                        row.get("floor") or "",
                        row.get("url") or "",
                        "恋文（第三方）",
                    ]
                )
            info = wb.create_sheet("说明")
            info.append(["数据来源", "恋文（第三方）+ 官方主题补全"])
            info.append(["已加载页", f"1~{self._love_page}"])
            info.append(["已加载条数", len(self._love_rows)])
            info.append(["是否还有更多", "是" if self._love_has_next else "否"])
            wb.save(path)
            self.love_status.configure(text="已导出已加载内容")
        except Exception as error:  # noqa: BLE001
            messagebox.showerror("导出失败", str(error), parent=self.root)


    def _on_love_level_changed(self) -> None:
        enabled = bool(self.love_level_var.get())
        if enabled and not self.settings.get("love_level_enabled", False):
            if not messagebox.askyesno(
                "恋文等级补充（第三方）",
                "数据来自恋文（第三方），非百度官方。\n"
                "恋文等级可能滞后或不准确，只能查询已知贴吧；"
                "需要贴吧ID或用户名。\n"
                "恋文与本软件无隶属关系。\n\n继续？",
                parent=self.root,
            ):
                self.love_level_var.set(False)
                enabled = False
        self.settings["love_level_enabled"] = enabled
        save_settings(self.settings)
        self._profile_cache = {}
        self._lookup_cache = {}
        self._all_levels_cache = {}
        self.general_hint.configure(text="已保存")


    def _profile_hit(self, profile, extra_bars=None):
        if not self._risk_enabled():
            return "", ""
        # 先按发言判定；只有拿到完整关注列表时，关注吧才参与。
        src = profile.get("post_sources") or {}
        if src and src.get("replies_ok") is False and src.get("threads_ok") is not True:
            return "", ""
        bars = [post.get("bar") for post in profile.get("recent_posts") or [] if post.get("bar")]
        bars += [bar for bar in (extra_bars or []) if bar]
        complete_follow = bool(profile.get("follow_loaded")) and not self._hidden_follows(profile)
        if complete_follow and not profile.get("follow_error"):
            bars += [forum["name"] for forum in profile.get("followed_forums") or []]
        return classify(bars, getattr(self, "blacklist", {}))


    def _reload_profile(self) -> None:
        user = self._selected_person() if hasattr(self, "people_tree") else None
        if user and self.last_threads:
            self._load_profile(user, force=True)


    def _load_profile(
        self,
        user,
        force: bool = False,
        force_follow: bool = False,
        topic_max_pages: Optional[int] = None,
        use_page_cache: Optional[bool] = None,
    ) -> None:
        key = user["uid"] or f"name:{user['name']}"
        cached = None if (force or force_follow) else getattr(self, "_profile_cache", {}).get(key)
        if cached and (cached.get("follow_loaded") or not force_follow):
            stored_at = float(cached.get("_ui_cached_at") or 0)
            if stored_at:
                cached["_ui_cache_age_seconds"] = max(0, int(time.time() - stored_at))
            self._render_profile(user, cached)
            return
        self.profile_status.configure(text=f"正在查询 {display_text(user['name'])} …")
        self._set_profile_cache_note("")
        for tree in (self.profile_forum_tree, self.profile_post_tree):
            self._clear_loaded_tree(tree)
        count = self._post_count()
        start_date, end_date = self._date_params()
        bar_name = self._user_bar(user)
        prefer = (
            "start"
            if self.post_date_mode_var.get() == "指定日期" and "最早" in self.post_prefer_var.get()
            else "end"
        )
        topic_limit = int(topic_max_pages or getattr(self, "_profile_topic_limit", 50) or 50)
        self._profile_topic_limit = topic_limit
        cached_posts = None
        if not force and not force_follow:
            cached_posts = cached_user_posts(user["uid"], count, start_date, end_date, prefer)

        def worker():
            profile: Dict[str, Any]
            try:
                fetcher = Fetcher(cookie=self._request_cookie(), delay=0, verbose=False)
                uid = str(user.get("uid") or "")
                username = str(user.get("username") or "")
                portrait = str(user.get("portrait") or "")
                if not uid:
                    info = {}
                    if username:
                        info = resolve_user(fetcher, username, "用户名")
                    elif user.get("name"):
                        info = resolve_user(fetcher, str(user["name"]), "昵称")
                    if not info.get("uid") and not portrait and user.get("name") and username:
                        info = resolve_user(fetcher, str(user["name"]), "昵称")
                    if info.get("uid"):
                        uid = str(info.get("uid") or "")
                        username = str(info.get("username") or username)
                        portrait = str(info.get("portrait") or portrait)
                        user["uid"] = uid
                        user["username"] = username
                        user["portrait"] = portrait
                    elif not portrait:
                        profile = {"ok": False, "error": "用户名/昵称无法唯一定位"}
                        self._post(("profile_ready", key, user, profile))
                        return
                local_cached = cached_posts
                if local_cached is None and uid and not force:
                    local_cached = cached_user_posts(uid, count, start_date, end_date, prefer)
                profile = fetch_user_profile(
                    fetcher,
                    uid,
                    username,
                    portrait,
                    bar_name,
                    count,
                    start_date,
                    end_date,
                    prefer,
                    load_follow=force_follow or self.settings.get("follow_load_mode", "auto") == "auto",
                    topic_max_pages=topic_limit,
                    use_cache=(not force if use_page_cache is None else bool(use_page_cache)),
                    cached_posts=local_cached,
                )
                if profile.get("ok"):
                    self._attach_love_levels(profile)
                    official = list(profile.get("recent_posts") or [])
                    current = merge_current_threads(official, getattr(self, "last_threads", []), profile)
                    posts = official + current["posts"]
                    posts.sort(key=lambda item: int(item.get("ts") or 0), reverse=(prefer != "start"))
                    profile["recent_posts"] = posts[:count]
                    profile["_supplement"] = {
                        "official": len(official),
                        "current": len(current["posts"]),
                        "retrace": 0,
                        "failed": 0,
                        "total": len(posts),
                        "shown": len(profile["recent_posts"]),
                    }
            except Exception as error:  # noqa: BLE001
                profile = {"ok": False, "error": str(error)}
            self._post(("profile_ready", key, user, profile))

        threading.Thread(target=worker, daemon=True).start()


    def _reload_profile_follow(self) -> None:
        user = getattr(self, "_profile_user", None)
        if user:
            self._load_profile(user, force=True, force_follow=True)


    def _continue_profile_topics(self) -> None:
        user = getattr(self, "_profile_user", None)
        if not user:
            return
        choice = self._ask_choice(
            "继续抓主题",
            "主题流已达到本次上限，继续会追加请求；不会重复抓已缓存的页。",
            [
                ("再抓 50 页", "more"),
                ("一直抓到开始日期", "all"),
                ("停止", None),
            ],
            default=0,
        )
        if choice is None:
            self.profile_topic_btn.pack_forget()
            return
        limit = 100000 if choice == "all" else int(getattr(self, "_profile_topic_limit", 50)) + 50
        self._load_profile(
            user, force=True, topic_max_pages=limit, use_page_cache=True
        )


    def _open_profile_forum(self) -> None:
        selection = self.profile_forum_tree.selection()
        if not selection:
            return
        name = self._profile_forum_names.get(selection[0])
        if name:
            webbrowser.open("https://tieba.baidu.com/f?kw=" + quote(name))


    def _open_profile_post(self) -> None:
        selection = self.profile_post_tree.selection()
        if not selection:
            return
        url = self._profile_post_urls.get(selection[0])
        if url:
            webbrowser.open(url)


    def _export_profile_report(self) -> None:
        user = getattr(self, "_profile_user", None)
        profile = getattr(self, "_profile_data", None)
        if not user or not profile:
            messagebox.showinfo("还没有资料", "先点开一个参与者。", parent=self.root)
            return
        mark = None
        if self._risk_enabled():
            level, bar = self._profile_hit(profile, [self._user_bar(user)])
            if level:
                mark = {"level": {"ban": "红", "warn": "黄", "safe": "绿"}.get(level, level), "bar": bar}
        self._choose_report_path(user, profile, mark)


    def _export_profile_posts(self) -> None:
        user = getattr(self, "_profile_user", None)
        profile = getattr(self, "_profile_data", None)
        if not user or not profile:
            messagebox.showinfo("还没有资料", "先点开一个参与者。", parent=self.root)
            return
        self._export_posts_for(user, profile, *self._date_params())


    def _render_profile(self, user, profile) -> None:
        self._set_profile_cache_note("")
        forum_tree = self.profile_forum_tree
        post_tree = self.profile_post_tree
        for tree in (forum_tree, post_tree):
            self._clear_loaded_tree(tree)
        self._profile_forum_names = {}
        self._profile_post_urls = {}

        if not profile.get("ok"):
            self.profile_status.configure(
                text="查询失败：" + friendly_error(RuntimeError(str(profile.get("error") or "未知错误")))
            )
            return

        level = str(profile.get("bar_level") or "").strip()
        if level in ("0", "None"):
            level = ""
        if not level:
            level = str(user.get("level") or "").strip()
            if level in ("0", "None"):
                level = ""

        if level:
            self._apply_level_to_threads(user, level)
        profile_tieba = str(profile.get("tieba_uid") or "")
        had_tieba = bool(user.get("tieba_uid"))
        if profile_tieba:
            user["tieba_uid"] = profile_tieba
            iid = user.get("_iid")
            if iid is not None:
                try:
                    values = list(self.people_tree.item(iid, "values"))
                    values[3] = profile_tieba
                    self.people_tree.item(iid, values=values)
                except tk.TclError:
                    pass
            if not had_tieba and getattr(self, "last_output", None):
                self._append_log("身份已更新，需要重新导出 Excel。", "muted")
                self.profile_status.configure(text="身份已更新，需要重新导出")
        if level and not user.get("level"):
            user["level"] = level
            iid = user.get("_iid")
            if iid is not None:
                try:
                    values = list(self.people_tree.item(iid, "values"))
                    values[5] = f"Lv.{level}"
                    self.people_tree.item(iid, values=values)
                except tk.TclError:
                    pass

        if profile.get("user_id") and not user.get("uid"):
            user["uid"] = str(profile.get("user_id") or "")
        if profile.get("username") and not user.get("username"):
            user["username"] = str(profile.get("username") or "")
        if profile.get("portrait") and not user.get("portrait"):
            user["portrait"] = str(profile.get("portrait") or "")
        self._profile_user = user
        self._profile_data = profile
        self._set_person_mark_row(user)
        posts = profile.get("recent_posts") or []
        combined = self._merge_forum_lists(profile, self._user_bar(user), level or "")
        self.profile_report_btn.pack(side="left", padx=(0, 6))
        self.profile_posts_btn.pack(side="left", padx=(0, 6))
        sources = profile.get("post_sources") or {}
        if sources.get("topic_truncated"):
            topic_text = f"继续抓主题（已 {sources.get('topic_pages') or 50} 页）"
            self.profile_topic_btn.configure(text=topic_text, width=max(12, len(topic_text) + 2))
            self.profile_topic_btn.pack(side="right", padx=(0, 4))
        else:
            self.profile_topic_btn.pack_forget()
        if profile.get("follow_loaded") is False:
            self.profile_follow_btn.pack(side="left", padx=(0, 6))
        else:
            self.profile_follow_btn.pack_forget()
        if self._follow_incomplete(profile):
            self.profile_probe_btn.pack(side="left", padx=(0, 6))
            self.profile_probe_btn.configure(state="normal", text="探测候选关注")
        else:
            self.profile_probe_btn.pack_forget()

        for index, forum in enumerate(combined):
            iid = str(index)
            forum_tree.insert(
                "",
                "end",
                iid=iid,
                values=(
                    forum["name"],
                    f"Lv.{forum['level']}" if forum.get("level") else "-",
                    forum.get("level_source") or "-",
                    forum.get("source") or "",
                ),
                tags=self._row_tags(forum["name"], "forum"),
            )
            self._profile_forum_names[iid] = forum["name"]
        self._request_love_levels(profile, combined, forum_tree, self._profile_forum_names, "profile")

        has_supplement = any(str(post.get("source") or "") not in ("", "官方") for post in posts)
        post_tree["displaycolumns"] = (
            ("time", "type", "bar", "title", "content", "source")
            if has_supplement
            else ("time", "type", "bar", "title", "content")
        )
        for index, post in enumerate(posts):
            iid = str(index)
            post_tree.insert(
                "",
                "end",
                iid=iid,
                values=(
                    post.get("time") or "",
                    post.get("type") or "",
                    post.get("bar") or "",
                    post.get("title") or "",
                    post.get("content") or "",
                    post.get("source") or "官方",
                ),
                tags=self._row_tags(post.get("bar"), "posts"),
            )
            self._profile_post_urls[iid] = post.get("url") or ""
        forum_note = self._forum_note(profile, self._user_bar(user), level or "")
        if self.mode_var.get() != "guest" and profile.get("follow_loaded") is False:
            forum_note += "（未加载公开关注）"
        self.profile_forum_note.configure(text=forum_note)

        # 本地标记：优先沿用批量结果，避免详情重算把 ⚠ 或等级改掉。
        if not self._risk_enabled():
            hit_level, hit_bar = "", ""
            hidden_follows = False
        elif user.get("mark_level"):
            hit_level = str(user.get("mark_level") or "")
            hit_bar = str(user.get("mark_bar") or "")
            hidden_follows = bool(user.get("mark_hidden")) and self.mode_var.get() != "guest"
        elif self.mode_var.get() == "guest":
            hit_level, hit_bar = classify([self._user_bar(user)], getattr(self, "blacklist", {}))
            hidden_follows = False
        else:
            bars = [post.get("bar") for post in posts if post.get("bar")]
            bars += [forum["name"] for forum in combined]
            hit_level, hit_bar = classify(bars, getattr(self, "blacklist", {}))
            hidden_follows = self._hidden_follows(profile)
            if hidden_follows and hit_level != "ban":
                if hit_level != "warn":
                    hit_bar = "关注列表隐藏"
                hit_level = "warn"
        label = LEVEL_LABEL.get(hit_level, "")
        for traffic_key, (traffic_label, traffic_color) in getattr(self, "traffic_labels", {}).items():
            traffic_label.configure(fg=traffic_color if traffic_key == hit_level else "#d1d5db")
        risk_reason = (
            "关注列表隐藏"
            if self.mode_var.get() != "guest" or hidden_follows
            else "关注列表未加载"
        )
        risk = f" · {risk_reason}，未命中警告时标黄" if hidden_follows else ""
        try:
            self.traffic_text.configure(
                text=(
                    (f"{label}：{hit_bar}" if hit_bar and hit_bar != "关注列表隐藏" else label)
                    if hit_level
                    else "未命中任何名单"
                )
                + risk
            )
        except tk.TclError:
            pass
        iid = user.get("_iid")
        if iid is not None:
            try:
                values = list(self.people_tree.item(iid, "values"))
                if user.get("mark_level"):
                    values[0] = self._person_mark_cell(user)
                elif hit_level:
                    values[0] = f"{marker(hit_level)} {label}" + (" ⚠" if hidden_follows else "")
                else:
                    values[0] = "⚪"
                self.people_tree.item(
                    iid,
                    values=values,
                    tags=(hit_level,) if hit_level and self.settings.get("color_people", True) else (),
                )
            except tk.TclError:
                pass

        status = (
            f"{display_text(user['name'])} | {user['username'] or '-'} | 贴吧ID {self._tieba_uid_cell(user)} | "
            f"用户ID {user['uid'] or '-'} | Lv.{level or '-'} | 发言 {len(posts)} | 关注 {len(combined)}"
        )
        if profile.get("fans") is not None:
            status += f" | 粉丝 {profile['fans']}"
        if profile.get("concern") is not None:
            status += f" | 关注数 {profile['concern']}"
        if profile.get("post_num"):
            status += f" | 发帖 {profile['post_num']}"
        if hit_level:
            status += f" · 命中{label}：{hit_bar}"
        if hidden_follows:
            status += f" · ⚠ {risk_reason}，未命中警告时标黄"
        note = self._post_source_note(profile)
        if note:
            status += " · " + note
        self.profile_status.configure(text=status)
        self._set_profile_cache_note(self._cache_age_note(profile))


class ToolsMixin:
    """日志/测试中心/使用说明/数据清理方法，从 App 拆出。"""

    _full_cookie_cache: Any
    blacklist: Any
    last_threads: Any
    last_output: Any
    last_outputs: Any
    _last_lookup_profile: Any
    _login_source: Any
    _profile_data: Any
    _profile_user: Any
    _webview_cookie_session: Any
    _usage_window: Any

    def __getattr__(self, name: str) -> Any:
        raise AttributeError(name)

    def _logs_dir(self) -> str:
        return os.path.join(APP_DIR, "logs")


    def _log_files(self) -> List[str]:
        folder = self._logs_dir()
        if not os.path.isdir(folder):
            return []
        paths = [
            os.path.join(folder, name)
            for name in os.listdir(folder)
            if name.startswith("app-") and name.endswith(".log")
        ]
        paths.sort(key=lambda path: os.path.getmtime(path), reverse=True)
        return paths


    def _refresh_log_summary(self) -> None:
        label = getattr(self, "log_summary_label", None)
        if label is None:
            return
        files = self._log_files()
        total = 0
        for path in files:
            try:
                total += os.path.getsize(path)
            except OSError:
                continue
        if not files:
            label.configure(text="日志：暂无")
        else:
            label.configure(text=f"日志：{len(files)} 个 / {total / 1024:.1f} KB")


    def _open_logs_dir(self) -> None:
        folder = self._logs_dir()
        try:
            os.makedirs(folder, exist_ok=True)
        except OSError as error:
            messagebox.showerror("打开失败", str(error), parent=self.root)
            return
        self._open_folder(folder)


    def _delete_logs(self) -> None:
        files = self._log_files()
        if not files:
            messagebox.showinfo("没有日志", "当前没有可删除的日志。", parent=self.root)
            return
        if not messagebox.askyesno(
            "删除日志",
            f"删除全部 {len(files)} 个日志文件？\n不会改变“保存运行日志”开关。",
            parent=self.root,
            default=messagebox.NO,
        ):
            return
        deleted = 0
        for path in files:
            try:
                os.remove(path)
                deleted += 1
            except OSError:
                continue
        self._refresh_log_summary()
        self._append_log(f"已删除日志：{deleted} 个", "muted")


    def _view_logs(self) -> None:
        files = self._log_files()
        if not files:
            messagebox.showinfo(
                "没有日志", "还没有日志。先开启“保存运行日志”并运行一次抓取。", parent=self.root
            )
            return
        win = tk.Toplevel(self.root)
        win.title("运行日志")
        win.transient(self.root)
        win.geometry("820x560")
        win.minsize(620, 400)
        win.configure(bg="#eef1f5")
        body = tk.Frame(win, bg="#eef1f5", padx=10, pady=10)
        body.pack(fill="both", expand=True)

        left = tk.Frame(body, bg="#eef1f5", width=210)
        left.pack(side="left", fill="y", padx=(0, 8))
        left.pack_propagate(False)
        tk.Label(left, text="日志文件", bg="#eef1f5", fg="#475569", anchor="w").pack(fill="x")
        listbox = tk.Listbox(left, exportselection=False, activestyle="none")
        listbox.pack(fill="both", expand=True, pady=(4, 0))

        right = tk.Frame(body, bg="#eef1f5")
        right.pack(side="left", fill="both", expand=True)
        text = tk.Text(right, wrap="none", state="disabled", bg="#ffffff", fg="#111827")
        y_scroll = ttk.Scrollbar(right, orient="vertical", command=text.yview)
        x_scroll = ttk.Scrollbar(right, orient="horizontal", command=text.xview)
        text.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        text.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        right.rowconfigure(0, weight=1)
        right.columnconfigure(0, weight=1)

        footer = tk.Frame(win, bg="#eef1f5", padx=10)
        footer.pack(fill="x", pady=(0, 10))

        def selected_path() -> str:
            selection = listbox.curselection()
            return files[int(selection[0])] if selection else ""

        def show_file(path: str) -> None:
            if not path:
                return
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as handle:
                    content = handle.read(512 * 1024)
            except OSError as error:
                content = f"读取失败：{error}"
            text.configure(state="normal")
            text.delete("1.0", "end")
            text.insert("end", content)
            text.configure(state="disabled")

        def on_select(_event=None) -> None:
            show_file(selected_path())

        def delete_current() -> None:
            path = selected_path()
            if not path:
                return
            if not messagebox.askyesno(
                "删除当前日志",
                f"删除 {os.path.basename(path)}？",
                parent=win,
                default=messagebox.NO,
            ):
                return
            try:
                os.remove(path)
            except OSError as error:
                messagebox.showerror("删除失败", str(error), parent=win)
                return
            self._refresh_log_summary()
            reload_list()

        def reload_list() -> None:
            listbox.delete(0, "end")
            current_files = self._log_files()
            files[:] = current_files
            for path in files:
                listbox.insert("end", os.path.basename(path))
            if files:
                listbox.selection_set(0)
                show_file(files[0])
            else:
                text.configure(state="normal")
                text.delete("1.0", "end")
                text.insert("end", "没有日志。")
                text.configure(state="disabled")

        for path in files:
            listbox.insert("end", os.path.basename(path))
        listbox.selection_set(0)
        show_file(files[0])
        listbox.bind("<<ListboxSelect>>", on_select)

        ttk.Button(footer, text="打开日志目录", command=self._open_logs_dir, width=12).pack(side="left")
        ttk.Button(footer, text="删除当前日志", command=delete_current, width=12).pack(
            side="left", padx=(6, 0)
        )
        ttk.Button(footer, text="关闭", command=win.destroy, width=8).pack(side="right")
        self._center_toplevel(win, 820, 560)
        try:
            self._theme_widget_tree(win)
        except Exception:  # noqa: BLE001
            pass


    def _on_log_toggle(self) -> None:
        self.settings["log_enabled"] = bool(self.log_enabled_var.get())
        save_settings(self.settings)


    def _show_test_center(self, _event=None) -> None:
        if self._network_task_active() or getattr(self, "_active_choice_window", None):
            return
        win = tk.Toplevel(self.root)
        win.title("测试中心")
        win.configure(bg="#eef1f5")
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        tk.Label(
            win,
            text="测试中心（不会写设置、不会覆盖结果）",
            bg="#eef1f5",
            fg="#1f2937",
            font=("Microsoft YaHei UI", 11, "bold"),
        ).pack(anchor="w", padx=16, pady=(14, 4))
        tk.Label(
            win,
            text="真实百度验证码无法自动触发；模拟测试不联网。",
            bg="#eef1f5",
            fg="#64748b",
            font=("Microsoft YaHei UI", 8),
        ).pack(anchor="w", padx=16, pady=(0, 10))

        def run(action) -> None:
            win.grab_release()
            win.destroy()
            action()

        ttk.Button(
            win,
            text="模拟游客验证码（可继续操作）",
            command=lambda: run(lambda: self._show_test_captcha_dialog("guest")),
            width=38,
        ).pack(fill="x", padx=16, pady=2)
        ttk.Button(
            win,
            text="模拟登录：完整 Cookie（可继续操作）",
            command=lambda: run(lambda: self._show_test_captcha_dialog("cookie")),
            width=38,
        ).pack(fill="x", padx=16, pady=2)
        ttk.Button(
            win,
            text="模拟登录：仅 BDUSS（可继续操作）",
            command=lambda: run(lambda: self._show_test_captcha_dialog("bduss")),
            width=38,
        ).pack(fill="x", padx=16, pady=2)
        ttk.Button(
            win,
            text="运行离线检查（PASS/FAIL）",
            command=lambda: run(self._show_offline_test_results),
            width=38,
        ).pack(fill="x", padx=16, pady=2)
        ttk.Button(
            win,
            text="测试内置验证窗口（不读取 Cookie）",
            command=lambda: run(self._test_embedded_verification),
            width=38,
        ).pack(fill="x", padx=16, pady=2)
        ttk.Button(
            win,
            text="打开内置登录（仅登录模式）",
            command=lambda: run(self._open_internal_login),
            width=38,
        ).pack(fill="x", padx=16, pady=2)
        ttk.Button(
            win,
            text="关闭",
            command=win.destroy,
            width=38,
        ).pack(fill="x", padx=16, pady=(10, 14))
        win.update_idletasks()
        width = max(420, win.winfo_reqwidth())
        height = max(240, win.winfo_reqheight())
        x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - width) // 2)
        y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - height) // 2)
        win.geometry(f"{width}x{height}+{x}+{y}")
        win.lift()
        win.focus_force()


    def _show_offline_test_results(self) -> None:
        results = self._offline_checks()
        win = tk.Toplevel(self.root)
        win.title("离线检查结果")
        win.configure(bg="#eef1f5")
        win.transient(self.root)
        win.grab_set()
        body = tk.Frame(win, bg="#eef1f5", padx=12, pady=10)
        body.pack(fill="both", expand=True)
        text = tk.Text(body, width=78, height=max(8, len(results) + 4), wrap="word")
        text.pack(fill="both", expand=True)
        text.insert("end", "测试项\t结果\t说明\n")
        for name, status, detail in results:
            text.insert("end", f"{name}\t{status}\t{detail}\n")
        text.configure(state="disabled")
        ttk.Button(body, text="关闭", command=win.destroy).pack(anchor="e", pady=(8, 0))
        win.update_idletasks()
        width = max(620, win.winfo_reqwidth())
        height = max(300, win.winfo_reqheight())
        x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - width) // 2)
        y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - height) // 2)
        win.geometry(f"{width}x{height}+{x}+{y}")
        win.lift()


    def _show_test_captcha_dialog(self, scenario: str) -> None:
        """测试模式：展示真实按钮，用假 Cookie 走后续界面，不联网。"""
        if scenario == "guest":
            message = (
                "当前是游客模式（测试）。\n"
                "真实流程会读取匿名 Cookie，不读取 BDUSS/STOKEN。"
            )
            choices = [
                ("模拟读取匿名 Cookie", "guest_embed"),
                ("改用登录模式（完整 Cookie）", "login"),
                ("停止本次", "stop"),
            ]
        elif scenario == "bduss":
            cached = bool(getattr(self, "_full_cookie_cache", ""))
            message = (
                "本次实际只发送了 BDUSS（测试）。\n"
                + ("有完整 Cookie 缓存，可改用。" if cached else "当前没有完整 Cookie 缓存。")
            )
            choices = [("模拟读取完整 Cookie", "full_embed")]
            if cached:
                choices.append(("改用已保存的完整 Cookie", "cached"))
            choices.extend([("手动粘贴完整 Cookie", "manual"), ("停止本次", "stop")])
        else:
            message = (
                "当前使用完整 Cookie（测试）。\n"
                "真实流程会读取完整 Cookie，并检查 BDUSS/TIEBA_SID。"
            )
            choices = [
                ("模拟读取完整 Cookie", "full_embed"),
                ("手动粘贴完整 Cookie", "manual"),
                ("停止本次", "stop"),
            ]
        choice = self._ask_choice("被安全验证拦截（测试模式）", message, choices)
        if not choice or choice == "stop":
            return
        if choice == "login":
            self._set_mode("login", persist=False)
            self.login_method_var.set("cookie")
            self._on_login_method_changed()
            self._append_log("测试：已切到登录模式，未写设置。", "muted")
            self.cookie_entry.focus_set()
            return
        if choice == "manual":
            self._set_mode("login", persist=False)
            self.login_method_var.set("cookie")
            self._on_login_method_changed()
            self.cookie_entry.focus_set()
            self._append_log("测试：请手动粘贴完整 Cookie。", "muted")
            return
        if choice == "cached":
            self._set_mode("login", persist=False)
            self.login_method_var.set("cookie")
            self.cookie_var.set(self._full_cookie_cache)
            self._on_login_method_changed()
            self.cookie_entry.focus_set()
            self._append_log("测试：已改用缓存的完整 Cookie。", "muted")
            return
        self._set_mode("guest" if choice == "guest_embed" else "login", persist=False)
        if choice == "guest_embed":
            self._guest_cookie = {"BAIDUID": "TEST_GUEST"}
            self._append_log("测试：已写入假匿名 Cookie，保持游客。", "muted")
            messagebox.showinfo("测试结果", "已模拟匿名 Cookie；不会联网。", parent=self.root)
            return
        fake = {"BDUSS": "TEST_BDUSS", "TIEBA_SID": "TEST_TIEBA_SID", "STOKEN": "TEST_STOKEN"}
        text = "; ".join(f"{key}={value}" for key, value in fake.items())
        self._login_source = "测试"
        self._full_cookie_cache = text
        if scenario == "bduss":
            self.login_method_var.set("bduss")
            self.bduss_var.set(fake["BDUSS"])
            self.cookie_var.set(f"BDUSS={fake['BDUSS']}")
        else:
            self.login_method_var.set("cookie")
            self.cookie_var.set(text)
            self.bduss_var.set(fake["BDUSS"])
        self._on_login_method_changed()
        self._webview_cookie_session = False
        self._append_log("测试：已写入假完整 Cookie，未联网、未写设置。", "muted")
        messagebox.showinfo(
            "测试结果",
            "已模拟完整 Cookie。\n不会预检、不会联网、不会自动开始抓取。",
            parent=self.root,
        )


    def _test_embedded_verification(self) -> None:
        if getattr(self, "_webview_running", False):
            messagebox.showinfo("正在测试", "内置验证窗口已经打开。", parent=self.root)
            return
        if not webview2.available():
            messagebox.showinfo(
                "内置验证窗口不可用",
                "未检测到 pywebview 或 WebView2 Runtime。\n"
                "仍可使用「转到登录模式（完整 Cookie）」。",
                parent=self.root,
            )
            return
        self._webview_running = True
        self.status_var.set("正在打开内置验证窗口（测试）…")
        center = self._webview_center()
        url = (getattr(self, "_current_urls", []) or ["https://tieba.baidu.com/"])[0]

        def worker():
            try:
                webview2.open_window(url, center=center)
                self._post(("webview_test_done", ""))
            except Exception as error:  # noqa: BLE001
                self._post(("webview_test_done", str(error)))

        threading.Thread(target=worker, daemon=True).start()


    def _copy_diagnostics(self) -> None:
        lines = [
            f"version={VERSION}",
            f"data_dir={APP_DIR}",
            f"follow_enhance={bool(self.settings.get('follow_enhance'))}",
            f"save_result={bool(self.settings.get('save_result'))}",
            f"encrypt_result={bool(self.settings.get('encrypt_result'))}",
            f"risk_enabled={bool(self.settings.get('risk_enabled'))}",
            f"page_size={self.settings.get('page_size')}",
        ]
        self.root.clipboard_clear()
        self.root.clipboard_append(chr(10).join(lines))
        self.status_var.set("诊断信息已复制（不含 Cookie/昵称/正文）")


    def _open_usage_window(self) -> None:
        existing = getattr(self, "_usage_window", None)
        try:
            if existing is not None and existing.winfo_exists():
                existing.lift()
                existing.focus_set()
                return
        except tk.TclError:
            pass
        win = tk.Toplevel(self.root)
        self._usage_window = win
        win.title("使用说明")
        win.configure(bg="#eef1f5")
        win.geometry("780x640")
        win.minsize(600, 420)
        win.transient(self.root)
        frame = tk.Frame(win, bg="#eef1f5", padx=10, pady=8)
        frame.pack(fill="both", expand=True)
        top = tk.Frame(frame, bg="#eef1f5")
        top.pack(fill="x", pady=(0, 6))
        tk.Label(
            top,
            text="使用说明",
            bg="#eef1f5",
            fg="#1F4E79",
            font=("Microsoft YaHei UI", 13, "bold"),
        ).pack(side="left")
        ttk.Button(top, text="关闭", command=win.destroy, width=8).pack(side="right")
        text = tk.Text(
            frame,
            wrap="word",
            font=("Microsoft YaHei UI", 9),
            relief="flat",
            padx=10,
            pady=8,
            bg="#fbfcfe",
            state="disabled",
        )
        scroll = ttk.Scrollbar(frame, orient="vertical", command=text.yview)
        text.configure(yscrollcommand=scroll.set)
        text.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        content = ""
        try:
            path = os.path.join(resource_dir(), "使用说明.txt")
            if os.path.isfile(path):
                with open(path, "r", encoding="utf-8", errors="replace") as handle:
                    content = handle.read()
        except OSError:
            content = ""
        shortcuts = (
            "快捷键\n"
            "  F1           打开使用说明\n"
            "  Ctrl+Enter   从链接输入框开始抓取\n"
            "  Ctrl+C       复制选中行\n"
            "  Esc          关闭当前弹窗\n"
            "  双击         打开选中项\n\n"
        )
        text.configure(state="normal")
        text.insert("1.0", shortcuts + (content or "使用说明文件未找到。"))
        text.configure(state="disabled")
        self._theme_widget_tree(win)
        self._center_toplevel(win, 780, 640)
        win.bind("<Escape>", lambda _event: win.destroy())
        win.protocol("WM_DELETE_WINDOW", win.destroy)
        win.focus_set()


    def _clear_all_data(self) -> None:
        if not messagebox.askyesno(
            "清空软件数据（不会删除导出文件）",
            "会删除：最近结果、保存的 BDUSS、本地标记、设置、临时文件、日志、"
            "内置 WebView2 登录态。\n"
            "不会删除：导出的 Excel、Markdown、你保存到其他位置的文件，"
            "以及系统浏览器里的百度登录。\n\n继续？",
            parent=self.root,
        ):
            return
        if self._network_task_active() or getattr(self, "_webview_running", False):
            messagebox.showinfo("请稍后", "抓取、补全或内置窗口进行中，暂时不能清空。", parent=self.root)
            return
        if simpledialog.askstring("确认清空", "输入“清空”后继续：", parent=self.root) != "清空":
            return
        if not webview2.clear_login_profile():
            messagebox.showerror(
                "清空失败",
                "WebView2 登录目录正在使用，无法删除。请关闭内置窗口后重试。",
                parent=self.root,
            )
            return
        clear_stored_result()
        clear_post_cache()
        post_cache().close()
        try:
            for path in stored_old_files():
                safe_remove_file(path)
            for path in (os.path.join(APP_DIR, "cookie.bin"), SETTINGS_PATH):
                safe_remove_file(path)
            safe_remove_tree(os.path.join(APP_DIR, "logs"))
            temp_dir = os.path.join(os.environ.get("TEMP", ""), "TiebaReplyStats")
            if os.path.isdir(temp_dir):
                safe_remove_tree(temp_dir, [temp_dir])
        except (OSError, ValueError) as error:
            post_cache().reopen()
            messagebox.showerror("清空失败", f"路径校验失败，已停止：{error}", parent=self.root)
            return
        post_cache().reopen()
        clear_cookie()
        save_blacklist({"safe": [], "warn": [], "ban": []})
        self._on_mark_rules_changed()
        save_settings({})
        self.settings = load_settings()
        self.blacklist = {"safe": [], "warn": [], "ban": []}
        self.last_threads = []
        self.last_output = None
        self.cookie_var.set("")
        self.bduss_var.set("")
        self._guest_cookie = {}
        self._full_cookie_cache = ""
        self._webview_cookie_session = False
        self._last_lookup_profile = None
        self._profile_user = None
        self._profile_data = None
        self._lookup_cache.clear()
        self._profile_cache.clear()
        self._apply_settings_to_ui()
        self._refresh_webview_status()
        self._set_lookup_text("软件数据已清空。")
        self.open_btn.configure(state="disabled")
        for _key, editor in getattr(self, "blacklist_editors", {}).items():
            editor.delete("1.0", "end")
        self._update_grid_selector()
        self._refresh_people()
        self._refresh_grid()
        self._refresh_old_files_label()
        self._refresh_log_summary()
        self._append_log("软件数据已清空；导出文件未受影响。", "ok")


    def _write_log(self, event: str, **fields) -> None:
        if not self.settings.get("log_enabled", False):
            return
        try:
            log_dir = os.path.join(APP_DIR, "logs")
            os.makedirs(log_dir, exist_ok=True)
            path = os.path.join(log_dir, f"app-{datetime.now():%Y%m%d}.log")
            record = {"time": datetime.now().isoformat(timespec="seconds"), "event": event}
            record.update({key: value for key, value in fields.items() if value not in (None, "")})
            with open(path, "a", encoding="utf-8") as handle:
                handle.write(json.dumps(record, ensure_ascii=False) + chr(10))
            files = sorted(
                (os.path.join(log_dir, name) for name in os.listdir(log_dir) if name.startswith("app-")),
                key=os.path.getmtime,
                reverse=True,
            )
            for old in files[7:]:
                try:
                    os.remove(old)
                except OSError:
                    pass
            files = files[:7]
            total = 0
            sizes = {}
            for item in files:
                try:
                    sizes[item] = os.path.getsize(item)
                except OSError:
                    sizes[item] = 0
                total += sizes[item]
            for old in reversed(files):
                if total <= 10 * 1024 * 1024:
                    break
                if old == path:
                    continue
                try:
                    os.remove(old)
                    total -= sizes.get(old, 0)
                except OSError:
                    pass
        except OSError:
            pass


    def _log_from_thread(self, message: str) -> None:
        self._post(("log", message, None))


    def _append_log(self, text: str, tag: str | None = None) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", text + "\n", tag or ())
        self.log_text.see("end")
        self.log_text.configure(state="disabled")


class LoginMixin:
    """登录 / Cookie / WebView2 相关方法，从 App 拆出。"""

    _login_probe_token: Any
    _guest_cookie: Any
    _last_webview_match: Any

    def __getattr__(self, name: str) -> Any:
        raise AttributeError(name)

    def _on_login_method_changed(self) -> None:
        method = self.login_method_var.get()
        if hasattr(self, "_refresh_webview_status"):
            self._refresh_webview_status()
        if not getattr(self, "_login_fields_expanded", True):
            return
        if method == "cookie":
            self.bduss_frame.grid_remove()
            self.cookie_frame.grid()
        else:
            self.cookie_frame.grid_remove()
            self.bduss_frame.grid()


    def _extract_bduss(self, raw: str) -> str:
        text = (raw or "").strip()
        if not text:
            return ""
        cookies = parse_cookie_input(text)
        if cookies.get("BDUSS"):
            return str(cookies["BDUSS"]).strip()
        if text.upper().startswith("BDUSS="):
            return text.split("=", 1)[1].strip()
        if ";" not in text and " " not in text:
            return text
        return ""


    def _paste_bduss(self) -> None:
        try:
            text = self.root.clipboard_get()
        except tk.TclError:
            return
        value = self._extract_bduss(text)
        if not value:
            messagebox.showwarning("不是 BDUSS", "剪贴板里没有识别到 BDUSS。", parent=self.root)
            return
        self.login_method_var.set("bduss")
        self.bduss_var.set(value)
        self._apply_bduss()


    def _apply_bduss(self) -> None:
        value = self._extract_bduss(self.bduss_var.get())
        if not value:
            messagebox.showwarning("需要 BDUSS", "请填写 BDUSS。", parent=self.root)
            self.bduss_entry.focus_set()
            return
        self.bduss_var.set(value)
        self.cookie_var.set(f"BDUSS={value}")
        self._full_cookie_cache = ""
        self._login_source = "手动粘贴"
        self._last_webview_match = None
        self._login_probe_token += 1
        if self.remember_bduss_var.get():
            if save_cookie(self.cookie_var.get()):
                self.status_var.set("BDUSS 已加密保存")
            else:
                self.status_var.set("BDUSS 保存失败，本次仍可使用")
        else:
            clear_cookie()
            self.status_var.set("BDUSS 仅本次使用")
        self._refresh_webview_status()
        self._append_log("登录方式：仅 BDUSS", "muted")


    def _cookie_completeness(self, cookies: Dict[str, str]) -> List[str]:
        missing: List[str] = []
        if not str(cookies.get("STOKEN") or "").strip():
            missing.append("STOKEN（关注检测可能不完整）")
        if not str(cookies.get("TIEBAUID") or "").strip() and not str(
            cookies.get("TIEBA_SID") or ""
        ).strip():
            missing.append("TIEBAUID/TIEBA_SID（贴吧身份字段可能缺失）")
        if not str(cookies.get("BAIDUID") or "").strip():
            missing.append("BAIDUID（匿名标识缺失）")
        return missing


    def _start_cookie_probe(self, cookies: Dict[str, str], token: int) -> None:
        if self._network_task_active():
            return
        self.status_var.set("正在预检登录信息…")
        missing = self._cookie_completeness(cookies)

        def worker() -> None:
            warning = ""
            logged = False
            try:
                probe = Fetcher(cookie=cookies, delay=0, verbose=False)
            except Exception as error:  # noqa: BLE001
                self._post(("manual_cookie_probe", token, f"预检初始化失败：{error}", False, missing))
                return
            try:
                probe.get("https://tieba.baidu.com/")
                info = probe.get("https://tieba.baidu.com/f/user/json_userinfo")
                try:
                    payload = json.loads(info.text)
                except (TypeError, ValueError):
                    warning = "预检返回无法解析，未能确认登录态"
                else:
                    data = payload.get("data") if isinstance(payload, dict) else None
                    if not isinstance(data, dict):
                        data = payload if isinstance(payload, dict) else {}
                    logged = bool(
                        data.get("user_id")
                        or data.get("user_name")
                        or data.get("name")
                    )
                    if not logged:
                        warning = "预检未确认登录态，Cookie 可能已失效或只有 BDUSS"
            except CaptchaError:
                warning = "预检仍被验证码拦截，Cookie 可能不可用"
            except Exception as error:  # noqa: BLE001
                warning = f"预检失败：{error}"
            finally:
                probe.close()
            self._post(("manual_cookie_probe", token, warning, logged, missing))

        threading.Thread(target=worker, daemon=True).start()


    def _apply_cookie(self) -> None:
        raw = self.cookie_var.get().strip()
        if not raw:
            messagebox.showwarning("需要 Cookie", "请粘贴完整 Cookie。", parent=self.root)
            self.cookie_entry.focus_set()
            return
        if not webview2.has_login_cookie(parse_cookie_input(raw)):
            messagebox.showwarning(
                "Cookie 不完整",
                "完整 Cookie 里没有 BDUSS/TIEBA_SID，无法确认登录态。",
                parent=self.root,
            )
            return
        self._full_cookie_cache = raw
        self._login_source = "手动粘贴"
        self._last_webview_match = None
        self.login_method_var.set("cookie")
        self._on_login_method_changed()
        if self.remember_cookie_var.get():
            if save_cookie(raw):
                self.status_var.set("完整 Cookie 已加密保存")
            else:
                self.status_var.set("完整 Cookie 保存失败，本次仍可使用")
        else:
            clear_cookie()
            self.status_var.set("完整 Cookie 仅本次使用")
        self._refresh_webview_status()
        self._append_log("登录方式：完整 Cookie", "muted")
        self._login_probe_token += 1
        self._start_cookie_probe(parse_cookie_input(raw), self._login_probe_token)


    def _remember_current_login(self) -> bool:
        if self.login_method_var.get() == "bduss":
            return bool(self.remember_bduss_var.get())
        return bool(self.remember_cookie_var.get())


    def _current_cookie_text(self) -> str:
        if self.login_method_var.get() == "bduss":
            value = self._extract_bduss(self.bduss_var.get())
            return f"BDUSS={value}" if value else ""
        return self.cookie_var.get().strip()


    def _sync_bduss_from_cookie(self, *_args) -> None:
        """完整 Cookie 变化时，同步提取 BDUSS，便于切换到仅 BDUSS。"""
        raw = self.cookie_var.get().strip()
        if not raw:
            return
        cookies = parse_cookie_input(raw)
        value = cookies.get("BDUSS", "")
        if value and set(cookies) != {"BDUSS"}:
            self._full_cookie_cache = raw
        if value and self.bduss_var.get() != value:
            self.bduss_var.set(value)


    def _sync_login_from_cookie(self) -> None:
        raw = self.cookie_var.get().strip()
        if not raw:
            return
        cookies = parse_cookie_input(raw)
        if cookies.get("BDUSS") and set(cookies) == {"BDUSS"}:
            self.login_method_var.set("bduss")
            self.bduss_var.set(str(cookies["BDUSS"]))
            self.remember_bduss_var.set(True)
        else:
            self.login_method_var.set("cookie")
            self.remember_cookie_var.set(True)
            self._full_cookie_cache = raw


    def _clear_cookie(self) -> None:
        self.cookie_var.set("")
        self.bduss_var.set("")
        self._full_cookie_cache = ""
        self._login_source = ""
        self._last_webview_match = None
        self._login_probe_token += 1
        if clear_cookie():
            self._append_log("已清除本机保存的 Cookie。", "muted")
        else:
            self._append_log("已清空输入框（没有可删除的已保存 Cookie）。", "muted")


    def _paste_cookie(self) -> None:
        try:
            text = self.root.clipboard_get()
        except tk.TclError:
            messagebox.showinfo(
                "剪贴板是空的",
                "先在浏览器里复制 Cookie，再点这个按钮。\n\n"
                "复制方法：F12 → 网络/Network → 刷新 → 点第一个帖子请求 → "
                "标头/Headers → 请求标头里的 Cookie，复制冒号后面的整段。",
                parent=self.root,
            )
            return
        text = text.strip()
        if text.lower().startswith("cookie:"):
            text = text.split(":", 1)[1].strip()
        # 只保留有效的 k=v 片段
        keep = [part.strip() for part in re.split(r"[;\n\r]+", text) if "=" in part]
        if not keep:
            messagebox.showwarning(
                "剪贴板内容不像 Cookie", "剪贴板里没有找到 k=v 形式的 Cookie，请重新复制。", parent=self.root
            )
            return
        self.cookie_var.set("; ".join(keep))
        self._append_log("已从剪贴板粘贴 Cookie。", "ok")


    def _show_bduss_help(self) -> None:
        messagebox.showinfo("如何获取 BDUSS", BDUSS_HELP, parent=self.root)
        self.bduss_entry.focus_set()


    def _clear_bduss(self) -> None:
        self.bduss_var.set("")
        if self.login_method_var.get() == "bduss":
            self.cookie_var.set("")
        clear_cookie()
        self.status_var.set("已清除 BDUSS")
        self._append_log("已清除 BDUSS。", "muted")


    def _show_cookie_help(self) -> None:
        messagebox.showinfo("如何获取完整 Cookie", COOKIE_HELP, parent=self.root)
        self.cookie_entry.focus_set()


    def _login_result_text(self, action: str) -> str:
        if getattr(self, "last_threads", None):
            return f"{action}。重新开始统计时，已完成的帖子不会重抓。"
        return f"{action}。请回到抓取页，粘贴帖子链接后点「开始统计」。"


    def _ask_go_fill_bduss(self) -> bool:
        win = tk.Toplevel(self.root)
        win.title("需要 BDUSS")
        win.configure(bg="#eef1f5")
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        result = {"go": False}
        tk.Label(
            win,
            text="关注检测已开启，但还没有完成登录。现在去抓取页填写吗？",
            bg="#eef1f5",
            fg="#334155",
            font=("Microsoft YaHei UI", 9),
            justify="left",
            wraplength=360,
        ).pack(fill="x", padx=14, pady=(14, 8))
        row = tk.Frame(win, bg="#eef1f5")
        row.pack(side="bottom", fill="x", padx=14, pady=(0, 12))

        def choose(go: bool) -> None:
            result["go"] = go
            win.destroy()

        ttk.Button(row, text="稍后", command=lambda: choose(False), width=9).pack(side="right")
        ttk.Button(row, text="去填写", command=lambda: choose(True), width=9).pack(
            side="right", padx=(0, 6)
        )
        win.protocol("WM_DELETE_WINDOW", lambda: choose(False))
        win.bind("<Escape>", lambda _event: choose(False))
        self._theme_widget_tree(win)
        win.update_idletasks()
        width, height = 420, max(140, win.winfo_reqheight() + 16)
        x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - width) // 2)
        y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - height) // 2)
        win.geometry(f"{width}x{height}+{x}+{y}")
        win.focus_set()
        self.root.wait_window(win)
        return bool(result["go"])


    def _go_to_login(self) -> None:
        self.notebook.select(self.tab_posts)
        self.post_nb.select(self.tab_fetch)
        self.root.update_idletasks()
        target = self.bduss_entry if self.login_method_var.get() == "bduss" else self.cookie_entry
        target.focus_set()


    def _toggle_login_fields(self) -> None:
        self._set_login_fields_expanded(not getattr(self, "_login_fields_expanded", True))


    def _set_login_fields_expanded(self, expanded: bool) -> None:
        self._login_fields_expanded = bool(expanded)
        if not hasattr(self, "login_method_row"):
            return
        if self._login_fields_expanded:
            self.login_method_row.grid()
            self.login_switch_row.grid()
            if self.login_method_var.get() == "bduss":
                self.bduss_frame.grid()
                self.cookie_frame.grid_remove()
            else:
                self.cookie_frame.grid()
                self.bduss_frame.grid_remove()
            self.login_toggle_btn.configure(text="收起")
        else:
            self.login_method_row.grid_remove()
            self.bduss_frame.grid_remove()
            self.cookie_frame.grid_remove()
            self.login_switch_row.grid_remove()
            self.login_toggle_btn.configure(text="展开")


    def _refresh_webview_status(self) -> None:
        if not hasattr(self, "webview_status_var"):
            return
        is_login = self.mode_var.get() == "login"
        mode = "登录" if is_login else "游客"
        available = webview2.available()
        if not available:
            login_state = "WebView2 不可用"
        elif webview2.login_profile_exists():
            login_state = "已保存本机登录态"
        else:
            login_state = "未保存本机登录态"
        if not is_login:
            login_info = "游客匿名 Cookie"
            has_login = False
        elif self.login_method_var.get() == "bduss":
            login_info = "仅 BDUSS" if self.bduss_var.get().strip() else "未填写"
            has_login = bool(self.bduss_var.get().strip())
        else:
            has_login = webview2.has_login_cookie(
                parse_cookie_input(self.cookie_var.get())
            )
            login_info = "完整 Cookie" if has_login else "未填写"
        self.webview_status_var.set(
            f"当前模式：{mode}\n当前登录信息：{login_info}\n内置登录：{login_state}"
        )
        if hasattr(self, "login_baidu_btn"):
            use_saved = (
                is_login
                and not has_login
                and webview2.login_profile_exists()
                and bool(self.settings.get("webview_remember_login", False))
            )
            if use_saved:
                self.login_baidu_btn.configure(
                    text="使用已保存的内置登录",
                    command=self._reuse_internal_login,
                )
            else:
                self.login_baidu_btn.configure(
                    text=(
                        "登录百度并获取 BDUSS"
                        if self.login_method_var.get() == "bduss"
                        else "登录百度并获取完整 Cookie"
                    ),
                    command=self._login_baidu,
                )
        if hasattr(self, "login_summary_var"):
            if not is_login:
                summary = "游客模式 · 匿名 Cookie"
            elif login_info == "未填写":
                summary = "登录信息未填写"
            else:
                summary = f"{login_info} · 已就绪"
            self.login_summary_var.set(summary)
        if hasattr(self, "login_source_label"):
            source = str(getattr(self, "_login_source", "") or "")
            match: Optional[bool] = None
            if has_login and webview2.login_profile_exists() and source in (
                "手动粘贴",
                "已保存 Cookie",
                "内置浏览器",
            ):
                if source == "内置浏览器" and self._last_webview_match is not None:
                    match = self._last_webview_match
                else:
                    current_cookies = parse_cookie_input(self.cookie_var.get())
                    if source == "内置浏览器" and self._full_cookie_cache:
                        current_cookies = parse_cookie_input(self._full_cookie_cache)
                    match = webview2.match_profile_fingerprint(current_cookies)
            self.login_source_label.unbind("<Button-1>")
            if match is False:
                self.login_source_label.configure(
                    text=(
                        f"当前登录来源：{source} · "
                        "内置登录：可能不是同一会话（点此沿用）"
                    ),
                    fg="#b45309",
                    cursor="hand2",
                )
                self.login_source_label.bind(
                    "<Button-1>", lambda _event: self._reuse_internal_login()
                )
            elif source:
                self.login_source_label.configure(
                    text=f"当前登录来源：{source}",
                    fg="#8a94a6",
                    cursor="",
                )
            else:
                self.login_source_label.configure(text="")
        can_open = available and is_login and not self._network_task_active()
        if hasattr(self, "webview_open_btn"):
            self.webview_open_btn.configure(state="normal" if can_open else "disabled")
        if hasattr(self, "webview_clear_btn"):
            can_clear = available and not self._network_task_active()
            self.webview_clear_btn.configure(state="normal" if can_clear else "disabled")
        if hasattr(self, "follow_enhance_check"):
            self.follow_enhance_check.configure(
                state="normal" if has_login and not self._network_task_active() else "disabled"
            )
        if hasattr(self, "follow_load_box"):
            self.follow_load_box.configure(
                state="readonly" if has_login and self.follow_enhance_var.get() else "disabled"
            )
        if hasattr(self, "follow_enhance_status"):
            self.follow_enhance_status.configure(
                text=(
                    "已具备登录信息"
                    if has_login
                    else "需要登录信息：先点「登录百度」或「手动填写」"
                )
            )
        if not has_login and getattr(self, "follow_enhance_var", None) is not None:
            if self.follow_enhance_var.get():
                self.follow_enhance_var.set(False)
                self.settings["follow_enhance"] = False
                save_settings(self.settings)


    def _on_webview_remember_changed(self) -> None:
        enabled = bool(self.webview_remember_var.get())
        if not enabled and webview2.login_profile_exists():
            if messagebox.askyesno(
                "关闭记住内置登录",
                "关闭后不会自动复用已保存的登录态。\n是否现在一并清除 WebView2 登录？",
                parent=self.root,
                default=messagebox.NO,
            ):
                if not webview2.clear_login_profile():
                    messagebox.showwarning(
                        "暂未清除",
                        "WebView2 登录目录可能正在使用，请稍后手动清除。",
                        parent=self.root,
                    )
                    enabled = True
                    self.webview_remember_var.set(True)
            else:
                enabled = True
                self.webview_remember_var.set(True)
        self.settings["webview_remember_login"] = enabled
        save_settings(self.settings)
        self._refresh_webview_status()


    def _go_manage_login(self) -> None:
        if self.mode_var.get() != "login":
            self._set_mode("login")
        self.notebook.select(self.tab_posts)
        self.post_nb.select(self.tab_fetch)
        self._set_login_fields_expanded(True)
        self.root.update_idletasks()
        target = self.bduss_entry if self.login_method_var.get() == "bduss" else self.cookie_entry
        target.focus_set()


    def _login_baidu(self) -> None:
        if not webview2.available():
            messagebox.showwarning(
                "WebView2 不可用",
                "未检测到 WebView2 Runtime，请使用「手动填写」。",
                parent=self.root,
            )
            return
        if self.mode_var.get() != "login":
            self._set_mode("login")
        self._set_login_fields_expanded(True)
        target = "bduss" if self.login_method_var.get() == "bduss" else "cookie"
        self._auto_get_login(
            target,
            url="https://passport.baidu.com/v2/?login",
        )


    def _save_login_fingerprint(self, cookies: dict, force: bool = False) -> None:
        if not webview2.login_profile_exists():
            return
        if not force and not bool(self.settings.get("webview_remember_login", False)):
            return
        if not webview2.save_profile_fingerprint(cookies):
            self._append_log("登录来源指纹保存失败，下次可能无法判断是否同一会话。", "warn")


    def _adopt_webview_login(self, cookies: dict, target: str) -> None:
        full = "; ".join(f"{key}={value}" for key, value in cookies.items())
        if target == "bduss" and cookies.get("BDUSS"):
            self._full_cookie_cache = ""
            self.login_method_var.set("bduss")
            self.bduss_var.set(str(cookies.get("BDUSS") or ""))
            self.cookie_var.set(f"BDUSS={cookies.get('BDUSS') or ''}")
        else:
            self._full_cookie_cache = full
            self.login_method_var.set("cookie")
            self.cookie_var.set(full)
            self._webview_cookie_session = True
        self._login_source = "内置浏览器"
        self._login_probe_token += 1
        self._on_login_method_changed()
        self._set_login_fields_expanded(False)
        if hasattr(self, "_refresh_webview_status"):
            self._refresh_webview_status()


    def _switch_baidu_account(self) -> None:
        if not webview2.available():
            messagebox.showwarning(
                "WebView2 不可用",
                "未检测到 WebView2 Runtime，无法切换百度账号。",
                parent=self.root,
            )
            return
        if self._network_task_active() or getattr(self, "_webview_running", False):
            messagebox.showinfo("请稍后", "当前有任务或内置窗口正在运行。", parent=self.root)
            return
        if not messagebox.askyesno(
            "切换百度账号",
            "将打开一个临时干净的内置窗口登录新账号。\n"
            "新账号成功登录前，旧登录和旧内置会话都会保留。\n\n继续？",
            parent=self.root,
            default=messagebox.NO,
        ):
            return
        temp_profile = webview2.prepare_switch_profile()
        self._webview_running = True
        self.status_var.set("正在打开切换账号窗口…")
        center = self._webview_center()

        def worker():
            try:
                result = webview2.open_verification(
                    "https://passport.baidu.com/v2/?login",
                    center=center,
                    persistent=True,
                    require_login=True,
                    profile_dir=temp_profile,
                )
                raw = result.get("cookies") if isinstance(result, dict) else result
                cookies = webview2.filter_cookies(raw, anonymous=False)
                if not webview2.has_login_cookie(cookies):
                    raise RuntimeError("没有读取到登录 Cookie（BDUSS/TIEBA_SID），可能尚未登录")
                self._post(("switch_login_done", cookies, "", ""))
            except Exception as error:  # noqa: BLE001
                self._post(("switch_login_done", {}, str(error), ""))

        threading.Thread(target=worker, daemon=True).start()


    def _reuse_internal_login(self) -> None:
        if self.mode_var.get() != "login":
            self._set_mode("login")
        target = "bduss" if self.login_method_var.get() == "bduss" else "cookie"
        url = (getattr(self, "_current_urls", []) or ["https://tieba.baidu.com/"])[0]
        self._auto_get_login(target, url=url)


    def _manual_fill_login(self) -> None:
        self._set_mode("login")
        self.login_method_var.set("cookie")
        self._set_login_fields_expanded(True)
        self._on_login_method_changed()
        try:
            self.notebook.select(self.tab_posts)
            self.post_nb.select(self.tab_fetch if hasattr(self, "tab_fetch") else 0)
        except tk.TclError:
            pass
        self.cookie_entry.focus_set()
        self._append_log("请手动粘贴完整 Cookie，然后点「应用」。", "muted")


    def _open_internal_login(self) -> None:
        if self.mode_var.get() != "login":
            messagebox.showinfo(
                "仅登录模式可用",
                "请先在抓取页切换到登录模式，或从游客验证码选择改用登录。",
                parent=self.root,
            )
            return
        if not webview2.available():
            messagebox.showwarning(
                "WebView2 不可用",
                "未检测到 WebView2 Runtime，请改用手动完整 Cookie。",
                parent=self.root,
            )
            return
        entry = self._ask_choice(
            "打开内置登录",
            "只开放以下百度入口，不提供任意网址。",
            [
                ("百度登录", "login"),
                ("当前帖子验证页", "thread"),
                ("贴吧首页", "home"),
                ("取消", "cancel"),
            ],
            default=0,
        )
        if not entry or entry == "cancel":
            return
        if entry == "login":
            target_url = "https://passport.baidu.com/v2/?login"
        elif entry == "home":
            target_url = "https://tieba.baidu.com/"
        else:
            target_url = (getattr(self, "_current_urls", []) or ["https://tieba.baidu.com/"])[0]
        self._start_webview_cookie_flow(anonymous=False, url=target_url)


    def _auto_get_login(self, target: str, url: str = "") -> None:
        if self.mode_var.get() != "login":
            messagebox.showinfo("仅登录模式可用", "请先切换到登录模式。", parent=self.root)
            return
        if not webview2.available():
            messagebox.showwarning(
                "WebView2 不可用",
                "未检测到 WebView2 Runtime，请手动粘贴 BDUSS 或完整 Cookie。",
                parent=self.root,
            )
            return
        if self._network_task_active() or getattr(self, "_webview_running", False):
            messagebox.showinfo("请稍后", "当前有任务或内置窗口正在运行。", parent=self.root)
            return
        if not messagebox.askyesno(
            "自动获取登录信息",
            "将打开内置百度登录窗口。\n"
            "登录完成后读取完整 Cookie"
            + ("，只提取 BDUSS。" if target == "bduss" else "。")
            + "\nCookie 默认只在内存；开启“记住内置登录”时 WebView2 会话会保留在本机。\n\n继续？",
            parent=self.root,
            default=messagebox.NO,
        ):
            return
        target_url = url or (getattr(self, "_current_urls", []) or ["https://tieba.baidu.com/"])[0]
        if not webview2.is_allowed_url(target_url):
            target_url = "https://tieba.baidu.com/"
        self._webview_running = True
        self.status_var.set("正在打开内置登录…")
        url = target_url
        center = self._webview_center()
        persistent = bool(self.settings.get("webview_remember_login", False))
        target_value = target

        def worker():
            try:
                result = webview2.open_verification(
                    url,
                    center=center,
                    persistent=persistent,
                    require_login=True,
                )
                raw = result.get("cookies") if isinstance(result, dict) else result
                cookies = webview2.filter_cookies(raw, anonymous=False)
                if not webview2.has_login_cookie(cookies):
                    raise RuntimeError("没有读取到登录 Cookie（BDUSS/TIEBA_SID），可能尚未登录")
                effective_target = (
                    "cookie"
                    if target_value == "bduss" and not cookies.get("BDUSS")
                    else target_value
                )
                warning = ""
                probe = Fetcher(cookie=cookies, delay=0, verbose=False)
                try:
                    probe.get("https://tieba.baidu.com/")
                except CaptchaError:
                    warning = "预检仍被验证码拦截，Cookie 可能不可用"
                except Exception as error:  # noqa: BLE001
                    warning = f"预检失败：{error}"
                finally:
                    probe.close()
                self._post(
                    ("auto_login_done", effective_target, cookies, "", warning)
                )
            except Exception as error:  # noqa: BLE001
                self._post(("auto_login_done", target_value, {}, str(error), ""))

        threading.Thread(target=worker, daemon=True).start()


    def _clear_internal_login(self) -> None:
        if self._network_task_active() or getattr(self, "_webview_running", False):
            messagebox.showinfo("请稍后", "抓取、补全或内置窗口进行中，暂时不能清除内置登录。", parent=self.root)
            return
        choice = self._ask_choice(
            "清除 WebView2 保存的登录",
            "只清除内置 WebView2 的百度登录态。\n"
            "当前 Cookie 输入框和缓存是否一并清除，请选择。",
            [
                ("只清除 WebView2 登录态", "webview"),
                ("同时清除当前登录信息", "all"),
                ("取消", "cancel"),
            ],
            default=0,
        )
        if not choice or choice == "cancel":
            return
        if not webview2.clear_login_profile():
            messagebox.showerror(
                "清除失败",
                "WebView2 登录目录可能正在使用。请关闭内置窗口和本软件后重试。",
                parent=self.root,
            )
            return
        if choice == "all":
            self._clear_cookie()
            self._guest_cookie = {}
            self.cookie_var.set("")
            self.bduss_var.set("")
            self._full_cookie_cache = ""
        self._append_log("已清除 WebView2 保存的登录。", "ok")
        self._refresh_webview_status()


    def _webview_center(self) -> tuple[int, int]:
        try:
            self.root.update_idletasks()
            x = self.root.winfo_rootx() + self.root.winfo_width() // 2 - 490
            y = self.root.winfo_rooty() + self.root.winfo_height() // 2 - 380
            return max(0, x), max(0, y)
        except tk.TclError:
            return 0, 0


    def _start_webview_cookie_flow(self, anonymous: bool, url: str = "") -> None:
        scope = (
            "只读取匿名 Cookie，不读取 BDUSS/STOKEN"
            if anonymous
            else "读取完整 Cookie，可能包含 BDUSS/STOKEN"
        )
        storage = (
            "Cookie 只在内存中使用，不写日志、不导出。"
            if anonymous
            else "Cookie 默认只在内存中使用；若开启“记住内置登录”，"
            "WebView2 登录态会写入本机独立目录。"
        )
        if not messagebox.askyesno(
            "读取验证 Cookie？",
            "内置验证完成后，软件将" + scope + "。\n"
            + storage
            + "\n\n继续？",
            parent=self.root,
            default=messagebox.NO,
        ):
            self._append_log("用户取消读取验证 Cookie。", "muted")
            return
        target_url = url or (getattr(self, "_current_urls", []) or ["https://tieba.baidu.com/"])[0]
        if not webview2.is_allowed_url(target_url):
            messagebox.showwarning(
                "地址不允许",
                "内置窗口只允许百度贴吧登录/验证页面。",
                parent=self.root,
            )
            return
        self._webview_running = True
        self.status_var.set("正在打开内置验证窗口…")
        center = self._webview_center()

        def worker():
            try:
                result = webview2.open_verification(
                    target_url,
                    center=center,
                    persistent=(not anonymous)
                    and bool(self.settings.get("webview_remember_login", False)),
                    require_login=not anonymous,
                )
                raw_cookies = result.get("cookies") if isinstance(result, dict) else result
                cookies = webview2.filter_cookies(raw_cookies, anonymous=anonymous)
                if not cookies:
                    raise RuntimeError("没有读取到可用 Cookie")
                warning = ""
                if not anonymous:
                    probe = Fetcher(cookie=cookies, delay=0, verbose=False)
                    try:
                        probe.get("https://tieba.baidu.com/")
                    except CaptchaError:
                        warning = "预检仍被验证码拦截，Cookie 可能不可用"
                    except Exception as error:  # noqa: BLE001
                        warning = f"预检失败：{error}"
                    finally:
                        probe.close()
                self._post(("embedded_done", cookies, "", warning))
            except Exception as error:  # noqa: BLE001
                self._post(("embedded_done", {}, str(error), ""))

        threading.Thread(target=worker, daemon=True).start()


    def _request_cookie(self) -> dict:
        if self.mode_var.get() == "guest":
            return dict(getattr(self, "_guest_cookie", {}) or {})
        return parse_cookie_input(self._current_cookie_text())


class App(GuiLayoutMixin, MessageMixin, SettingsMixin, GridMixin, LookupMixin, WorkerMixin, ExportMixin, BatchIdentityMixin, ThemeMixin, ProfileMixin, ToolsMixin, LoginMixin):
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.log_queue: "queue.Queue[tuple]" = queue.Queue()
        self.stop_event = threading.Event()
        self.worker: threading.Thread | None = None
        self.running = False
        self.last_output: str | None = None
        self.last_outputs: List[str] = []
        self.last_threads: list = []
        self._pending_retry = False
        self._auto_filled = False
        self._identity_choice_queue: "queue.Queue[tuple]" = queue.Queue()
        self._rate_choice_queue: "queue.Queue[str]" = queue.Queue()
        self._last_stage_times: Dict[str, float] = {}
        self._manual_timing: Dict[str, float] = {}
        self._supplement_stop = threading.Event()
        self._guest_cookie: Dict[str, str] = {}
        self._full_cookie_cache = ""
        self._webview_cookie_session = False
        self._tieba_uid_failed: set[str] = set()
        self._user_id_failed: set[str] = set()
        self._profile_topic_limit = 50
        self._lookup_normalize_job: Optional[str] = None
        self._emoji_window: Optional[tk.Toplevel] = None
        self.lookup_entry: Optional[ttk.Entry] = None
        self.lookup_native_edit: Optional[int] = None
        self.lookup_native_host: Optional[tk.Frame] = None
        self.love_native_edit: Optional[int] = None
        self.love_native_host: Optional[tk.Frame] = None
        self.emoji_btn: Optional[tk.Button] = None
        self.blacklist = load_blacklist()
        self.settings = load_settings()
        post_cache().set_limit_mb(int(self.settings.get("post_cache_limit_mb", DEFAULT_LIMIT_MB)))
        self.palette = resolve_palette(
            str(self.settings.get("theme_mode", "light")),
            str(self.settings.get("accent", "blue")),
        )
        self.mode_var = tk.StringVar(value="guest")
        self._login_source = ""
        self._last_webview_match: Optional[bool] = None
        self._login_probe_token = 0
        self._active_choice_window: Optional[tk.Toplevel] = None
        self._usage_window: Optional[tk.Toplevel] = None
        webview2.cleanup_profile_backup()
        webview2.clear_switch_profile()

        self._build_ui()
        self._apply_window_icon()
        self._apply_theme()
        self._apply_table_density()
        self.notebook.bind("<<NotebookTabChanged>>", self._on_main_tab_changed)
        saved_geometry = str(self.settings.get("window_geometry") or "")
        match = re.match(r"(\d+)x(\d+)", saved_geometry)
        if match:
            saved_w, saved_h = int(match.group(1)), int(match.group(2))
            # 旧版本可能记住了窄窗口；小于新默认值时恢复默认尺寸。
            screen_w = self.root.winfo_screenwidth()
            screen_h = self.root.winfo_screenheight()
            width = max(1300, min(saved_w, screen_w - 100)) if saved_w >= 1300 else min(1400, screen_w - 100)
            height = max(820, min(saved_h, screen_h - 120)) if saved_h >= 820 else min(900, screen_h - 120)
            x = max(0, (screen_w - width) // 2)
            y = max(0, (screen_h - height) // 2)
            try:
                self.root.geometry(f"{width}x{height}+{x}+{y}")
            except tk.TclError:
                pass
        self._set_output(default_output_path())
        self.root.after(120, self._poll_queue)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self.root.bind("<F1>", lambda _event: self._open_usage_window())
        self.root.bind("<Control-Shift-C>", self._show_test_center)
        self.root.bind("<Escape>", self._on_escape)

        saved = load_cookie()
        if saved:
            self.cookie_var.set(saved)
            self._sync_login_from_cookie()
            self._login_source = "已保存 Cookie"
            self._append_log("已自动填入上次保存的登录信息（如果失效，请重新粘贴一次）。", "muted")
        if self.settings.get("follow_enhance"):
            self.mode_var.set("login")
            if not parse_cookie_input(self.cookie_var.get()).get("BDUSS"):
                self._append_log("关注检测已开启，但尚未完成登录；请先登录后再开始统计。", "muted")
        else:
            self.mode_var.set("guest")
        self._apply_mode_ui()
        self._refresh_webview_status()
        self._apply_risk_ui()
        self._apply_proxy()
        self._restore_last_result()
        self.root.after(300, self._maybe_onboarding)

    def _apply_window_icon(self) -> None:
        for path in (
            os.path.join(resource_dir(), "assets", "app.ico"),
            os.path.join(resource_dir(), "app.ico"),
        ):
            if not os.path.isfile(path):
                continue
            try:
                self.root.iconbitmap(default=path)
                self.root.iconbitmap(path)
                return
            except tk.TclError:
                continue

    # ------------------------------------------------------------------ UI

    # ------------------------------------------------------------- actions
    def _set_output(self, path: str) -> None:
        self.output_var.set(path)

    def _choose_output(self) -> None:
        path = filedialog.asksaveasfilename(
            title="选择保存位置",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile=os.path.basename(self.output_var.get() or "OK的贴吧玩具箱.xlsx"),
            initialdir=os.path.dirname(self.output_var.get() or ""),
        )
        if path:
            self._set_output(path)

















    def _open_readme(self) -> None:
        for name in ("使用说明.txt", "README.md"):
            path = os.path.join(resource_dir(), name)
            if os.path.isfile(path):
                try:
                    os.startfile(path)  # type: ignore[attr-defined]
                    return
                except OSError:
                    break
        messagebox.showinfo(
            "使用说明",
            "1. 粘贴帖子链接（每行一个）\n"
            "2. 选择保存位置\n"
            "3. 点「开始统计」\n"
            "4. 完成后点「打开 Excel」\n\n"
            "如果提示被安全验证拦截，点「如何获取？」，按说明把 Cookie 粘贴进来。",
            parent=self.root,
        )


    # ------------------------------------------------------- 参与者与查用户
    def _aggregate_users(self, threads=None):
        buckets: Dict[str, Dict[str, Any]] = {}
        for thread in (self.last_threads if threads is None else threads):
            for item in thread.user_stats():
                key = (
                    f"uid:{item['uid']}"
                    if item.get("uid")
                    else (
                        f"portrait:{str(item.get('portrait') or '').split('?')[0]}"
                        if item.get("portrait")
                        else (f"name:{item['username']}" if item.get("username") else f"anon:{item['name']}")
                    )
                )
                bucket = buckets.setdefault(
                    key,
                    {
                        "name": item["name"],
                        "username": item["username"],
                        "uid": item["uid"],
                        "tieba_uid": item.get("tieba_uid") or "",
                        "portrait": item["portrait"],
                        "level": item["level"],
                        "badge": item["badge"],
                        "floor": 0,
                        "comment": 0,
                        "total": 0,
                        "aliases": set(),
                    },
                )
                bucket["floor"] += item["floor_count"]
                bucket["comment"] += item["comment_count"]
                bucket["total"] += item["total"]
                if item.get("name") and item["name"] != bucket["name"]:
                    bucket["aliases"].add(bucket["name"])
                    bucket["name"] = item["name"]
                bucket["aliases"].update(item.get("aliases") or [])
                if not bucket["username"] and item["username"]:
                    bucket["username"] = item["username"]
                if not bucket["portrait"] and item["portrait"]:
                    bucket["portrait"] = item["portrait"]
                if not bucket["tieba_uid"] and item.get("tieba_uid"):
                    bucket["tieba_uid"] = item["tieba_uid"]
                if not bucket["level"] and item["level"]:
                    bucket["level"] = item["level"]
                if not bucket["badge"] and item["badge"]:
                    bucket["badge"] = item["badge"]
        users = list(buckets.values())
        for user in users:
            user["aliases"] = sorted(a for a in user.get("aliases", ()) if a and a != user["name"])
        users.sort(key=lambda item: (-item["total"], -item["floor"], item["name"]))
        return users


    def _on_tab_changed(self, _event=None) -> None:
        try:
            current = self.post_nb.index("current")
            if current == self.post_nb.index(self.tab_people):
                if self.last_threads:
                    self._refresh_people()
                    self._auto_fill_levels()
            elif current == self.post_nb.index(self.tab_grid):
                self._refresh_grid()
        except tk.TclError:
            pass

    def _date_group(self, parent, label: str, months_ago: int = 0):
        import calendar

        base = datetime.now()
        if months_ago:
            year = base.year
            month = base.month - months_ago
            while month <= 0:
                month += 12
                year -= 1
            day = min(base.day, calendar.monthrange(year, month)[1])
            base = base.replace(year=year, month=month, day=day)
        tk.Label(parent, text=f"{label}", bg="#ffffff", fg="#475569", font=("Microsoft YaHei UI", 9)).pack(
            side="left"
        )
        year_var = tk.StringVar(value=str(base.year))
        month_var = tk.StringVar(value=f"{base.month:02d}")
        day_var = tk.StringVar(value=f"{base.day:02d}")
        ttk.Combobox(
            parent,
            values=[str(year) for year in range(2003, datetime.now().year + 2)],
            textvariable=year_var,
            width=5,
            state="readonly",
        ).pack(side="left")
        tk.Label(parent, text="年", bg="#ffffff").pack(side="left")
        ttk.Combobox(
            parent,
            values=[f"{month:02d}" for month in range(1, 13)],
            textvariable=month_var,
            width=3,
            state="readonly",
        ).pack(side="left")
        tk.Label(parent, text="月", bg="#ffffff").pack(side="left")
        ttk.Combobox(
            parent,
            values=[f"{day:02d}" for day in range(1, 32)],
            textvariable=day_var,
            width=3,
            state="readonly",
        ).pack(side="left")
        tk.Label(parent, text="日", bg="#ffffff").pack(side="left")
        return year_var, month_var, day_var

    def _group_date(self, parts) -> str:
        import calendar

        try:
            year, month = int(parts[0].get()), int(parts[1].get())
            day = min(int(parts[2].get()), calendar.monthrange(year, month)[1])
            return f"{year:04d}-{month:02d}-{day:02d}"
        except (ValueError, IndexError):
            return ""

    def _toggle_date_entries(self) -> None:
        enabled = getattr(self, "post_date_mode_var", None) and self.post_date_mode_var.get() == "指定日期"
        try:
            if enabled:
                if not self.post_date_row.winfo_ismapped():
                    self.post_date_row.pack(fill="x", before=self.people_paned)
            else:
                self.post_date_row.pack_forget()
        except tk.TclError:
            pass

    def _on_post_parts_changed(self, *_args) -> None:
        if getattr(self, "_batch_running", False):
            return
        self.post_start_var.set(self._group_date(self.post_start_parts))
        self.post_end_var.set(self._group_date(self.post_end_parts))
        if self.post_date_mode_var.get() == "指定日期":
            self._on_post_date_changed()

    def _on_post_date_changed(self) -> None:
        if getattr(self, "_batch_running", False):
            return
        self._toggle_date_entries()
        if self.post_date_mode_var.get() == "指定日期":
            self.post_start_var.set(self._group_date(self.post_start_parts))
            self.post_end_var.set(self._group_date(self.post_end_parts))
        else:
            self.post_start_var.set("")
            self.post_end_var.set("")
            self.post_prefer_var.set("优先最近")
        if hasattr(self, "_profile_cache"):
            self._profile_cache.clear()
        self.profile_status.configure(text="条件已变更，点「刷新资料」生效")

    def _open_emoji_panel(self) -> None:
        if os.name == "nt":
            try:
                self._toggle_emoji_inline()
                return
            except Exception:  # noqa: BLE001
                pass
        self._open_emoji_picker()

    def _build_emoji_inline(self) -> tk.Frame:
        frame = tk.Frame(self.tab_lookup, bg="#eef1f5", padx=10, pady=6)
        self._emoji_inline_frame = frame
        tk.Label(
            frame,
            text="输入 Emoji（支持微软输入法 / Win+. / Ctrl+V）",
            bg="#eef1f5",
            fg="#475569",
            font=("Microsoft YaHei UI", 8),
            anchor="w",
        ).pack(fill="x")
        row = tk.Frame(frame, bg="#eef1f5")
        row.pack(fill="x", pady=(2, 4))
        host = tk.Frame(row, bg="#ffffff", height=34)
        host.pack(side="left", fill="x", expand=True)
        host.pack_propagate(False)
        self._emoji_host = host
        ttk.Button(row, text="取消", command=self._hide_emoji_inline, width=8).pack(
            side="right", padx=(6, 0)
        )
        ttk.Button(row, text="插入", command=self._commit_emoji_inline, width=8).pack(
            side="right", padx=(6, 0)
        )
        self._theme_widget_tree(frame)
        return frame

    def _toggle_emoji_inline(self) -> None:
        frame = getattr(self, "_emoji_inline_frame", None)
        if frame is not None and frame.winfo_manager():
            self._hide_emoji_inline()
            return
        if frame is None:
            frame = self._build_emoji_inline()
        frame.pack(fill="x", after=self.lookup_query_bar)
        self.root.update_idletasks()
        if not self._create_inline_edit():
            self._hide_emoji_inline()
            self._open_emoji_picker()
            return
        import ctypes

        ctypes.windll.user32.SetFocus(self._emoji_edit)

    def _create_inline_edit(self) -> bool:
        if getattr(self, "_emoji_edit", None):
            return True
        import ctypes
        from ctypes import wintypes

        user32 = ctypes.WinDLL("user32", use_last_error=True)
        gdi32 = ctypes.WinDLL("gdi32", use_last_error=True)
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        user32.CreateWindowExW.restype = wintypes.HWND
        user32.CreateWindowExW.argtypes = [
            wintypes.DWORD, wintypes.LPCWSTR, wintypes.LPCWSTR, wintypes.DWORD,
            ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int,
            wintypes.HWND, wintypes.HMENU, wintypes.HINSTANCE, wintypes.LPVOID,
        ]
        user32.SendMessageW.restype = ctypes.c_ssize_t
        user32.SendMessageW.argtypes = [wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]
        user32.MoveWindow.argtypes = [wintypes.HWND, ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int, wintypes.BOOL]
        gdi32.CreateFontW.restype = wintypes.HFONT
        gdi32.CreateFontW.argtypes = [
            ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int,
            wintypes.DWORD, wintypes.DWORD, wintypes.DWORD, wintypes.DWORD,
            wintypes.DWORD, wintypes.DWORD, wintypes.DWORD, wintypes.DWORD, wintypes.LPCWSTR,
        ]
        kernel32.GetModuleHandleW.restype = wintypes.HMODULE
        hinst = kernel32.GetModuleHandleW(None)
        style = 0x40000000 | 0x10000000 | 0x00800000 | 0x00000080
        edit = user32.CreateWindowExW(
            0, "EDIT", "", style, 0, 0, 100, 32,
            wintypes.HWND(self._emoji_host.winfo_id()), None, hinst, None,
        )
        if not edit:
            return False
        self._emoji_edit = edit
        font = gdi32.CreateFontW(
            -16, 0, 0, 0, 400, 0, 0, 0, 1, 0, 0, 0, 0, "Segoe UI Emoji"
        )
        self._emoji_edit_font = font
        if font:
            user32.SendMessageW(edit, 0x0030, font, 1)

        def resize(event) -> None:
            user32.MoveWindow(edit, 0, 0, max(10, event.width), max(20, event.height), True)

        self._emoji_host.bind("<Configure>", resize)
        resize(type("E", (), {"width": 600, "height": 32})())
        return True

    def _commit_emoji_inline(self) -> None:
        edit = getattr(self, "_emoji_edit", None)
        if not edit:
            return
        import ctypes

        user32 = ctypes.windll.user32
        length = user32.GetWindowTextLengthW(edit)
        buffer = ctypes.create_unicode_buffer(length + 1)
        user32.GetWindowTextW(edit, buffer, length + 1)
        text = buffer.value.strip()
        if text and self.lookup_entry is not None:
            self.lookup_entry.insert("insert", text)
        self._hide_emoji_inline()

    def _hide_emoji_inline(self) -> None:
        frame = getattr(self, "_emoji_inline_frame", None)
        if frame is not None:
            frame.pack_forget()
        edit = getattr(self, "_emoji_edit", None)
        if edit:
            try:
                import ctypes

                ctypes.windll.user32.SetWindowTextW(edit, "")
            except Exception:  # noqa: BLE001
                pass
        if self.lookup_entry is not None:
            self.lookup_entry.focus_set()

    def _open_emoji_picker(self) -> None:
        existing = getattr(self, "_emoji_window", None)
        try:
            if existing is not None and existing.winfo_exists():
                existing.lift()
                return
        except tk.TclError:
            pass
        win = tk.Toplevel(self.root)
        self._emoji_window = win
        win.title("选择 Emoji")
        win.configure(bg="#eef1f5")
        win.resizable(False, False)
        win.transient(self.root)
        frame = tk.Frame(win, bg="#eef1f5", padx=8, pady=8)
        frame.pack()
        for index, char in enumerate(EMOJI_PICKER):
            def choose(value: str = char) -> None:
                self._insert_emoji(value, win)

            button = tk.Button(
                frame,
                text=char,
                font=("Segoe UI Emoji", 14),
                width=2,
                relief="flat",
                bg="#ffffff",
                activebackground="#dbeafe",
                cursor="hand2",
                command=choose,
            )
            button.grid(row=index // 8, column=index % 8, padx=2, pady=2)
        self._theme_widget_tree(win)
        self._center_toplevel(win)
        win.bind("<Escape>", lambda _event: win.destroy())
        win.focus_set()

    def _insert_emoji(self, char: str, win=None) -> None:
        if self.lookup_entry is not None:
            self.lookup_entry.insert("insert", char)
            self.lookup_entry.focus_set()
        if win is not None:
            win.destroy()




    def _make_sortable(self, tree, headers, sortable=None) -> None:
        """点表头在升序/降序之间切换；sortable 限定真正可排序的列。"""
        state: Dict[str, Any] = {"column": None, "reverse": False}
        allowed = set(sortable) if sortable is not None else set(headers)

        def sort_by(column: str) -> None:
            if state["column"] == column:
                state["reverse"] = not state["reverse"]
            else:
                state["column"] = column
                state["reverse"] = False
            reverse = bool(state["reverse"])

            def key(item):
                value = tree.set(item, column)
                text = value[3:] if value.startswith("Lv.") else value
                try:
                    return (0, float(text))
                except ValueError:
                    return (1, str(value))

            items = list(tree.get_children())
            items.sort(key=key, reverse=reverse)
            for index, item in enumerate(items):
                tree.move(item, "", index)
            for col, title in headers.items():
                arrow = ""
                if col == column:
                    arrow = " ▼" if reverse else " ▲"
                tree.heading(col, text=title + arrow)

        for col, title in headers.items():
            if col in allowed:
                tree.heading(col, text=title, command=lambda c=col: sort_by(c))
            else:
                tree.heading(col, text=title)

    def _selected_person(self):
        selection = self.people_tree.selection()
        if not selection:
            return None
        try:
            return self.people_users[int(selection[0])]
        except (IndexError, ValueError):
            return None

    def _copy_text(self, content: str) -> None:
        self.root.clipboard_clear()
        self.root.clipboard_append(content or "")
        self.people_hint.configure(text="已复制到剪贴板")

    def _copy_tree_page(self, tree) -> None:
        lines = []
        for iid in tree.get_children():
            lines.append("	".join(str(value) for value in tree.item(iid, "values")))
        self._copy_text(chr(10).join(lines))

    def _copy_people_page(self) -> None:
        self._copy_tree_page(self.people_tree)

    def _copy_person_row(self) -> None:
        user = self._selected_person()
        if not user:
            return
        level = f"Lv.{user['level']}" if user.get("level") else "-"
        self._copy_text(
            "\t".join(
                str(item)
                for item in (
                    user.get("name", ""),
                    user.get("username", ""),
                    user.get("tieba_uid", ""),
                    user.get("uid", ""),
                    level,
                )
            )
        )

    def _copy_person_field(self, key: str) -> None:
        user = self._selected_person()
        if user:
            self._copy_text(str(user.get(key) or ""))

    def _person_menu(self, event) -> None:
        if getattr(self, "_batch_running", False):
            return
        if self.people_tree.identify_region(event.x, event.y) == "heading":
            column = self.people_tree.identify_column(event.x)
            try:
                index = int(column[1:]) - 1
                key = self.people_columns[index]
            except (ValueError, IndexError):
                return
            if key in ("mark", "name", "username", "tieba_uid", "uid"):
                self._filter_people_column(key, (event.x_root, event.y_root))
            return
        row = self.people_tree.identify_row(event.y)
        if row:
            self.people_tree.selection_set(row)
            try:
                self._person_popup.tk_popup(event.x_root, event.y_root)
            finally:
                self._person_popup.grab_release()


    def _offer_retry_incomplete(self, source: str) -> None:
        targets = [
            thread
            for thread in getattr(self, "last_threads", [])
            if not getattr(thread, "complete", True)
        ]
        if not targets:
            return
        if messagebox.askyesno(
            f"{source}已更新",
            f"有 {len(targets)} 个帖子未完成。\n"
            "是否只重试失败部分（不会重抓已完成页/楼层）？",
            parent=self.root,
            default=messagebox.YES,
        ):
            # 登录更新回调可能仍处于原任务收尾阶段；延后到主循环再启动。
            self._pending_retry = True
            self.root.after(200, self._retry_incomplete_threads)

    def _retry_incomplete_threads(self) -> None:
        if getattr(self, "_supplement_running", False) or self.running:
            self._pending_retry = True
            self._append_log("当前任务尚未结束，稍后重试失败部分。", "muted")
            return
        self._pending_retry = False
        targets = [thread for thread in self.last_threads if not getattr(thread, "complete", True)]
        if not targets:
            messagebox.showinfo("没有失败部分", "当前帖子都是完整的。", parent=self.root)
            return
        mode = self.mode_var.get()
        if mode == "login":
            cookie = self._request_cookie()
            if not webview2.has_login_cookie(cookie):
                messagebox.showwarning(
                    "需要登录 Cookie",
                    "登录模式续抓需要 BDUSS/TIEBA_SID，请先重新获取登录信息。",
                    parent=self.root,
                )
                return
        else:
            cookie = {}
        resume_requests()
        configure_gate(float(self.delay_var.get() or 0.8))
        self._set_supplement_running(True)
        self._supplement_stop.clear()
        set_stop_checker(self._supplement_stop.is_set)
        self.retry_incomplete_btn.state(["disabled"])
        self.people_hint.configure(text=f"正在重试 {len(targets)} 个帖子的失败部分…")

        def worker():
            try:
                fetcher = Fetcher(cookie=cookie, delay=0, verbose=False)
                options = CollectOptions(
                    mode=mode,
                    include_comments=self.comments_var.get(),
                    comment_depth=self._comment_depth_value(),
                    comment_concurrency=max(
                        1, min(3, int(self.comment_concurrency_var.get() or 1))
                    ),
                    verbose=True,
                    log=self._log_from_thread,
                    should_stop=self._supplement_stop.is_set,
                )
                for thread in targets:
                    url = thread.meta.url or f"https://tieba.baidu.com/p/{thread.meta.tid}"
                    collect_thread(url, fetcher, options, resume_thread=thread)
                self._persist_result(self.last_threads)
                self._post(("incomplete_retry_done", len(targets)))
            except CollectCancelled:
                self._persist_result(self.last_threads)
                self._post(("incomplete_retry_stopped",))
            except Exception as error:  # noqa: BLE001
                self._post(("incomplete_retry_error", str(error)))

        threading.Thread(target=worker, daemon=True).start()





























    def _start_probe(self, user, profile, button) -> None:
        uid = str(profile.get("user_id") or user.get("uid") or "")
        portrait = str(profile.get("portrait") or user.get("portrait") or "")
        if not uid or not portrait:
            messagebox.showinfo("缺少资料", "没有用户 ID 或头像，无法探测。", parent=self.root)
            return
        candidates = candidate_bars(user, profile, self.blacklist)
        if not candidates:
            messagebox.showinfo(
                "没有候选吧",
                "本帖发言、发言、本地标记里没有可探测的吧。",
                parent=self.root,
            )
            return
        self._probe_running = True
        self._probe_stop = threading.Event()
        button.configure(state="normal", text="停止探测")
        self._append_log(f"开始探测 {len(candidates)} 个候选吧（最多 30 个）…", "muted")

        def worker() -> None:
            fetcher = Fetcher(cookie={}, delay=0.3, verbose=False)
            results = probe_forums(
                fetcher,
                uid,
                portrait,
                candidates,
                log=self._log_from_thread,
                should_stop=self._probe_stop.is_set,
            )
            self._post(("probe_done", profile, results, button))

        threading.Thread(target=worker, daemon=True).start()

    def _post_source_note(self, profile) -> str:
        src = profile.get("post_sources") or {}
        if src.get("replies_ok") is False:
            return "回复流查询失败，当前只显示主题帖"
        if src.get("replies_empty") and src.get("threads_ok"):
            return "回复流为空，当前只显示主题帖"
        return ""

    def _cache_age_note(self, profile) -> str:
        src = profile.get("post_sources") or {}
        if src.get("cached"):
            age = src.get("cached_age_minutes")
            if age is not None:
                return f"缓存于 {age} 分钟前（刷新资料可强制更新）"
        age_seconds = profile.get("_ui_cache_age_seconds")
        if age_seconds is not None:
            return f"缓存于 {max(0, int(age_seconds) // 60)} 分钟前（刷新资料可强制更新）"
        return ""





    def _apply_level_results(self, kind, profile, forums, results) -> None:
        by_name = {str(forum.get("name") or ""): forum for forum in forums}
        for name, data in results.items():
            forum = by_name.get(name)
            if forum is None:
                continue
            forum["level"] = int(data.get("level") or 0)
            forum["level_source"] = "恋文（第三方）"
            forum["third_party_level"] = True
            for key in ("followed_forums", "hidden_forums", "like_forum", "probed_forums"):
                for saved in profile.get(key) or []:
                    if str(saved.get("name") or "") == name:
                        saved["level"] = forum["level"]
                        saved["level_source"] = "恋文（第三方）"
        tree = self.lookup_forum_tree if kind == "lookup" else self.profile_forum_tree
        names = self._lookup_forum_names if kind == "lookup" else self._profile_forum_names
        for iid, name in list(names.items()):
            forum = by_name.get(name)
            if not forum:
                continue
            try:
                tree.set(iid, "level", f"Lv.{forum['level']}" if forum.get("level") else "-")
                tree.set(iid, "level_source", forum.get("level_source") or "-")
            except tk.TclError:
                pass



    def _open_participants(self) -> None:
        if not self.last_threads:
            messagebox.showinfo("还没有数据", "先点「开始统计」抓一个帖子，再查看参与者。", parent=self.root)
            return
        self._refresh_people()
        self.notebook.select(self.tab_posts)
        self.post_nb.select(self.tab_people)

    def _update_people_sort_headings(self) -> None:
        if not hasattr(self, "people_tree"):
            return
        current, reverse = getattr(self, "_people_sort", ("", False))
        for column, label in getattr(self, "_people_headings", {}).items():
            if column not in ("mark", "level", "floor", "comment", "total"):
                continue
            arrow = ""
            if column == current:
                arrow = "  ▲" if not reverse else "  ▼"
            try:
                self.people_tree.heading(column, text=label + arrow)
            except tk.TclError:
                pass

    def _sort_people(self, column: str) -> None:
        if getattr(self, "_batch_running", False):
            return
        if column not in ("mark", "level", "floor", "comment", "total"):
            return
        current, reverse = self._people_sort
        self._people_sort = (column, (not reverse) if current == column else False)
        self._people_page = 0
        self._update_people_sort_headings()
        self._refresh_people()

    def _people_column_value(self, user, column: str) -> str:
        if column == "mark":
            level = str(user.get("mark_level") or "")
            if not level:
                level, _bar = classify(user.get("bars") or [], self.blacklist)
            return LEVEL_LABEL.get(level, "") if self._risk_enabled() else ""
        if column == "name":
            return self._display_name(user)
        if column == "username":
            return user.get("username") or "-"
        if column == "tieba_uid":
            return self._tieba_uid_cell(user)
        if column == "uid":
            return user.get("uid") or "-"
        return ""

    def _filter_people_column(self, column: str, anchor: Optional[tuple] = None) -> None:
        if getattr(self, "_batch_running", False):
            return
        values = sorted({self._people_column_value(user, column) for user in self.people_users})
        chosen = self._filter_text_dialog(
            {"mark": "标记", "name": "昵称", "username": "用户名", "tieba_uid": "贴吧ID", "uid": "用户ID"}.get(column, column),
            values,
            self._people_col_filters.get(column),
            anchor,
        )
        if chosen is None:
            return
        if chosen:
            self._people_col_filters[column] = chosen
        else:
            self._people_col_filters.pop(column, None)
        self._people_page = 0
        self._refresh_people()

    def _on_people_filter_changed(self) -> None:
        if getattr(self, "_batch_running", False):
            return
        self._people_page = 0
        if getattr(self, "_filter_resetting", False):
            return
        if self.people_level_var.get() == "自定义":
            self.people_custom_frame.pack(side="left", padx=(4, 0))
        else:
            self.people_custom_frame.pack_forget()
        self._refresh_people()

    def _reset_people_filters(self) -> None:
        if getattr(self, "_batch_running", False):
            return
        self._filter_resetting = True
        try:
            self.people_keyword_var.set("")
            self.people_field_var.set("全部字段")
            self.people_level_var.set("全部")
            self.people_mark_var.set("全部")
            self.people_min_var.set("")
            self.people_max_var.set("")
            self._people_col_filters = {}
            self._people_sort = ("", False)
            self.people_custom_frame.pack_forget()
        finally:
            self._filter_resetting = False
        self._update_people_sort_headings()
        self._refresh_people()

    def _display_name(self, item) -> str:
        name = display_text(item.get("name") or "")
        aliases = [display_text(value) for value in item.get("aliases") or []]
        if not aliases:
            return name
        shown = "、".join(aliases[:2])
        if len(aliases) > 2:
            shown += f" 等{len(aliases)}个"
        return f"{name}（曾用：{shown}）"

    def _people_threads(self):
        value = self.people_range_var.get()
        if value.startswith("全部"):
            return list(self.last_threads)
        for index, thread in enumerate(self.last_threads, start=1):
            if value.startswith(f"{index}. "):
                return [thread]
        return list(self.last_threads)

    def _refresh_supplement_rows(self) -> None:
        """只更新已补到的贴吧ID，不重建整张参与者表。"""
        by_uid: Dict[str, str] = {}
        for thread in getattr(self, "last_threads", []):
            for item in list(thread.floors) + list(thread.comments):
                uid = str(getattr(item, "author_id", "") or "")
                tieba_uid_value = str(getattr(item, "tieba_uid", "") or "")
                if uid and tieba_uid_value:
                    by_uid[uid] = tieba_uid_value
        if not by_uid:
            return
        tree = getattr(self, "people_tree", None)
        for user in getattr(self, "people_users", []):
            uid = str(user.get("uid") or "")
            value = by_uid.get(uid)
            if not value or user.get("tieba_uid") == value:
                continue
            user["tieba_uid"] = value
            iid = user.get("_iid")
            if tree is None or iid is None:
                continue
            try:
                values = list(tree.item(iid, "values"))
                if len(values) > 3:
                    values[3] = value
                    tree.item(iid, values=values)
            except tk.TclError:
                pass


    def _tieba_uid_cell(self, item) -> str:
        value = str(item.get("tieba_uid") or "")
        if value:
            return value
        uid = str(item.get("uid") or item.get("user_id") or "")
        if item.get("tieba_uid_failed") or (uid and uid in self._tieba_uid_failed):
            return "-"
        if uid or item.get("username") or item.get("portrait"):
            return "待补充"
        return "-"

    def _apply_tieba_failed_states(self, threads) -> None:
        failed = set(getattr(self, "_tieba_uid_failed", set()) or set())
        for thread in threads or []:
            for item in list(thread.floors) + list(thread.comments):
                uid = str(getattr(item, "author_id", "") or "")
                item.tieba_uid_failed = bool(uid and uid in failed and not item.tieba_uid)

    def _tieba_uid_obj_cell(self, item) -> str:
        value = str(getattr(item, "tieba_uid", "") or "")
        if value:
            return value
        uid = str(getattr(item, "author_id", "") or "")
        if getattr(item, "tieba_uid_failed", False) or (uid and uid in self._tieba_uid_failed):
            return "-"
        if uid or getattr(item, "author_name", "") or getattr(item, "portrait", ""):
            return "待补充"
        return "-"

    def _apply_user_id_failed_states(self, threads) -> None:
        failed = set(getattr(self, "_user_id_failed", set()) or set())
        for thread in threads or []:
            for item in list(thread.floors) + list(thread.comments):
                if not getattr(item, "author_id", ""):
                    item.user_id_failed = self._identity_fail_key(item) in failed

    def _person_uid_text(self, user) -> str:
        uid = str(user.get("uid") or "")
        if uid:
            return uid
        if user.get("user_id_failed") or self._identity_fail_key(user) in self._user_id_failed:
            return "-"
        if user.get("username") or user.get("portrait"):
            return "待补充"
        return "-"

    def _person_mark_cell(self, user) -> str:
        if not self._risk_enabled():
            return ""
        level = str(user.get("mark_level") or "")
        if not level:
            return "⚪"
        label = LEVEL_LABEL.get(level, "")
        suffix = " ⚠" if user.get("mark_hidden") else ""
        return f"{marker(level)} {label}{suffix}"

    def _set_person_mark_row(self, user) -> None:
        iid = user.get("_iid")
        if iid is None:
            return
        try:
            values = list(self.people_tree.item(iid, "values"))
            if not values:
                return
            values[0] = self._person_mark_cell(user)
            values[2] = str(user.get("username") or "-")
            values[3] = self._tieba_uid_cell(user)
            values[4] = self._person_uid_text(user)
            level = str(user.get("level") or "")
            if level:
                values[5] = f"Lv.{level}"
            self.people_tree.item(
                iid,
                values=values,
                tags=(str(user.get("mark_level") or ""),)
                if user.get("mark_level") and self.settings.get("color_people", True)
                else (),
            )
        except tk.TclError:
            pass

    def _refresh_people(self) -> None:
        values = ["全部帖子"]
        for index, thread in enumerate(self.last_threads, start=1):
            values.append(f"{index}. {display_text(thread.meta.title or thread.meta.tid)}")
        if hasattr(self, "people_range_box"):
            self.people_range_box.configure(values=values)
            if self.people_range_var.get() not in values:
                self.people_range_var.set("全部帖子")
        previous = {
            str(user.get("uid") or user.get("username") or user.get("name") or ""): user
            for user in getattr(self, "people_users", [])
        }
        users = self._aggregate_users(self._people_threads())
        for user in users:
            key = str(user.get("uid") or user.get("username") or user.get("name") or "")
            old = previous.get(key)
            if not old:
                continue
            for field in ("mark_level", "hit", "mark_bar", "mark_hidden", "mark_at", "mark_error"):
                if field in old:
                    user[field] = old[field]
            if not user.get("level") and old.get("level"):
                user["level"] = old["level"]
        self.people_users = users
        self._people_index = {id(user): index for index, user in enumerate(users)}
        for index, user in enumerate(users):
            user["_iid"] = str(index)
        filtered = filter_users(
            users,
            keyword=self.people_keyword_var.get(),
            field=self.people_field_var.get(),
            level=self.people_level_var.get(),
            mark=self.people_mark_var.get() if self._risk_enabled() else "全部",
            min_level=_optional_int(self.people_min_var.get()),
            max_level=_optional_int(self.people_max_var.get()),
        )
        for column, selected in self._people_col_filters.items():
            if selected:
                filtered = [
                    user for user in filtered
                    if self._people_column_value(user, column) in selected
                ]
        sort_column, sort_reverse = getattr(self, "_people_sort", ("", False))
        if sort_column:
            def sort_key(user):
                if sort_column == "mark":
                    order = {"警告": 3, "可疑": 2, "安全": 1, "": 0}
                    value = self._people_column_value(user, sort_column)
                    return order.get(value, 0)
                if sort_column == "level":
                    return int(user.get("level") or 0)
                return int(user.get(sort_column) or 0)
            filtered.sort(key=sort_key, reverse=bool(sort_reverse))
        page_size = max(1, int(self.settings.get("page_size", 1000) or 1000))
        pages = max(1, (len(filtered) + page_size - 1) // page_size)
        self._people_page = min(max(0, getattr(self, "_people_page", 0)), pages - 1)
        start = self._people_page * page_size
        for item in self.people_tree.get_children():
            self.people_tree.delete(item)
        for user in filtered[start : start + page_size]:
            index = self._people_index[id(user)]
            self.people_tree.insert(
                "",
                "end",
                iid=str(index),
                values=(
                    self._person_mark_cell(user),
                    self._display_name(user),
                    user["username"] or "-",
                    self._tieba_uid_cell(user),
                    self._person_uid_text(user),
                    f"Lv.{user['level']}" if user["level"] else "-",
                    user["floor"],
                    user["comment"],
                    user["total"],
                ),
            )
        if hasattr(self, "people_page_label"):
            self.people_page_label.configure(
                text=f"第 {self._people_page + 1}/{pages} 页 · 筛选后 {len(filtered)} / 共 {len(users)} 人"
            )
        with_level = sum(1 for user in filtered if user.get("level"))
        mode_note = ""
        if self._risk_enabled():
            mode_note = " · 红绿灯：仅发言" if self.mode_var.get() == "guest" else " · 红绿灯：发言+关注"
        warning_count = sum(len(thread.warnings) for thread in self.last_threads)
        warning_note = f" · ⚠ {warning_count} 条数据警告" if warning_count else ""
        stale_note = " · 标记规则已变更，可点重新标记（本地）" if getattr(self, "_marks_stale", False) else ""
        self.people_hint.configure(
            text=f"v{VERSION} · 当前 {len(filtered)} / 共 {len(users)} 人 · 有等级 {with_level}{mode_note}{warning_note}{stale_note}"
        )
        if hasattr(self, "retry_incomplete_btn"):
            incomplete = any(not getattr(thread, "complete", True) for thread in self.last_threads)
            if incomplete:
                self.retry_incomplete_btn.state(["!disabled"])
            else:
                self.retry_incomplete_btn.state(["disabled"])
        if hasattr(self, "supplement_btn"):
            missing_identity = any(
                (not user.get("uid") and (user.get("username") or user.get("portrait")))
                or (user.get("uid") and not user.get("tieba_uid"))
                for user in self.people_users
            )
            self.supplement_btn.configure(state="normal" if missing_identity else "disabled")
        self.profile_status.configure(text="选择左边一个人查看详情")
        self._set_profile_cache_note("")
        for tree in (self.profile_forum_tree, self.profile_post_tree):
            for item in tree.get_children():
                tree.delete(item)

    def _people_go(self, value) -> None:
        users = getattr(self, "people_users", [])
        page_size = max(1, int(self.settings.get("page_size", 1000) or 1000))
        pages = max(1, (len(users) + page_size - 1) // page_size)
        if value == "first":
            self._people_page = 0
        elif value == "last":
            self._people_page = pages - 1
        else:
            self._people_page = min(max(0, self._people_page + int(value)), pages - 1)
        self._refresh_people()

    def _on_people_paned_configure(self, event=None) -> None:
        if getattr(self, "_people_sash_set", False):
            return
        if event is not None and event.width < 500:
            return
        self._people_sash_set = True
        self.root.after_idle(self._set_people_sash)

    def _set_people_sash(self) -> None:
        try:
            width = max(520, int(self.root.winfo_width() * 0.58))
            self.people_paned.sashpos(0, width)
        except tk.TclError:
            pass

    def _on_people_range_changed(self) -> None:
        if getattr(self, "_batch_running", False):
            return
        self._people_page = 0
        self._refresh_people()

    def _on_person_select(self, _event=None) -> None:
        selection = self.people_tree.selection()
        if not selection:
            return
        try:
            user = self.people_users[int(selection[0])]
        except (IndexError, ValueError):
            return
        self._profile_topic_limit = 50
        self._load_profile(user)


    # ------------------------------------------------------------- risk marks
    def _risk_enabled(self) -> bool:
        return bool(self.settings.get("risk_enabled", False))

    def _has_mark_rules(self) -> bool:
        return any(self.blacklist.get(key) for key in ("safe", "warn", "ban"))

    def _apply_risk_ui(self) -> None:
        enabled = self._risk_enabled()
        if hasattr(self, "people_tree"):
            columns = getattr(self, "people_columns", ())
            display = columns if enabled else [column for column in columns if column != "mark"]
            try:
                self.people_tree.configure(displaycolumns=display)
            except tk.TclError:
                pass
        if hasattr(self, "batch_btn") and not getattr(self, "_batch_running", False):
            self.batch_btn.configure(state="normal" if enabled else "disabled", text="一键红绿灯")
        if hasattr(self, "remark_btn") and not getattr(self, "_batch_running", False):
            self.remark_btn.state(["!disabled"] if enabled else ["disabled"])
        if hasattr(self, "include_follows_check") and not getattr(self, "_batch_running", False):
            allow_follow = enabled and bool(self.settings.get("follow_enhance"))
            self.include_follows_check.configure(state="normal" if allow_follow else "disabled")
            if not allow_follow:
                self.include_follows_var.set(False)


    def _on_include_follows_changed(self) -> None:
        if getattr(self, "_batch_running", False):
            return
        if self.include_follows_var.get() and not self.settings.get("follow_enhance"):
            self.include_follows_var.set(False)

    def _on_follow_toggle(self) -> None:
        enabled = bool(self.follow_enhance_var.get())
        if not enabled:
            self.settings["follow_enhance"] = False
            save_settings(self.settings)
            self._refresh_webview_status()
            return
        has_login = self.mode_var.get() == "login" and webview2.has_login_cookie(
            self._request_cookie()
        )
        if not has_login:
            self.follow_enhance_var.set(False)
            self.settings["follow_enhance"] = False
            save_settings(self.settings)
            messagebox.showinfo(
                "需要登录信息",
                "请先点「登录百度」或「手动填写」完成登录，再启用关注检测。",
                parent=self.root,
            )
            self._refresh_webview_status()
            return
        self.settings["follow_enhance"] = True
        save_settings(self.settings)
        self._refresh_webview_status()



    def _on_follow_load_mode_changed(self) -> None:
        value = "manual" if self.follow_load_mode_var.get() == "手动" else "auto"
        self.settings["follow_load_mode"] = value
        save_settings(self.settings)

    def _on_risk_toggle(self) -> None:
        enabled = bool(self.risk_var.get())
        if enabled:
            ok = messagebox.askyesno(
                "启用本地标记",
                """本地标记只用于你自己的记录，请勿用于挂人、骚扰或公开曝光。

确认启用？""",
                parent=self.root,
            )
            if not ok:
                self.risk_var.set(False)
                return
        self.settings["risk_enabled"] = enabled
        save_settings(self.settings)
        self._apply_risk_ui()
        if getattr(self, "people_users", None):
            self._refresh_people()

    # ------------------------------------------------------------------ mode
    def _apply_mode_ui(self) -> None:
        login = self.mode_var.get() == "login"
        state = "normal" if login else "disabled"
        for widget in getattr(self, "cookie_controls", []):
            try:
                widget.configure(state=state)
            except tk.TclError:
                pass
        if hasattr(self, "cookie_box"):
            if login:
                self.cookie_box.grid()
            else:
                self.cookie_box.grid_remove()
        if hasattr(self, "mode_hint"):
            self.mode_hint.configure(
                text="完整楼中楼 · 关注列表 · 完整资料"
                if login
                else "公开帖子 · 楼中楼按设置补全 · 红绿灯只看发言"
            )
        self._apply_risk_ui()
        self._update_risk_summary()



    def _set_mode(self, mode: str, persist: bool = True) -> None:
        mode = mode if mode in ("guest", "login") else "guest"
        self.mode_var.set(mode)
        if persist:
            self.settings["mode"] = mode
            save_settings(self.settings)
        self._apply_mode_ui()
        if hasattr(self, "_refresh_webview_status"):
            self._refresh_webview_status()

    def _make_scrollable_settings_page(self, parent) -> tuple:
        outer = tk.Frame(parent, bg="#eef1f5")
        canvas = tk.Canvas(outer, bg="#eef1f5", highlightthickness=0)
        scrollbar = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        body = tk.Frame(canvas, bg="#eef1f5")
        window = canvas.create_window((0, 0), window=body, anchor="nw")
        body.bind(
            "<Configure>",
            lambda _event: canvas.configure(scrollregion=canvas.bbox("all")),
        )
        canvas.bind(
            "<Configure>",
            lambda event: canvas.itemconfigure(window, width=event.width),
        )

        def on_wheel(event) -> None:
            first, last = canvas.yview()
            if event.delta > 0 and first <= 0.0:
                return
            if event.delta < 0 and last >= 1.0:
                return
            canvas.yview_scroll(int(-event.delta / 120), "units")

        canvas.bind("<MouseWheel>", on_wheel)
        body.bind("<MouseWheel>", on_wheel)
        return outer, body

    def _bind_mousewheel_tree(self, widget, canvas) -> None:
        if isinstance(widget, (tk.Text, tk.Listbox, ttk.Treeview)):
            return

        def on_wheel(event) -> str:
            first, last = canvas.yview()
            if (event.delta > 0 and first <= 0.0) or (event.delta < 0 and last >= 1.0):
                return "break"
            canvas.yview_scroll(int(-event.delta / 120), "units")
            return "break"

        widget.bind("<MouseWheel>", on_wheel)
        for child in widget.winfo_children():
            self._bind_mousewheel_tree(child, canvas)

    def _build_links_tab(self) -> None:
        page = self.tab_links
        top = tk.Frame(page, bg="#ffffff", padx=12, pady=8)
        top.pack(fill="x")
        tk.Label(
            top,
            text="每行一个帖子链接（支持主题 / 楼层 / 楼中楼）",
            bg="#ffffff",
            fg="#475569",
            font=("Microsoft YaHei UI", 9),
        ).pack(anchor="w")
        links_wrap = tk.Frame(top, bg="#ffffff")
        links_wrap.pack(fill="x", pady=(4, 4))
        self.links_text = tk.Text(links_wrap, height=6, wrap="none", font=("Consolas", 9))
        links_yscroll = ttk.Scrollbar(links_wrap, orient="vertical", command=self.links_text.yview)
        links_xscroll = ttk.Scrollbar(links_wrap, orient="horizontal", command=self.links_text.xview)
        self.links_text.configure(yscrollcommand=links_yscroll.set, xscrollcommand=links_xscroll.set)
        self.links_text.grid(row=0, column=0, sticky="nsew")
        links_yscroll.grid(row=0, column=1, sticky="ns")
        links_xscroll.grid(row=1, column=0, sticky="ew")
        links_wrap.rowconfigure(0, weight=1)
        links_wrap.columnconfigure(0, weight=1)
        bar = tk.Frame(top, bg="#ffffff")
        bar.pack(fill="x")
        ttk.Button(bar, text="从剪贴板导入", command=self._links_from_clipboard).pack(side="left")
        ttk.Button(bar, text="开始解析", command=self._start_links).pack(side="left", padx=(8, 0))
        ttk.Button(bar, text="停止", command=self._stop_links).pack(side="left", padx=(8, 0))
        ttk.Button(bar, text="打开选中链接", command=self._open_links_link).pack(side="left", padx=(8, 0))
        ttk.Button(bar, text="导出 Excel", command=self._export_links).pack(side="left", padx=(8, 0))
        self.links_status = tk.Label(
            bar, text="", bg="#ffffff", fg="#64748b", font=("Microsoft YaHei UI", 9)
        )
        self.links_status.pack(side="right")

        wrap = tk.Frame(page, bg="#eef1f5")
        wrap.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        self.links_tree = ttk.Treeview(
            wrap,
            columns=("time", "type", "bar", "title", "content", "floor", "link"),
            show="headings",
            height=18,
        )
        for key, title, width in (
            ("time", "日期", 120),
            ("type", "类型", 60),
            ("bar", "贴吧", 100),
            ("title", "标题", 220),
            ("content", "内容", 300),
            ("floor", "楼层", 60),
            ("link", "链接", 260),
        ):
            self.links_tree.heading(key, text=title)
            self.links_tree.column(key, width=width, anchor="center" if key in ("time", "type", "floor") else "w")
        yscroll = ttk.Scrollbar(wrap, orient="vertical", command=self.links_tree.yview)
        xscroll = ttk.Scrollbar(wrap, orient="horizontal", command=self.links_tree.xview)
        self.links_tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.links_tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        wrap.rowconfigure(0, weight=1)
        wrap.columnconfigure(0, weight=1)
        self._links_urls: Dict[str, str] = {}
        self.links_tree.bind("<Double-1>", lambda _event: self._open_links_link())
        self._links_rows: List[Dict] = []
        self._links_errors: List[Dict] = []
        self._links_stop = threading.Event()
        self._links_running = False

    def _open_links_link(self) -> None:
        selection = self.links_tree.selection()
        if selection:
            url = self._links_urls.get(selection[0], "")
            if url:
                webbrowser.open(url)

    def _links_from_clipboard(self) -> None:
        try:
            text = self.root.clipboard_get()
        except tk.TclError:
            return
        self.links_text.delete("1.0", "end")
        self.links_text.insert("1.0", text)

    def _stop_links(self) -> None:
        self._links_stop.set()
        self.links_status.configure(text="正在停止…")

    def _start_links(self) -> None:
        if self._links_running:
            return
        raw = self.links_text.get("1.0", "end").strip()
        if not raw:
            self.links_status.configure(text="请先粘贴链接")
            return
        self._links_running = True
        self._links_stop.clear()
        self._links_rows = []
        self._links_errors = []
        self._links_urls = {}
        for item in self.links_tree.get_children():
            self.links_tree.delete(item)
        self.links_status.configure(text="解析中…")

        def worker():
            try:
                fetcher = Fetcher(cookie={}, delay=0.3, verbose=False)
                result = parse_links(
                    raw,
                    fetcher,
                    should_stop=self._links_stop.is_set,
                )
                self._post(("links_ready", result, ""))
            except Exception as error:  # noqa: BLE001
                self._post(("links_ready", None, str(error)))

        threading.Thread(target=worker, daemon=True).start()

    def _render_links(self, result: Dict, error: str) -> None:
        self._links_running = False
        if result is None:
            self.links_status.configure(text=f"失败：{error}")
            return
        self._links_rows = list(result.get("rows") or [])
        self._links_errors = list(result.get("errors") or [])
        for index, row in enumerate(self._links_rows):
            iid = str(index)
            self.links_tree.insert(
                "",
                "end",
                iid=iid,
                values=(
                    row.get("time") or "",
                    row.get("type") or "",
                    row.get("bar") or "",
                    row.get("title") or "",
                    row.get("content") or "",
                    row.get("floor") or "",
                    row.get("url") or "",
                ),
            )
            self._links_urls[iid] = row.get("url") or ""
        extra = f" / 失败 {len(self._links_errors)}" if self._links_errors else ""
        self.links_status.configure(
            text=f"链接 {result.get('links', 0)} / 主题 {result.get('threads', 0)} / 成功 {len(self._links_rows)}{extra}"
        )

    def _export_links(self) -> None:
        if not self._links_rows:
            self.links_status.configure(text="没有可导出的结果")
            return
        path = filedialog.asksaveasfilename(
            parent=self.root,
            title="导出链接解析",
            defaultextension=".xlsx",
            initialfile="链接解析.xlsx",
            filetypes=[("Excel 表格", "*.xlsx")],
        )
        if not path:
            return
        try:
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "链接解析"
            headers = ["日期", "类型", "贴吧", "标题", "内容", "楼层", "链接"]
            ws.append(headers)
            for row in self._links_rows:
                ws.append(
                    [
                        row.get("time") or "",
                        row.get("type") or "",
                        row.get("bar") or "",
                        row.get("title") or "",
                        row.get("content") or "",
                        row.get("floor") or "",
                        row.get("url") or "",
                    ]
                )
            ws2 = wb.create_sheet("说明")
            ws2.append(["链接数", len(self._links_rows) + len(self._links_errors)])
            ws2.append(["成功", len(self._links_rows)])
            ws2.append(["失败", len(self._links_errors)])
            ws2.append(["数据来源", "官方主题抓取"])
            wb.save(path)
            self.links_status.configure(text="已导出 Excel")
        except Exception as error:  # noqa: BLE001
            messagebox.showerror("导出失败", str(error), parent=self.root)


    def _popup_tree_menu(self, tree, menu, event) -> None:
        row = tree.identify_row(event.y)
        if row:
            tree.selection_set(row)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()












    def _show_settings_page(self, index: int) -> None:
        for position, page in enumerate(self.settings_pages):
            if position == index:
                page.pack(fill="both", expand=True)
            else:
                page.pack_forget()
        self._set_cache_stats_timer(index == 4)

    def _on_settings_start_changed(self) -> None:
        self.settings["settings_start"] = "last" if self.settings_start_var.get() == "上次分类" else "general"
        save_settings(self.settings)

    def _on_main_tab_changed(self, _event=None) -> None:
        if self.notebook.select() != str(self.tab_settings):
            self._set_cache_stats_timer(False)
            return
        index = 0
        if self.settings.get("settings_start") == "last":
            index = getattr(self, "_last_settings_index", 0)
        if 0 <= index < self.settings_nav.size():
            self.settings_nav.selection_clear(0, "end")
            self.settings_nav.selection_set(index)
            self._show_settings_page(index)

    def _on_settings_nav(self, _event=None) -> None:
        selection = self.settings_nav.curselection()
        if selection:
            self._last_settings_index = selection[0]
            self._show_settings_page(selection[0])

    def _focus_local_marks(self) -> None:
        self.notebook.select(self.tab_settings)
        self.root.after_idle(self._show_local_marks_page)

    def _show_local_marks_page(self) -> None:
        self.settings_nav.selection_clear(0, "end")
        self.settings_nav.selection_set(3)
        self._last_settings_index = 3
        self._show_settings_page(3)
        editor = self.blacklist_editors.get("safe")
        if editor is not None:
            editor.focus_set()
            editor.see("1.0")


    def _on_settings_mode_selected(self) -> None:
        mode = self.mode_var.get()
        if mode == "guest" and bool(self.settings.get("follow_enhance", False)):
            if not messagebox.askyesno(
                "切换到游客模式",
                "游客模式不使用关注检测。\n切换会关闭关注检测，是否继续？",
                parent=self.root,
                default=messagebox.NO,
            ):
                self.mode_var.set("login")
                return
            self.settings["follow_enhance"] = False
            if hasattr(self, "follow_enhance_var"):
                self.follow_enhance_var.set(False)
            save_settings(self.settings)
        self._set_mode(mode)
        self._refresh_webview_status()















    def _prepare_accent_popup(self) -> None:
        """让强调色下拉列表只使用当前主题的选中色，避免残留粉色/橘色行。"""
        try:
            popdown = str(
                self.root.tk.call("ttk::combobox::PopdownWindow", self.accent_combo)
            )
            listbox = f"{popdown}.f.l"
            self.root.tk.call(
                listbox,
                "configure",
                "-background",
                self.palette["surface"],
                "-foreground",
                self.palette["text"],
                "-selectbackground",
                self.palette["selection_bg"],
                "-selectforeground",
                self.palette["selection_text"],
                "-activestyle",
                "none",
            )
        except tk.TclError:
            pass

    def _on_accent_changed(self) -> None:
        accent = {
            "蓝色": "blue",
            "绿色": "green",
            "紫色": "purple",
            "橘黄": "orange",
            "粉色": "pink",
            "青色": "cyan",
            "红色": "red",
        }.get(self.accent_var.get(), "blue")
        self.settings["accent"] = accent
        save_settings(self.settings)
        self._apply_theme()









    def _center_toplevel(self, win: tk.Toplevel, width: int = 0, height: int = 0) -> None:
        win.update_idletasks()
        width = width or win.winfo_reqwidth()
        height = height or win.winfo_reqheight()
        x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - width) // 2)
        y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - height) // 2)
        win.geometry(f"+{x}+{y}")


    def _commit_cache_limit(self) -> None:
        now = time.monotonic()
        if getattr(self, "_cache_limit_busy", False) or now < getattr(
            self, "_cache_limit_guard_until", 0
        ):
            current = str(int(self.settings.get("post_cache_limit_mb", DEFAULT_LIMIT_MB)))
            try:
                if self.cache_limit_var.get() != current:
                    self.cache_limit_var.set(current)
            except tk.TclError:
                pass
            return
        self._cache_limit_busy = True
        box = getattr(self, "cache_limit_box", None)
        try:
            if box is not None:
                box.configure(state="disabled")
            try:
                value = int(float(self.cache_limit_var.get().strip()))
            except (TypeError, ValueError):
                value = int(self.settings.get("post_cache_limit_mb", DEFAULT_LIMIT_MB))
            value = max(MIN_LIMIT_MB, min(MAX_LIMIT_MB, value))
            old = int(self.settings.get("post_cache_limit_mb", DEFAULT_LIMIT_MB))
            if value < old:
                if not messagebox.askokcancel(
                    "降低发言缓存上限",
                    "降低上限会立即清除部分发言缓存，之后查询可能变慢。\n"
                    "不会删除 Excel、本地标记或最近结果。\n\n继续？",
                    parent=self.root,
                ):
                    self.cache_limit_var.set(str(old))
                    return
            self.cache_limit_var.set(str(value))
            self.settings["post_cache_limit_mb"] = value
            save_settings(self.settings)
            post_cache().set_limit_mb(value)
            self._refresh_cache_stats()
        finally:
            if box is not None:
                try:
                    box.configure(state="normal")
                except tk.TclError:
                    pass
            self._cache_limit_busy = False
            self._cache_limit_guard_until = time.monotonic() + 0.35

    def _set_cache_stats_timer(self, enabled: bool) -> None:
        timer = getattr(self, "_cache_stats_timer", None)
        if timer is not None:
            try:
                self.root.after_cancel(timer)
            except tk.TclError:
                pass
            self._cache_stats_timer = None
        if enabled:
            self._cache_stats_timer = self.root.after(2000, self._cache_stats_tick)

    def _cache_stats_tick(self) -> None:
        self._cache_stats_timer = None
        self._notify_cache_eviction()
        self._refresh_cache_stats()
        self._cache_stats_timer = self.root.after(2000, self._cache_stats_tick)

    def _notify_cache_eviction(self) -> None:
        if getattr(self, "_cache_eviction_notified", False):
            post_cache().take_eviction_notice()
            return
        count = post_cache().take_eviction_notice()
        if not count:
            return
        self._cache_eviction_notified = True
        text = "缓存已达上限，已淘汰部分旧缓存；不影响本次标记和导出。"
        self.status_var.set(text)
        self._append_log(text, "muted")

    def _refresh_cache_stats_if_visible(self) -> None:
        if getattr(self, "_cache_stats_timer", None) is not None:
            self._refresh_cache_stats()

    def _refresh_cache_stats(self) -> None:
        label = getattr(self, "cache_stats_label", None)
        if label is None:
            return
        try:
            stats = post_cache().stats()
            label.configure(
                text=(
                    f"热内存：{stats['hot_users']} 人 / {stats['hot_bytes'] / 1024 / 1024:.1f} MB；"
                    f"冷库：{stats['cold_users']} 人 / {stats['cold_posts']} 条 / "
                    f"{stats['cold_bytes'] / 1024 / 1024:.1f} MB（上限 {stats['limit_mb']} MB）"
                )
            )
        except Exception as error:  # noqa: BLE001
            label.configure(text=f"缓存状态读取失败：{error}")

    def _clear_post_cache(self) -> None:
        if not messagebox.askyesno(
            "清除发言缓存",
            "清除本次运行的发言缓存？\n不会删除 Excel、本地标记或最近结果。",
            parent=self.root,
        ):
            return
        clear_post_cache()
        for name in ("_lookup_cache", "_profile_cache", "_batch_cache"):
            value = getattr(self, name, None)
            if isinstance(value, dict):
                value.clear()
        self._refresh_cache_stats()
        self._append_log("发言缓存已清除。", "muted")

    def _build_about_tab(self) -> None:
        box = tk.Frame(self.tab_about, bg="#eef1f5", padx=14, pady=10)
        box.pack(fill="both", expand=True)
        tk.Label(
            box,
            text=f"{APP_TITLE} v{VERSION}",
            bg="#eef1f5",
            fg="#1F4E79",
            font=("Microsoft YaHei UI", 12, "bold"),
        ).pack(anchor="w")
        tk.Label(
            box,
            text="作者：OK · GitHub：Goto-OK",
            bg="#eef1f5",
            fg="#64748b",
            font=("Microsoft YaHei UI", 8),
        ).pack(anchor="w", pady=(2, 2))
        tk.Label(
            box,
            text="默认匿名 · 结果只存本机 · 不保存 STOKEN · 支持浅色/深色",
            bg="#eef1f5",
            fg="#8a94a6",
            font=("Microsoft YaHei UI", 8),
        ).pack(anchor="w", pady=(2, 8))

        help_box = ttk.LabelFrame(box, text=" 帮助与诊断 ")
        help_box.pack(fill="x", pady=(0, 6))
        help_row = tk.Frame(help_box, bg="#eef1f5")
        help_row.pack(anchor="w", padx=8, pady=(6, 2))
        ttk.Button(help_row, text="打开使用说明", command=self._open_usage_window, width=14).pack(side="left")
        ttk.Button(help_row, text="第三方许可", command=self._open_notice, width=12).pack(
            side="left", padx=(8, 0)
        )
        ttk.Button(help_row, text="复制诊断信息", command=self._copy_diagnostics, width=14).pack(
            side="left", padx=(8, 0)
        )
        tk.Label(
            help_box,
            text="第三方组件：Python / PyInstaller / lxml / openpyxl / Tcl-Tk，各自按其许可证分发。",
            bg="#eef1f5",
            fg="#8a94a6",
            font=("Microsoft YaHei UI", 8),
            wraplength=760,
            justify="left",
        ).pack(anchor="w", padx=8, pady=(2, 4))

    def _on_save_result_toggle(self) -> None:
        self.settings["save_result"] = bool(self.save_result_var.get())
        save_settings(self.settings)
        if not self.save_result_var.get():
            clear_stored_result()

    def _on_encrypt_toggle(self) -> None:
        self.settings["encrypt_result"] = bool(self.encrypt_var.get())
        save_settings(self.settings)
        if self.last_threads:
            self._persist_result(self.last_threads)


    def _on_restore_mode_changed(self) -> None:
        self.settings["restore_mode"] = "ask" if self.restore_mode_var.get() == "询问" else "auto"
        save_settings(self.settings)

    def _clear_result(self) -> None:
        if not messagebox.askyesno(
            "删除当前结果",
            "删除本机保存的当前结果？不会删除已导出的 Excel、Markdown 或本地标记；删除后不可撤销。",
            parent=self.root,
            default=messagebox.NO,
        ):
            return
        clear_stored_result()
        self.last_threads = []
        self.last_output = None
        self.open_btn.configure(state="disabled")
        self._refresh_people()
        self._refresh_grid()
        self._append_log("当前结果已删除。", "muted")

    def _cleanup_old_results(self) -> None:
        files = stored_old_files()
        if not files:
            self.status_var.set("没有需要清理的旧版本结果")
            return
        if not messagebox.askyesno(
            "清理旧版本结果",
            f"删除 {len(files)} 个旧版本结果文件？不会影响当前结果。",
            parent=self.root,
            default=messagebox.NO,
        ):
            return
        for path in files:
            try:
                safe_remove_file(path)
            except (OSError, ValueError):
                pass
        self._refresh_old_files_label()
        self._append_log(f"已清理旧版本结果：{len(files)} 个", "muted")

    def _refresh_old_files_label(self) -> None:
        if not hasattr(self, "old_files_label"):
            return
        files = stored_old_files()
        size = 0
        for path in files:
            try:
                size += os.path.getsize(path)
            except OSError:
                pass
        self.old_files_label.configure(
            text=f"旧版本残留：{len(files)} 个，{size / 1024 / 1024:.1f} MB"
        )

    def _open_notice(self) -> None:
        path = os.path.join(resource_dir(), "THIRD_PARTY_NOTICES.md")
        if not os.path.isfile(path):
            messagebox.showinfo("第三方许可", "许可文件未找到。", parent=self.root)
            return
        existing = getattr(self, "_notice_window", None)
        try:
            if existing is not None and existing.winfo_exists():
                existing.lift()
                existing.focus_set()
                return
        except tk.TclError:
            pass
        try:
            with open(path, "r", encoding="utf-8") as handle:
                content = handle.read()
        except OSError as error:
            messagebox.showerror("打不开许可文件", friendly_error(error), parent=self.root)
            return
        win = tk.Toplevel(self.root)
        self._notice_window = win
        win.title("第三方许可")
        win.geometry("760x600")
        win.minsize(560, 360)
        win.configure(bg="#eef1f5")
        win.transient(self.root)
        frame = tk.Frame(win, bg="#eef1f5", padx=10, pady=8)
        frame.pack(fill="both", expand=True)
        top = tk.Frame(frame, bg="#eef1f5")
        top.pack(fill="x", pady=(0, 6))
        tk.Label(
            top, text="第三方组件与许可", bg="#eef1f5", fg="#1F4E79",
            font=("Microsoft YaHei UI", 12, "bold"),
        ).pack(side="left")
        ttk.Button(top, text="关闭", command=win.destroy, width=8).pack(side="right")
        text = tk.Text(
            frame, wrap="word", font=("Microsoft YaHei UI", 9), relief="flat",
            padx=10, pady=8, bg="#fbfcfe", state="disabled",
        )
        scroll = ttk.Scrollbar(frame, orient="vertical", command=text.yview)
        text.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        text.pack(side="left", fill="both", expand=True)
        text.configure(state="normal")
        text.insert("1.0", content)
        text.configure(state="disabled")
        self._theme_widget_tree(win)
        self._center_toplevel(win, 760, 600)
        win.protocol("WM_DELETE_WINDOW", win.destroy)
        win.focus_set()

    def _open_folder(self, path: str) -> None:
        try:
            os.makedirs(path, exist_ok=True)
            os.startfile(path)  # type: ignore[attr-defined]
        except OSError as error:
            messagebox.showerror("打不开目录", friendly_error(error), parent=self.root)



    def _offline_checks(self) -> list[tuple[str, str, str]]:
        results: list[tuple[str, str, str]] = []

        def check(name: str, condition: bool, detail: str = "") -> None:
            results.append((name, "PASS" if condition else "FAIL", detail))

        check(
            "允许百度登录域名",
            webview2.is_allowed_url("https://passport.baidu.com/v2/?login"),
        )
        check(
            "拒绝非白名单域名",
            not webview2.is_allowed_url("https://example.com/"),
        )
        check(
            "接受 TIEBA_SID 登录态",
            webview2.has_login_cookie({"TIEBA_SID": "test"}),
        )
        check(
            "匿名 Cookie 不算登录",
            not webview2.has_login_cookie({"BAIDUID": "test"}),
        )
        filtered = webview2.filter_cookies(
            [
                {"name": "BDUSS", "value": "secret", "domain": ".baidu.com", "path": "/"},
                {"name": "BAIDUID", "value": "anonymous", "domain": ".baidu.com", "path": "/"},
            ],
            anonymous=True,
        )
        check("游客过滤 BDUSS", "BDUSS" not in filtered and "BAIDUID" in filtered)
        merged = webview2.filter_cookies(
            [
                {"name": "X", "value": "root", "domain": ".baidu.com", "path": "/"},
                {"name": "X", "value": "tieba", "domain": ".tieba.baidu.com", "path": "/p/1"},
            ],
            anonymous=False,
        )
        check("同名 Cookie 取更具体 path", merged.get("X") == "tieba")
        check("WebView2 组件可用", webview2.available(), "仅信息，不强制安装")
        return results



    def _simulate_captcha(self, _event=None, scenario: str = "") -> None:
        """隐藏测试入口：不发网络请求，只走验证码处理流程。"""
        if self._network_task_active() or getattr(self, "_active_choice_window", None):
            return
        restore = None
        if scenario in ("guest", "cookie", "bduss"):
            restore = (
                self.mode_var.get(),
                self.login_method_var.get(),
            )
            if scenario == "guest":
                self.mode_var.set("guest")
            else:
                self.mode_var.set("login")
                self.login_method_var.set(scenario)
                self._on_login_method_changed()
        self._simulated_restore = restore
        self._append_log("模拟验证码：未发出网络请求，仅测试界面流程。", "muted")
        self._simulated_captcha = True
        self._set_running(True)
        self.root.after(
            20,
            lambda: self._handle_message(
                (
                    "captcha",
                    "模拟验证码：未发出网络请求，仅用于测试界面流程。",
                    [],
                )
            ),
        )





    def _ask_choice(self, title: str, message: str, choices, default: int = 0):
        """带明确按钮文案的模态选择框；choices=[(label,value),...]，None 表示取消。"""
        win = tk.Toplevel(self.root)
        win.title(title)
        win.configure(bg="#ffffff")
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        self._active_choice_window = win
        frame = tk.Frame(win, bg="#ffffff", padx=20, pady=16)
        frame.pack(fill="both", expand=True)
        tk.Label(
            frame,
            text=title,
            bg="#ffffff",
            fg="#1f2937",
            font=("Microsoft YaHei UI", 12, "bold"),
        ).pack(anchor="w")
        tk.Label(
            frame,
            text=message,
            bg="#ffffff",
            fg="#475569",
            font=("Microsoft YaHei UI", 9),
            wraplength=440,
            justify="left",
            anchor="w",
        ).pack(fill="x", pady=(8, 14))
        result: Dict[str, Any] = {"value": None}
        buttons = tk.Frame(frame, bg="#ffffff")
        buttons.pack(fill="x")

        def finish(value) -> None:
            result["value"] = value
            win.destroy()

        def command(value):
            return lambda: finish(value)

        for index in range(len(choices) - 1, -1, -1):
            label, value = choices[index]
            if index == default:
                tk.Button(
                    buttons,
                    text=label,
                    command=command(value),
                    bg="#1F4E79",
                    fg="#ffffff",
                    activebackground="#163b5c",
                    activeforeground="#ffffff",
                    relief="flat",
                    padx=14,
                    pady=5,
                    cursor="hand2",
                    font=("Microsoft YaHei UI", 9, "bold"),
                ).pack(side="right", padx=(8, 0))
            else:
                ttk.Button(buttons, text=label, command=command(value)).pack(
                    side="right", padx=(8, 0)
                )
        win.bind("<Return>", lambda _event: finish(choices[default][1]))
        win.bind("<Escape>", lambda _event: finish(None))
        win.protocol("WM_DELETE_WINDOW", lambda: finish(None))
        win.update_idletasks()
        width = max(win.winfo_reqwidth(), buttons.winfo_reqwidth() + 40)
        height = win.winfo_reqheight()
        x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - width) // 2)
        y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - height) // 2)
        self._theme_widget_tree(win)
        win.geometry(f"{width}x{height}+{x}+{y}")
        self.root.wait_window(win)
        if getattr(self, "_active_choice_window", None) is win:
            self._active_choice_window = None
        return result["value"]

    def _hidden_advanced_settings(self) -> List[str]:
        names: List[str] = []
        if str(self.settings.get("follow_load_mode") or "auto") == "manual":
            names.append("关注加载：手动")
        strategy = str(self.settings.get("identity_strategy") or "full")
        labels = {
            "ask": "身份策略：询问",
            "skip": "身份策略：跳过",
            "user_id": "身份策略：只补用户ID",
            "detail": "身份策略：详情再补",
        }
        if strategy in labels:
            names.append(labels[strategy])
        try:
            concurrency = int(self.settings.get("identity_concurrency", 3))
        except (TypeError, ValueError):
            concurrency = 3
        if concurrency != 3:
            names.append(f"身份并发 {concurrency}")
        try:
            batch = int(self.settings.get("batch_concurrency", 1))
        except (TypeError, ValueError):
            batch = 1
        if batch > 1:
            names.append(f"红绿灯并发 {batch}")
        if str(self.settings.get("proxy_mode") or "system") == "custom":
            names.append("自定义代理")
        if bool(self.settings.get("lookup_fill_threads")):
            names.append("查用户回抓主题")
        if bool(self.settings.get("log_enabled")):
            names.append("运行日志")
        if bool(self.settings.get("encrypt_result")):
            names.append("结果加密")
        try:
            cache_mb = int(self.settings.get("post_cache_limit_mb", DEFAULT_LIMIT_MB))
        except (TypeError, ValueError):
            cache_mb = DEFAULT_LIMIT_MB
        if cache_mb != DEFAULT_LIMIT_MB:
            names.append(f"缓存上限 {cache_mb} MB")
        try:
            comment_concurrency = int(self.comment_concurrency_var.get() or 1)
        except (TypeError, ValueError):
            comment_concurrency = 1
        if comment_concurrency > 1:
            names.append(f"楼中楼并发 {comment_concurrency}")
        try:
            delay = float(self.delay_var.get() or 0.8)
        except (TypeError, ValueError):
            delay = 0.8
        if abs(delay - 0.8) > 0.001:
            names.append(f"请求间隔 {delay:g} 秒")
        if str(self.comment_depth_var.get() or "全部") != "全部":
            names.append(f"楼中楼：{self.comment_depth_var.get()}")
        return names

    def _refresh_advanced_summary(self) -> None:
        if not hasattr(self, "ui_mode_var"):
            return
        simple = self.ui_mode_var.get() != "advanced"
        names = self._hidden_advanced_settings() if simple else []
        if not simple or not names:
            self.fetch_advanced_summary.pack_forget()
            self.fetch_advanced_summary.unbind("<Button-1>")
            self.ui_mode_hint.configure(
                text="简洁模式隐藏调参和高风险操作", cursor=""
            )
            self.ui_mode_hint.unbind("<Button-1>")
            return
        shown = "、".join(names[:2])
        if len(names) > 2:
            shown += f" 等 {len(names)} 项"
        text = f"高级设置生效中：{shown}  [管理]"
        self.fetch_advanced_summary.configure(text=text)
        self.fetch_advanced_summary.pack(
            fill="x", padx=8, pady=(4, 0), before=self.fetch_page
        )
        self.fetch_advanced_summary.configure(cursor="hand2")
        self.fetch_advanced_summary.bind(
            "<Button-1>", lambda _event: self._manage_advanced_settings()
        )
        self.ui_mode_hint.configure(text=text, cursor="hand2")
        self.ui_mode_hint.bind(
            "<Button-1>", lambda _event: self._manage_advanced_settings()
        )

    def _manage_advanced_settings(self) -> None:
        if self.ui_mode_var.get() == "advanced":
            self.notebook.select(self.tab_settings)
            return
        if not messagebox.askyesno(
            "开启高级模式",
            "高级模式会显示更多选项和高风险操作。\n不会自动改变现有设置。\n\n继续？",
            parent=self.root,
            default=messagebox.NO,
        ):
            return
        self.ui_mode_var.set("advanced")
        self._apply_ui_mode()

    def _reset_hidden_behavior_settings(self) -> None:
        self.follow_load_mode_var.set("自动")
        self.identity_strategy_var.set("补全用户ID + 贴吧ID")
        self.identity_timing_var.set(False)
        self.identity_concurrency_var.set("3")
        self.batch_concurrency_var.set("1")
        self.proxy_mode_var.set("跟随系统")
        self.proxy_url_var.set("")
        self.lookup_fill_var.set(False)
        self.lookup_fill_all_var.set(False)
        self.comment_concurrency_var.set("1")
        self.delay_var.set("0.8")
        self.comment_depth_var.set("全部")
        self.settings.update(
            {
                "follow_load_mode": "auto",
                "identity_strategy": "full",
                "identity_timing": False,
                "identity_concurrency": 3,
                "batch_concurrency": 1,
                "proxy_mode": "system",
                "proxy_url": "",
                "lookup_fill_threads": False,
                "lookup_fill_all": False,
            }
        )
        save_settings(self.settings)

    def _capture_ui_pack_state(self) -> None:
        if hasattr(self, "_ui_pack_widgets"):
            return
        widgets = [
            self.follow_load_row,
            self.proxy_row,
            self.batch_row,
            self.lookup_fill_box,
            self.cache_box,
            self.log_box,
            self.encrypt_check,
            self.old_result_buttons,
        ]
        options: Dict[Any, Dict[str, Any]] = {}
        next_widgets: Dict[Any, Any] = {}
        for widget in widgets:
            info = dict(widget.pack_info())
            info.pop("in", None)
            info.pop("in_", None)
            options[widget] = info
            siblings = widget.master.pack_slaves()
            try:
                index = siblings.index(widget)
            except ValueError:
                continue
            next_widgets[widget] = siblings[index + 1] if index + 1 < len(siblings) else None
        self._ui_pack_widgets = widgets
        self._ui_pack_options = options
        self._ui_pack_next = next_widgets

    def _restore_ui_pack_widgets(self) -> None:
        self._capture_ui_pack_state()
        for widget in reversed(self._ui_pack_widgets):
            info = dict(self._ui_pack_options.get(widget) or {})
            if not info:
                continue
            next_widget = self._ui_pack_next.get(widget)
            while next_widget is not None and not next_widget.winfo_manager():
                next_widget = self._ui_pack_next.get(next_widget)
            if next_widget is not None and next_widget.winfo_manager():
                widget.pack(**info, before=next_widget)
            else:
                widget.pack(**info)

    def _apply_ui_mode(self, persist: bool = True) -> None:
        simple = self.ui_mode_var.get() != "advanced"
        self.settings["ui_mode"] = "simple" if simple else "advanced"
        if persist:
            save_settings(self.settings)
        self._capture_ui_pack_state()
        pack_widgets = getattr(self, "_ui_pack_widgets", [])

        if simple:
            self.identity_row.grid_remove()
            self.delay_row.grid_remove()
            for widget in pack_widgets:
                widget.pack_forget()
        else:
            self.identity_row.grid()
            self.delay_row.grid()
            self._restore_ui_pack_widgets()
        self._refresh_advanced_summary()

    def _on_ui_mode_selected(self) -> None:
        target = self.ui_mode_var.get()
        current = "advanced" if self.settings.get("ui_mode") == "advanced" else "simple"
        if target == current:
            return
        if target == "advanced":
            if not messagebox.askyesno(
                "开启高级模式",
                "高级模式会显示更多选项和高风险操作。\n不会自动改变现有设置。\n\n继续？",
                parent=self.root,
                default=messagebox.NO,
            ):
                self.ui_mode_var.set(current)
                return
            self._apply_ui_mode()
            return

        names = self._hidden_advanced_settings()
        if names:
            detail = "\n".join(f"- {name}" for name in names)
            choice = self._ask_choice(
                "切回简洁模式",
                "以下高级设置仍在生效：\n\n"
                + detail
                + "\n\n切回简洁模式不会自动关闭它们。",
                [
                    ("保留设置并切换", "keep"),
                    ("关闭行为类设置并切换", "close"),
                    ("取消", "cancel"),
                ],
                default=0,
            )
            if not choice or choice == "cancel":
                self.ui_mode_var.set(current)
                return
            if choice == "close":
                self._reset_hidden_behavior_settings()
        self._apply_ui_mode()


    def _open_guide_from_banner(self, banner) -> None:
        banner.destroy()
        self._open_usage_window()

    def _maybe_onboarding(self) -> None:
        if self.settings.get("onboarding_seen"):
            return
        self.settings["onboarding_seen"] = True
        save_settings(self.settings)
        banner = tk.Frame(self.root, bg="#fff7e6", padx=14, pady=7)
        banner.pack(fill="x", before=self.notebook)
        tk.Label(
            banner,
            text="首次使用：默认匿名；红绿灯/本地标记只存本机；关注增强默认关闭，可在设置里开启。",
            bg="#fff7e6",
            fg="#7c4a03",
            font=("Microsoft YaHei UI", 9),
            anchor="w",
        ).pack(side="left", fill="x", expand=True)
        ttk.Button(
            banner,
            text="查看说明",
            command=lambda: self._open_guide_from_banner(banner),
            width=8,
        ).pack(side="right", padx=(6, 0))
        ttk.Button(banner, text="关闭", command=banner.destroy, width=6).pack(side="right", padx=(6, 0))
        self._theme_widget_tree(banner)

    def _apply_settings_to_ui(self, reset_mode: bool = True) -> None:
        settings = self.settings
        if hasattr(self, "risk_var"):
            self.risk_var.set(bool(settings.get("risk_enabled", True)))
        if hasattr(self, "follow_enhance_var"):
            self.follow_enhance_var.set(bool(settings.get("follow_enhance", False)))
        if hasattr(self, "follow_load_mode_var"):
            self.follow_load_mode_var.set(
                "手动" if settings.get("follow_load_mode") == "manual" else "自动"
            )
        if hasattr(self, "webview_remember_var"):
            self.webview_remember_var.set(bool(settings.get("webview_remember_login", False)))
        if hasattr(self, "_refresh_webview_status"):
            self._refresh_webview_status()
        if hasattr(self, "auto_refresh_var"):
            self.auto_refresh_var.set(bool(settings.get("auto_refresh_old", False)))
        if hasattr(self, "identity_strategy_var"):
            self.identity_strategy_var.set(
                {
                    "full": "补全用户ID + 贴吧ID",
                    "ask": "每次询问",
                    "skip": "跳过补全",
                    "user_id": "只补用户ID",
                    "detail": "详情再补",
                }.get(str(settings.get("identity_strategy", "full")), "补全用户ID + 贴吧ID")
            )
            self.identity_timing_var.set(bool(settings.get("identity_timing", True)))
            self.identity_concurrency_var.set(str(settings.get("identity_concurrency", 3)))
        if hasattr(self, "lookup_fill_var"):
            self.lookup_fill_var.set(bool(settings.get("lookup_fill_threads", False)))
            self.lookup_fill_pages_var.set(
                "完整主题" if settings.get("lookup_fill_pages") == "full" else "第 1 页"
            )
            self.lookup_fill_limit_var.set(str(settings.get("lookup_fill_limit", 50)))
            self.lookup_fill_all_var.set(bool(settings.get("lookup_fill_all", False)))
            self.lookup_fill_comments_var.set(bool(settings.get("lookup_fill_comments", False)))
            all_topics = bool(settings.get("lookup_fill_all", False))
            try:
                self.lookup_fill_pages_box.configure(state="disabled" if all_topics else "readonly")
                self.lookup_fill_limit_box.configure(state="disabled" if all_topics else "normal")
            except tk.TclError:
                pass
        if hasattr(self, "love_level_var"):
            self.love_level_var.set(bool(settings.get("love_level_enabled", False)))
        if hasattr(self, "removed_action_var"):
            self.removed_action_var.set(
                {"keep": "保留旧结果", "remove": "移除旧结果", "ask": "每次询问"}.get(
                    settings.get("removed_thread_action", "keep"), "保留旧结果"
                )
            )
        if hasattr(self, "proxy_mode_var"):
            self.proxy_mode_var.set(
                {"system": "跟随系统", "none": "不使用", "custom": "自定义"}.get(
                    settings.get("proxy_mode", "system"), "跟随系统"
                )
            )
        if hasattr(self, "proxy_url_var"):
            self.proxy_url_var.set(str(settings.get("proxy_url") or ""))
        if hasattr(self, "batch_concurrency_var"):
            self.batch_concurrency_var.set(str(settings.get("batch_concurrency", 1)))
        if hasattr(self, "settings_start_var"):
            self.settings_start_var.set(
                "上次分类" if settings.get("settings_start") == "last" else "常规"
            )
        if hasattr(self, "theme_mode_var"):
            self.theme_mode_var.set(
                {"light": "浅色", "dark": "深色", "system": "跟随系统"}.get(
                    settings.get("theme_mode", "light"), "浅色"
                )
            )
        if hasattr(self, "accent_var"):
            self.accent_var.set(
                {
                    "blue": "蓝色",
                    "green": "绿色",
                    "purple": "紫色",
                    "orange": "橘黄",
                    "pink": "粉色",
                    "cyan": "青色",
                    "red": "红色",
                }.get(settings.get("accent", "blue"), "蓝色")
            )
        if hasattr(self, "table_density_var"):
            self.table_density_var.set(
                "紧凑" if settings.get("table_density") == "compact" else "舒适"
            )
        if hasattr(self, "save_result_var"):
            self.save_result_var.set(bool(settings.get("save_result", True)))
        if hasattr(self, "encrypt_var"):
            self.encrypt_var.set(bool(settings.get("encrypt_result", False)))
        if hasattr(self, "log_enabled_var"):
            self.log_enabled_var.set(bool(settings.get("log_enabled", False)))
        if hasattr(self, "restore_mode_var"):
            self.restore_mode_var.set("询问" if settings.get("restore_mode") == "ask" else "自动")
        if hasattr(self, "cache_limit_var"):
            limit = int(settings.get("post_cache_limit_mb", DEFAULT_LIMIT_MB))
            self.cache_limit_var.set(str(limit))
            post_cache().set_limit_mb(limit)
        for key, var in getattr(self, "setting_vars", {}).items():
            var.set(bool(settings.get(key, True)))
        if reset_mode:
            self._set_mode("guest")
        self._apply_risk_ui()
        self._apply_theme()
        self._apply_table_density()
        self._apply_proxy()

    def _restore_defaults(self) -> None:
        if not messagebox.askyesno(
            "恢复默认设置",
            "只重置设置，保留结果、本地标记和 BDUSS。继续？",
            parent=self.root,
            default=messagebox.NO,
        ):
            return
        keep_onboarding = bool(self.settings.get("onboarding_seen"))
        keep_geometry = self.settings.get("window_geometry", "")
        keep_webview_remember = bool(self.settings.get("webview_remember_login", False))
        self.settings = load_settings()
        self.settings["onboarding_seen"] = keep_onboarding
        self.settings["window_geometry"] = keep_geometry
        self.settings["webview_remember_login"] = keep_webview_remember
        save_settings(self.settings)
        self._apply_settings_to_ui()
        self._append_log("设置已恢复默认（结果、标记、BDUSS 保留）。", "muted")



    def _apply_proxy(self) -> None:
        var = getattr(self, "proxy_mode_var", None)
        mode = {"跟随系统": "system", "不使用": "none", "自定义": "custom"}.get(
            var.get() if var is not None else "跟随系统",
            "system",
        )
        if mode == "custom":
            Fetcher.DEFAULT_PROXY = self.proxy_url_var.get().strip() or None
        elif mode == "none":
            Fetcher.DEFAULT_PROXY = ""
        else:
            Fetcher.DEFAULT_PROXY = None

    def _on_proxy_changed(self) -> None:
        mode = {"跟随系统": "system", "不使用": "none", "自定义": "custom"}.get(
            self.proxy_mode_var.get(), "system"
        )
        self.settings["proxy_mode"] = mode
        self.settings["proxy_url"] = self.proxy_url_var.get().strip()
        save_settings(self.settings)
        self._apply_proxy()
        self._update_risk_summary()

    def _on_batch_concurrency_changed(self) -> None:
        try:
            value = int(self.batch_concurrency_var.get())
        except ValueError:
            value = 1
        self.settings["batch_concurrency"] = max(1, min(3, value))
        save_settings(self.settings)



    def _update_risk_summary(self) -> None:
        if not hasattr(self, "risk_summary_label"):
            return
        items = self._high_risk_items()
        self.risk_summary_label.configure(
            text=("风险：" + " · ".join(items)) if items else "风险：无高风险设置"
        )





    def _on_fetch_settings_changed(self) -> None:
        self.settings["auto_refresh_old"] = bool(self.auto_refresh_var.get())
        value = {"保留旧结果": "keep", "移除旧结果": "remove", "每次询问": "ask"}.get(
            self.removed_action_var.get(), "keep"
        )
        self.settings["removed_thread_action"] = value
        save_settings(self.settings)

    def _on_settings_changed(self) -> None:
        for key, var in getattr(self, "setting_vars", {}).items():
            self.settings[key] = bool(var.get())
        save_settings(self.settings)
        selection = self.people_tree.selection() if hasattr(self, "people_tree") else ()
        if selection:
            try:
                user = self.people_users[int(selection[0])]
            except (IndexError, ValueError):
                user = None
            if user:
                cached = getattr(self, "_profile_cache", {}).get(user["uid"] or f"name:{user['name']}")
                if cached:
                    self._render_profile(user, cached)
        if getattr(self, "_last_lookup_profile", None):
            self._render_lookup(self._last_lookup_profile)
        self.general_hint.configure(text="已保存")

    def _backup_config(self) -> None:
        path = filedialog.asksaveasfilename(
            title="备份设置",
            defaultextension=".json",
            filetypes=[("JSON 配置", "*.json"), ("ZIP 压缩包", "*.zip")],
            initialfile="贴吧统计设置.json",
            parent=self.root,
        )
        if not path:
            return
        try:
            saved = write_backup(path, self.settings, self.blacklist)
        except Exception as error:  # noqa: BLE001
            messagebox.showerror("备份失败", friendly_error(error), parent=self.root)
            return
        self._append_log(f"设置已备份：{saved}", "ok")
        messagebox.showinfo(
            "备份完成",
            "已保存到：" + chr(10) + saved,
            parent=self.root,
        )

    def _restore_config(self) -> None:
        path = filedialog.askopenfilename(
            title="恢复设置",
            filetypes=[("设置备份", "*.json *.zip"), ("所有文件", "*.*")],
            parent=self.root,
        )
        if not path:
            return
        try:
            data = read_backup(path)
        except ValueError as error:
            messagebox.showerror("无法恢复", friendly_error(error), parent=self.root)
            return
        if not messagebox.askyesno(
            "恢复设置",
            backup_preview(data) + chr(10) + chr(10) + "会覆盖当前设置和本地标记，继续吗？",
            parent=self.root,
        ):
            return
        try:
            import tempfile

            write_backup(
                os.path.join(tempfile.gettempdir(), "tieba-config-before-restore.json"),
                self.settings,
                self.blacklist,
            )
        except Exception:  # noqa: BLE001
            pass
        self.settings.update(data["settings"])
        save_settings(self.settings)
        self.blacklist = data["blacklist"]
        save_blacklist(self.blacklist)
        self._on_mark_rules_changed()
        follow = bool(self.settings.get("follow_enhance", False))
        if follow and not parse_cookie_input(self.cookie_var.get()).get("BDUSS"):
            follow = False
            self.settings["follow_enhance"] = False
        if hasattr(self, "follow_enhance_var"):
            self.follow_enhance_var.set(follow)
        self.mode_var.set("login" if follow else "guest")
        self._apply_settings_to_ui(reset_mode=False)
        for key, editor in getattr(self, "blacklist_editors", {}).items():
            editor.delete("1.0", "end")
            editor.insert("1.0", chr(10).join(self.blacklist.get(key, [])))
        self._apply_mode_ui()
        self._refresh_people()
        if getattr(self, "_last_lookup_profile", None):
            self._render_lookup(self._last_lookup_profile)
        self._append_log("配置已恢复。", "ok")
        self.status_var.set("配置已恢复")

    def _on_mark_rules_changed(self) -> None:
        self._marks_stale = True
        if hasattr(self, "people_hint"):
            self.people_hint.configure(
                text="标记规则已变更，可点「重新标记（本地）」；不会自动联网"
            )

    def _save_blacklist_from_tab(self) -> None:
        data = {
            key: [line.strip() for line in editor.get("1.0", "end").splitlines() if line.strip()]
            for key, editor in self.blacklist_editors.items()
        }
        if save_blacklist(data):
            self.blacklist = data
            self.settings_hint.configure(text="已保存")
            self._on_mark_rules_changed()
            self._append_log("本地标记已保存。", "ok")
        else:
            self.settings_hint.configure(text="保存失败")

    def _blacklist_from_editors(self) -> dict:
        return {
            key: [line.strip() for line in editor.get("1.0", "end").splitlines() if line.strip()]
            for key, editor in self.blacklist_editors.items()
        }

    def _export_blacklist(self) -> None:
        path = filedialog.asksaveasfilename(
            parent=self.root,
            title="导出本地标记",
            defaultextension=".xlsx",
            initialfile="贴吧本地标记.xlsx",
            filetypes=[("Excel 表格", "*.xlsx")],
        )
        if not path:
            return
        if export_blacklist_xlsx(path, self._blacklist_from_editors()):
            self.settings_hint.configure(text=f"已导出：{os.path.basename(path)}")
            self._append_log(f"本地标记已导出：{path}", "ok")
        else:
            messagebox.showerror("导出失败", "无法写入文件，请换个位置重试。", parent=self.root)

    def _import_blacklist(self) -> None:
        path = filedialog.askopenfilename(
            parent=self.root,
            title="导入本地标记",
            filetypes=[("Excel 表格", "*.xlsx"), ("所有文件", "*.*")],
        )
        if not path:
            return
        data = import_blacklist_xlsx(path)
        total = sum(len(value) for value in data.values())
        if not total:
            messagebox.showwarning(
                "没有读到数据",
                "格式应为两列：\n第一列 类别（安全 / 可疑 / 警告）\n第二列 贴吧名\n\n"
                "可以用「导出 Excel」生成模板后再填。",
                parent=self.root,
            )
            return
        existing_total = sum(len(value) for value in self.blacklist.values())
        merged = merge_blacklists(self.blacklist, data)
        merged_total = sum(len(value) for value in merged.values())
        imported_total = total
        choice = self._ask_choice(
            "导入方式",
            f"现有 {existing_total} 条；导入 {imported_total} 条。"
            f"合并后 {merged_total} 条；覆盖后 {imported_total} 条。",
            [("覆盖（默认）", False), ("合并（保留现有）", True), ("取消", None)],
            default=0,
        )
        if choice is None:
            return
        if choice:
            data = merge_blacklists(self.blacklist, data)
        if save_blacklist(data):
            self.blacklist = data
            for key, editor in self.blacklist_editors.items():
                editor.delete("1.0", "end")
                editor.insert("1.0", "\n".join(data[key]))
            self.settings_hint.configure(text="已导入")
            self._on_mark_rules_changed()
            self._append_log(f"本地标记已导入：{path}", "ok")
        else:
            messagebox.showerror("导入失败", "保存失败，请检查配置目录权限。", parent=self.root)

    def _open_manage_results(self) -> None:
        if not self.last_threads:
            messagebox.showinfo("还没有数据", "先抓取帖子，再管理帖子结果。", parent=self.root)
            return
        win = tk.Toplevel(self.root)
        win.title("管理帖子结果")
        win.geometry("620x420")
        win.transient(self.root)
        win.grab_set()
        win.configure(bg="#eef1f5")
        tk.Label(
            win,
            text="勾选要删除的帖子结果；只删结果，不删原帖、Excel 或本地标记。",
            bg="#eef1f5",
            fg="#475569",
            font=("Microsoft YaHei UI", 9),
            anchor="w",
        ).pack(fill="x", padx=12, pady=(10, 4))
        body = tk.Frame(win, bg="#ffffff")
        body.pack(fill="both", expand=True, padx=12)
        canvas = tk.Canvas(body, bg="#ffffff", highlightthickness=0)
        scrollbar = ttk.Scrollbar(body, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        inner = tk.Frame(canvas, bg="#ffffff")
        canvas.create_window((0, 0), window=inner, anchor="nw")
        inner.bind("<Configure>", lambda _event: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda event: canvas.itemconfigure(1, width=event.width))
        variables = []
        for index, thread in enumerate(self.last_threads, start=1):
            var = tk.BooleanVar(value=False)
            variables.append(var)
            title = display_text(thread.meta.title or thread.meta.tid)
            count = len(thread.user_stats())
            ttk.Checkbutton(
                inner,
                text=f"{index}. {title}    {count} 人",
                variable=var,
            ).pack(anchor="w", padx=6, pady=3)
        footer = tk.Frame(win, bg="#eef1f5")
        footer.pack(fill="x", padx=12, pady=10)
        cancel_btn = ttk.Button(footer, text="取消", command=win.destroy, width=8)
        cancel_btn.pack(side="right")
        def select_all(value: bool) -> None:
            for var in variables:
                var.set(value)

        ttk.Button(
            footer,
            text="全选",
            command=lambda: select_all(True),
            width=8,
        ).pack(side="left")
        ttk.Button(
            footer,
            text="取消全选",
            command=lambda: select_all(False),
            width=10,
        ).pack(side="left", padx=(6, 0))

        def delete_selected() -> None:
            selected = [index for index, var in enumerate(variables) if var.get()]
            if not selected:
                messagebox.showinfo("没有选择", "先勾选要删除的帖子结果。", parent=win)
                return
            if not messagebox.askyesno(
                "删除帖子结果",
                f"删除 {len(selected)} 个帖子结果？删除后不可撤销，不影响原帖、Excel 或本地标记。",
                parent=win,
                default=messagebox.NO,
            ):
                return
            remaining = [thread for index, thread in enumerate(self.last_threads) if index not in selected]
            try:
                if self.settings.get("save_result", True):
                    if remaining:
                        save_threads(
                            remaining,
                            VERSION,
                            encrypted=bool(self.settings.get("encrypt_result", False)),
                        )
                    else:
                        if not clear_stored_result():
                            raise OSError("无法删除当前结果文件")
            except Exception as error:  # noqa: BLE001
                messagebox.showerror("保存失败", f"删除未执行：{error}", parent=win)
                return
            active = self.grid_thread_var.get()
            active_tid = None
            if not active.startswith("全部"):
                for old_thread in self.last_threads:
                    if active.endswith(display_text(old_thread.meta.title or old_thread.meta.tid)):
                        active_tid = old_thread.meta.tid
                        break
            self.last_threads = remaining
            if not remaining:
                self.last_output = None
                self.open_btn.configure(state="disabled")
            elif active.startswith("全部"):
                self._update_grid_selector()
            elif active_tid in {thread.meta.tid for thread in remaining}:
                self._update_grid_selector()
                for index, thread in enumerate(remaining, start=1):
                    if thread.meta.tid == active_tid:
                        self.grid_thread_var.set(
                            f"{index}. {display_text(thread.meta.title or thread.meta.tid)}"
                        )
                        break
            else:
                self.grid_thread_var.set("")
                self._update_grid_selector()
            self._refresh_people()
            self._refresh_grid()
            self._append_log(f"已删除 {len(selected)} 个帖子结果。", "muted")
            win.destroy()

        ttk.Button(footer, text="删除选中", command=delete_selected, width=10).pack(
            side="right", padx=(0, 6)
        )
        win.protocol("WM_DELETE_WINDOW", win.destroy)
        win.update_idletasks()
        width = max(620, win.winfo_width())
        height = max(420, win.winfo_height())
        x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - width) // 2)
        y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - height) // 2)
        win.geometry(f"{width}x{height}+{x}+{y}")
        cancel_btn.focus_set()


    def _apply_table_density(self) -> None:
        compact = self.settings.get("table_density") == "compact"
        style = ttk.Style(self.root)
        style.configure("Treeview", rowheight=25 if compact else 31)

    def _on_table_density_changed(self) -> None:
        self.settings["table_density"] = "compact" if self.table_density_var.get() == "紧凑" else "comfortable"
        save_settings(self.settings)
        self._apply_table_density()



    def _aggregate_threads(self, threads):
        buckets: Dict[str, Dict[str, Any]] = {}
        for thread in threads:
            for item in thread.user_stats():
                key = (
                    f"uid:{item['uid']}"
                    if item.get("uid")
                    else (
                        f"portrait:{str(item.get('portrait') or '').split('?')[0]}"
                        if item.get("portrait")
                        else (f"name:{item['username']}" if item.get("username") else f"anon:{item['name']}")
                    )
                )
                bucket = buckets.setdefault(
                    key,
                    {
                        "name": item["name"],
                        "username": item["username"],
                        "uid": item["uid"],
                        "tieba_uid": item.get("tieba_uid") or "",
                        "portrait": item["portrait"],
                        "level": item["level"],
                        "badge": item["badge"],
                        "floor": 0,
                        "comment": 0,
                        "total": 0,
                        "threads": 0,
                        "aliases": set(),
                    },
                )
                bucket["floor"] += item["floor_count"]
                bucket["comment"] += item["comment_count"]
                bucket["total"] += item["total"]
                bucket["threads"] += 1
                if item.get("name") and item["name"] != bucket["name"]:
                    bucket["aliases"].add(bucket["name"])
                    bucket["name"] = item["name"]
                bucket["aliases"].update(item.get("aliases") or [])
                if not bucket["username"] and item["username"]:
                    bucket["username"] = item["username"]
                if not bucket["portrait"] and item["portrait"]:
                    bucket["portrait"] = item["portrait"]
                if not bucket["tieba_uid"] and item.get("tieba_uid"):
                    bucket["tieba_uid"] = item["tieba_uid"]
                if not bucket["level"] and item["level"]:
                    bucket["level"] = item["level"]
                if not bucket["badge"] and item["badge"]:
                    bucket["badge"] = item["badge"]
        users = list(buckets.values())
        for user in users:
            user["aliases"] = sorted(a for a in user.get("aliases", ()) if a and a != user["name"])
        users.sort(key=lambda item: (-item["total"], -item["floor"], item["name"]))
        return users

        # 参与者对比已确认砍掉，不再创建入口。

    def _add_compare_tab(self) -> None:
        frame = tk.Frame(self.grid_nb)
        self.grid_nb.add(frame, text="  参与者对比  ")
        top = tk.Frame(frame, bg="#ffffff", padx=8, pady=6)
        top.pack(fill="x")
        tk.Label(top, text="勾选帖子：", bg="#ffffff", fg="#475569", font=("Microsoft YaHei UI", 9)).pack(
            side="left"
        )
        self.compare_vars = []
        for index, thread in enumerate(self.last_threads):
            var = tk.BooleanVar(value=index < 2)
            label = display_text(thread.meta.title or thread.meta.tid)[:16]
            tk.Checkbutton(
                top, text=label, variable=var, bg="#ffffff", anchor="w", font=("Microsoft YaHei UI", 8)
            ).pack(side="left", padx=(2, 0))
            self.compare_vars.append(var)
        self.compare_mode_var = tk.StringVar(value="交集")
        ttk.Combobox(top, values=MODES, textvariable=self.compare_mode_var, width=8, state="readonly").pack(
            side="left", padx=(8, 0)
        )
        ttk.Button(top, text="生成对比", command=self._generate_compare, width=10).pack(
            side="left", padx=(6, 0)
        )
        ttk.Button(top, text="导出结果", command=self._export_compare, width=10).pack(
            side="left", padx=(4, 0)
        )
        self.compare_tree = ttk.Treeview(frame, columns=(), show="headings")
        self.compare_tree.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self.compare_result: Optional[tuple] = None

    def _generate_compare(self) -> None:
        selected = [index for index, var in enumerate(self.compare_vars) if var.get()]
        if len(selected) < 2:
            self.grid_hint.configure(text="至少勾选 2 个帖子才能对比")
            return
        try:
            result = compare_threads(self.last_threads, selected, self.compare_mode_var.get())
        except ValueError as error:
            self.grid_hint.configure(text=friendly_error(error))
            return
        headers = ["昵称", "用户名", "用户ID"] + result["threads"] + ["合计", "类型"]
        columns = [str(index) for index in range(len(headers))]
        self.compare_tree.configure(columns=columns)
        for index, header in enumerate(headers):
            self.compare_tree.heading(columns[index], text=header)
            self.compare_tree.column(
                columns[index],
                width=110 if header in ("昵称", "用户名") else 70,
                anchor="w" if header in ("昵称", "用户名") else "center",
            )
        for item in self.compare_tree.get_children():
            self.compare_tree.delete(item)
        rows = []
        for row in result["rows"]:
            values = [
                display_text(row["name"]),
                row["username"],
                row["uid"],
            ] + row["counts"] + [row["total"], row["type"]]
            rows.append([str(value) for value in values])
            self.compare_tree.insert("", "end", values=values)
        self.compare_result = (headers, rows)

    def _compare_label(self, index: int) -> str:
        selected = [i for i, var in enumerate(getattr(self, "compare_vars", [])) if var.get()]
        if index < len(selected):
            thread = self.last_threads[selected[index]]
            return display_text(thread.meta.title or thread.meta.tid)
        return ""

    def _export_compare(self) -> None:
        if not self.compare_result:
            self.grid_hint.configure(text="先生成对比结果")
            return
        from openpyxl import Workbook

        path = filedialog.asksaveasfilename(
            title="导出对比结果",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile="参与者对比.xlsx",
            parent=self.root,
        )
        if not path:
            return
        headers, rows = self.compare_result
        book = Workbook()
        sheet = book.active
        sheet.title = "参与者对比"
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        notes = book.create_sheet("导出说明")
        notes.append(["项目", "内容"])
        notes.append(["对比模式", self.compare_mode_var.get()])
        notes.append(["勾选帖子数", str(len(rows))])
        notes.append(["A", self._compare_label(0)])
        notes.append(["B", self._compare_label(1)])
        notes.append(["软件版本", VERSION])
        notes.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        notes.append(["结果行数", str(len(rows))])
        selected = [i for i, var in enumerate(getattr(self, "compare_vars", [])) if var.get()]
        warnings = sum(len(self.last_threads[i].warnings) for i in selected)
        notes.append(["数据警告", str(warnings) if warnings else "无"])
        book.save(path)
        self._append_log(f"对比结果已导出：{path}", "ok")

    _SORT_ONLY_HEADERS = {"日期", "等级", "楼层", "楼中楼", "合计", "所在楼层", "#"}
    _TEXT_FILTER_HEADERS = {
        "类型", "标记", "昵称", "回复者", "用户名", "贴吧ID", "用户ID", "头衔",
        "内容", "帖子", "贴吧", "标题", "来源",
    }

    def _anchor_point(self, anchor: Optional[tuple]) -> tuple:
        if anchor:
            return int(anchor[0]), int(anchor[1])
        return self.root.winfo_pointerx(), self.root.winfo_pointery()

    def _place_filter_window(self, win, anchor: Optional[tuple], fallback_width: int = 320) -> None:
        x, y = self._anchor_point(anchor)
        win.update_idletasks()
        width = max(fallback_width, win.winfo_width())
        height = max(120, win.winfo_height())
        screen_w = win.winfo_screenwidth()
        screen_h = win.winfo_screenheight()
        left = x
        top = y + 6
        if left + width > screen_w - 8:
            left = max(8, screen_w - width - 8)
        if top + height > screen_h - 8:
            top = max(8, y - height - 6)
        win.geometry(f"{width}x{height}+{left}+{top}")

    def _filter_text_dialog(
        self, title: str, values: List[str], selected, anchor: Optional[tuple] = None
    ) -> Optional[set]:
        win = tk.Toplevel(self.root)
        win.title("筛选：" + title)
        win.configure(bg="#ffffff")
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        frame = tk.Frame(win, bg="#ffffff", padx=12, pady=10)
        frame.pack(fill="both", expand=True)
        tk.Label(frame, text="筛选：" + title, bg="#ffffff", fg="#1f2937",
                 font=("Microsoft YaHei UI", 10, "bold")).pack(anchor="w")
        search = tk.StringVar()
        tk.Entry(frame, textvariable=search, width=38).pack(fill="x", pady=(6, 6))
        tk.Label(
            frame,
            text="单击切换选中，可多选",
            bg="#ffffff",
            fg="#8a94a6",
            font=("Microsoft YaHei UI", 8),
        ).pack(anchor="w", pady=(0, 3))
        listbox = tk.Listbox(frame, selectmode="multiple", width=42, height=14, exportselection=False)
        listbox.pack(fill="both", expand=True)
        for value in values:
            listbox.insert("end", value)
        chosen = set(str(v) for v in selected) if selected else set(values)
        for index, value in enumerate(values):
            if value in chosen:
                listbox.selection_set(index)

        def take_selection() -> set:
            return {str(listbox.get(i)) for i in listbox.curselection()}

        select_all_btn: Optional[ttk.Button] = None

        def update_select_all_label() -> None:
            if select_all_btn is None:
                return
            size = listbox.size()
            selected = len(listbox.curselection())
            select_all_btn.configure(text="取消全选" if size and selected >= size else "全选")

        def toggle_all() -> None:
            size = listbox.size()
            if not size:
                return
            if len(listbox.curselection()) >= size:
                listbox.selection_clear(0, "end")
            else:
                listbox.selection_set(0, "end")
            update_select_all_label()

        def refresh(*_args) -> None:
            nonlocal chosen
            chosen = take_selection()
            query = search.get().strip().lower()
            listbox.delete(0, "end")
            for value in values:
                if not query or query in value.lower():
                    listbox.insert("end", value)
            for index, value in enumerate(listbox.get(0, "end")):
                if value in chosen:
                    listbox.selection_set(index)
            update_select_all_label()

        search.trace_add("write", refresh)
        listbox.bind("<<ListboxSelect>>", lambda _event: update_select_all_label())
        result: Dict[str, Any] = {"value": None}
        buttons = tk.Frame(frame, bg="#ffffff")
        buttons.pack(fill="x", pady=(8, 0))

        def apply() -> None:
            result["value"] = take_selection()
            win.destroy()

        def clear_filter() -> None:
            result["value"] = set()
            win.destroy()

        select_all_btn = ttk.Button(buttons, text="全选", command=toggle_all, width=8)
        select_all_btn.pack(side="left")
        ttk.Button(buttons, text="清除筛选", command=clear_filter, width=10).pack(
            side="left", padx=(6, 0)
        )
        ttk.Button(buttons, text="取消", command=lambda: win.destroy(), width=8).pack(side="right")
        ttk.Button(buttons, text="应用", command=apply, width=8).pack(side="right", padx=(0, 6))
        update_select_all_label()
        self._place_filter_window(win, anchor, 380)
        self.root.wait_window(win)
        return result["value"]

    def _filter_date_dialog(
        self, title: str, selected, anchor: Optional[tuple] = None
    ) -> Optional[tuple]:
        win = tk.Toplevel(self.root)
        win.title("筛选：" + title)
        win.configure(bg="#ffffff")
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        frame = tk.Frame(win, bg="#ffffff", padx=14, pady=12)
        frame.pack(fill="both", expand=True)
        tk.Label(frame, text="日期范围（留空表示不限）", bg="#ffffff", fg="#1f2937",
                 font=("Microsoft YaHei UI", 10, "bold")).pack(anchor="w")
        start_var = tk.StringVar(value=(selected or ("", ""))[0])
        end_var = tk.StringVar(value=(selected or ("", ""))[1])
        row = tk.Frame(frame, bg="#ffffff")
        row.pack(fill="x", pady=(8, 8))
        tk.Label(row, text="开始", bg="#ffffff", fg="#475569").pack(side="left")
        tk.Entry(row, textvariable=start_var, width=12).pack(side="left", padx=(4, 10))
        tk.Label(row, text="结束", bg="#ffffff", fg="#475569").pack(side="left")
        tk.Entry(row, textvariable=end_var, width=12).pack(side="left", padx=(4, 0))
        preset = tk.Frame(frame, bg="#ffffff")
        preset.pack(fill="x", pady=(0, 8))
        for label, days in (("全部", 0), ("近7天", 7), ("近30天", 30), ("近90天", 90)):
            def apply_preset(value=days) -> None:
                if not value:
                    start_var.set("")
                    end_var.set("")
                    return
                today = datetime.now()
                start_var.set((today - timedelta(days=value)).strftime("%Y-%m-%d"))
                end_var.set(today.strftime("%Y-%m-%d"))
            ttk.Button(preset, text=label, command=apply_preset, width=7).pack(side="left", padx=(0, 4))
        result: Dict[str, Any] = {"value": None}
        buttons = tk.Frame(frame, bg="#ffffff")
        buttons.pack(fill="x")

        def apply() -> None:
            for value in (start_var.get().strip(), end_var.get().strip()):
                if value and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                    messagebox.showerror("日期格式", "日期请用 YYYY-MM-DD。", parent=win)
                    return
            result["value"] = (start_var.get().strip(), end_var.get().strip())
            win.destroy()

        def clear() -> None:
            result["value"] = ("", "")
            win.destroy()

        ttk.Button(buttons, text="取消", command=lambda: win.destroy(), width=8).pack(side="right")
        ttk.Button(buttons, text="应用", command=apply, width=8).pack(side="right", padx=(0, 6))
        ttk.Button(buttons, text="清空", command=clear, width=8).pack(side="left")
        self._place_filter_window(win, anchor, 430)
        self.root.wait_window(win)
        return result["value"]

    def _loaded_filter_state(self, tree) -> dict:
        state = getattr(tree, "_loaded_filter_state", None)
        if state is None:
            state = {"all": list(tree.get_children()), "filters": {}}
            setattr(tree, "_loaded_filter_state", state)  # noqa: B010
        return state

    def _clear_loaded_tree(self, tree) -> None:
        state = getattr(tree, "_loaded_filter_state", None)
        for iid in (state["all"] if state else list(tree.get_children())):
            try:
                tree.delete(iid)
            except tk.TclError:
                pass
        if hasattr(tree, "_loaded_filter_state"):
            delattr(tree, "_loaded_filter_state")

    def _apply_loaded_filters(self, tree) -> None:
        state = self._loaded_filter_state(tree)
        filters = state["filters"]
        for iid in state["all"]:
            tree.detach(iid)
        for iid in state["all"]:
            keep = True
            for column, selected in filters.items():
                value = tree.set(iid, column)
                if isinstance(selected, tuple):
                    start_date, end_date = selected
                    day = value[:10]
                    if (start_date and (not day or day < start_date)) or (
                        end_date and (not day or day > end_date)
                    ):
                        keep = False
                        break
                elif selected and value not in selected:
                    keep = False
                    break
            if keep:
                tree.reattach(iid, "", "end")

    def _filter_loaded_tree(
        self, tree, column: str, title: str, date_column: bool = False, anchor: Optional[tuple] = None
    ) -> None:
        state = self._loaded_filter_state(tree)
        chosen: Any
        if date_column:
            current = state["filters"].get(column)
            chosen = self._filter_date_dialog(title, current, anchor)
            if chosen is None:
                return
            if chosen[0] or chosen[1]:
                state["filters"][column] = chosen
            else:
                state["filters"].pop(column, None)
        else:
            values = sorted({tree.set(iid, column) for iid in state["all"]})
            current = state["filters"].get(column)
            chosen = self._filter_text_dialog(title, values, current, anchor)
            if chosen is None:
                return
            if chosen:
                state["filters"][column] = chosen
            else:
                state["filters"].pop(column, None)
        self._apply_loaded_filters(tree)

    def _bind_loaded_filter(self, tree, mapping: Dict[str, Optional[tuple]]) -> None:
        def event_handler(event):
            if tree.identify_region(event.x, event.y) != "heading":
                return None
            column = tree.identify_column(event.x)
            keys = list(mapping)
            try:
                index = int(column[1:]) - 1
                key = keys[index]
                spec = mapping.get(key)
                if not spec:
                    return None
                title, is_date = spec
            except (ValueError, IndexError, KeyError):
                return None
            self._filter_loaded_tree(tree, key, title, is_date, (event.x_root, event.y_root))
            return "break"

        def left_handler(event):
            if tree.identify_region(event.x, event.y) != "heading":
                return None
            column = tree.identify_column(event.x)
            keys = list(mapping)
            try:
                index = int(column[1:]) - 1
                key = keys[index]
                spec = mapping.get(key)
                if not spec or spec[1]:
                    return None
            except (ValueError, IndexError, KeyError):
                return None
            self._filter_loaded_tree(tree, key, spec[0], False, (event.x_root, event.y_root))
            return "break"

        tree.bind("<Button-3>", event_handler, add="+")
        tree.bind("<Button-1>", left_handler, add="+")

    def _header_filterable(self, header: str) -> bool:
        return header in self._TEXT_FILTER_HEADERS or header == "日期"


    def _people_filter_active(self) -> bool:
        if self.people_keyword_var.get().strip():
            return True
        if self.people_level_var.get() != "全部":
            return True
        if self._risk_enabled() and self.people_mark_var.get() != "全部":
            return True
        if self.people_min_var.get().strip() or self.people_max_var.get().strip():
            return True
        if getattr(self, "_people_col_filters", {}):
            return True
        return False

    def _visible_people(self) -> list:
        users = filter_users(
            self.people_users,
            keyword=self.people_keyword_var.get(),
            field=self.people_field_var.get(),
            level=self.people_level_var.get(),
            mark=self.people_mark_var.get() if self._risk_enabled() else "全部",
            min_level=_optional_int(self.people_min_var.get()),
            max_level=_optional_int(self.people_max_var.get()),
        )
        for column, selected in getattr(self, "_people_col_filters", {}).items():
            if selected:
                users = [
                    user for user in users
                    if self._people_column_value(user, column) in selected
                ]
        return users

    def _filtered_threads(self) -> list:
        from dataclasses import replace

        selected = set()
        filtered = self._visible_people()
        for user in filtered:
            for value in (user.get("uid"), user.get("username"), user.get("name")):
                if value:
                    selected.add(str(value))
        result = []
        for thread in self.last_threads:
            floors = [
                floor
                for floor in thread.floors
                if str(floor.author_id) in selected
                or str(floor.author_name) in selected
                or str(floor.author) in selected
            ]
            comments = [
                comment
                for comment in thread.comments
                if str(comment.author_id) in selected
                or str(comment.author_name) in selected
                or str(comment.author) in selected
            ]
            result.append(replace(thread, floors=floors, comments=comments))
        return result

    def _filter_summary(self) -> str:
        if not self._people_filter_active():
            return "无"
        parts = []
        keyword = self.people_keyword_var.get().strip()
        if keyword:
            parts.append(f"关键词「{keyword}」")
        level = self.people_level_var.get()
        if level and level not in ("全部", "无等级"):
            parts.append(f"等级 {level}")
        elif level == "无等级":
            parts.append("等级 无")
        mark = self.people_mark_var.get()
        if mark and mark != "全部":
            parts.append(f"标记 {mark}")
        for column, selected in getattr(self, "_people_col_filters", {}).items():
            if selected:
                title = {"mark": "标记", "name": "昵称", "username": "用户名", "uid": "用户ID"}.get(column, column)
                parts.append(f"{title}筛选 {len(selected)} 项")
        return "；".join(parts) or "无"

    def _expand_large_threads(self, threads, limit: int = 500_000):
        result: list = []
        for thread in threads:
            size = len(thread.floors) + len(thread.comments)
            if size <= limit:
                result.append(thread)
                continue
            by_floor: Dict[int, List] = {}
            for comment in thread.comments:
                by_floor.setdefault(comment.floor_no, []).append(comment)
            groups: list = []
            part_floors: list = []
            part_comments: list = []
            rows = 0
            for floor in thread.floors:
                comments = by_floor.pop(floor.post_no, [])
                block = 1 + len(comments)
                if block > limit:
                    if part_floors or part_comments:
                        groups.append((part_floors, part_comments))
                        part_floors, part_comments, rows = [], [], 0
                    chunk_size = max(1, limit - 1)
                    for start in range(0, len(comments), chunk_size):
                        groups.append(([floor], comments[start : start + chunk_size]))
                    continue
                if part_floors and rows + block > limit:
                    groups.append((part_floors, part_comments))
                    part_floors, part_comments, rows = [], [], 0
                part_floors.append(floor)
                part_comments.extend(comments)
                rows += block
            for comments in by_floor.values():
                if not comments:
                    continue
                if part_floors or part_comments:
                    groups.append((part_floors, part_comments))
                    part_floors, part_comments, rows = [], [], 0
                chunk_size = max(1, limit)
                for start in range(0, len(comments), chunk_size):
                    groups.append(([], comments[start : start + chunk_size]))
            if part_floors or part_comments:
                groups.append((part_floors, part_comments))
            total = len(groups)
            for index, (floors, comments) in enumerate(groups, start=1):
                meta = replace(
                    thread.meta,
                    title=f"{thread.meta.title or thread.meta.tid}（分片 {index}/{total}）",
                )
                result.append(
                    replace(
                        thread,
                        meta=meta,
                        floors=list(floors),
                        comments=list(comments),
                    )
                )
        return result






    def _user_bar(self, user) -> str:
        """这个用户参与的是哪个吧（多帖子时取他出现过的第一个帖子的吧）。"""
        uid = user.get("uid") or ""
        name = user.get("name") or ""
        for thread in self.last_threads:
            for floor in thread.floors:
                if (uid and floor.author_id == uid) or (not uid and floor.author == name):
                    return thread.meta.forum
            for comment in thread.comments:
                if (uid and comment.author_id == uid) or (not uid and comment.author == name):
                    return thread.meta.forum
        return self.last_threads[0].meta.forum if self.last_threads else ""

    def _post_count(self, fallback: int = 50) -> int:
        try:
            count = int(self.post_count_var.get())
        except (AttributeError, ValueError):
            count = fallback
        return max(5, min(1000, count))

    def _date_params(self):
        if self.post_date_mode_var.get() == "指定日期":
            return self._group_date(self.post_start_parts), self._group_date(self.post_end_parts)
        return "", ""

    def _set_people_controls_enabled(self, enabled: bool) -> None:
        readonly = "readonly" if enabled else "disabled"
        disabled = "normal" if enabled else "disabled"
        for name, state in (
            ("post_count_box", "normal" if enabled else "disabled"),
            ("post_date_mode_box", readonly),
            ("post_prefer_box", readonly),
            ("people_range_box", readonly),
            ("people_field_box", readonly),
            ("people_keyword_entry", "normal" if enabled else "disabled"),
            ("people_level_box", readonly),
            ("people_min_entry", "normal" if enabled else "disabled"),
            ("people_max_entry", "normal" if enabled else "disabled"),
            ("people_reset_btn", disabled),
        ):
            widget = getattr(self, name, None)
            if widget is not None:
                try:
                    widget.configure(state=state)
                except tk.TclError:
                    pass



    def _hidden_follows(self, profile) -> bool:
        return bool(profile.get("followed_private")) or profile.get("priv_like") == 3

    def _follow_incomplete(self, profile) -> bool:
        """关注列表是否不完整；游客、无 BDUSS、隐藏关注都需要候选探测。"""
        if profile.get("followed_forums"):
            return False
        if profile.get("follow_loaded") and not self._hidden_follows(profile) and not profile.get("follow_error"):
            return False
        return True



    def _remark_local_marks(self) -> None:
        if getattr(self, "_batch_running", False):
            return
        if not self._risk_enabled():
            return
        if not self._has_mark_rules():
            self.people_hint.configure(text="没有本地标记规则，请先在设置→本地标记中填写")
            self._focus_local_marks()
            return
        users = self._visible_people()
        if not users:
            self.people_hint.configure(text="当前筛选没有用户")
            return
        count = self._post_count()
        start_date, end_date = self._date_params()
        prefer = (
            "start"
            if self.post_date_mode_var.get() == "指定日期" and "最早" in self.post_prefer_var.get()
            else "end"
        )
        marked = 0
        missing = 0
        for user in users:
            uid = user.get("uid") or ""
            posts = cached_user_posts(uid, count, start_date, end_date, prefer) if uid else None
            if posts is None:
                missing += 1
                continue
            profile = {
                "ok": True,
                "recent_posts": posts,
                "bar_level": str(user.get("level") or ""),
                "follow_loaded": False,
                "followed_forums": [],
                "hidden_forums": [],
                "post_sources": {"cached": True},
            }
            self._apply_batch_mark(user, profile)
            marked += 1
        self._marks_stale = missing > 0
        self._refresh_people()
        note = f"重新标记完成 {marked}/{len(users)} 人"
        if missing:
            note += f" · 缓存不足 {missing} 人（未联网）"
        self.people_hint.configure(text=note)





    def _apply_level_to_threads(self, user, level: str) -> None:
        """把某个人的本吧等级写回楼层/楼中楼数据，刷新和导出时都不会丢。"""
        level = str(level or "").strip()
        if level in ("", "0", "None"):
            return
        uid = user.get("uid") or ""
        username = user.get("username") or ""
        name = user.get("name") or ""
        portrait = str(user.get("portrait") or "").split("?")[0]
        for thread in self.last_threads:
            for floor in thread.floors:
                match = (
                    (uid and floor.author_id == uid)
                    or (portrait and str(floor.portrait or "").split("?")[0] == portrait)
                    or (not uid and username and floor.author_name == username)
                    or (not uid and not username and floor.author == name)
                )
                if match and not floor.level:
                    floor.level = level
            for comment in thread.comments:
                match = (
                    (uid and comment.author_id == uid)
                    or (portrait and str(comment.portrait or "").split("?")[0] == portrait)
                    or (not uid and username and comment.author_name == username)
                    or (not uid and not username and comment.author == name)
                )
                if match and not comment.level:
                    comment.level = level



    def _configure_colors(self, tree) -> None:
        tree.tag_configure("ban", background=self.palette["danger_bg"])
        tree.tag_configure("warn", background=self.palette["warning_bg"])
        tree.tag_configure("safe", background=self.palette["safe_bg"])

    def _row_tags(self, bar, kind: str):
        if not self._risk_enabled():
            return ()
        key = "color_forum" if kind == "forum" else "color_posts"
        if not getattr(self, "settings", {}).get(key, True):
            return ()
        level, _ = classify([bar], getattr(self, "blacklist", {}))
        return (level,) if level else ()

    def _merge_forum_lists(self, profile, bar_name: str = "", bar_level: str = "") -> list:
        forums = profile.get("followed_forums") or []
        hidden = profile.get("hidden_forums") or []
        base = forums or hidden
        merged = {}
        for forum in base:
            name = forum.get("name")
            if not name:
                continue
            item = dict(forum)
            if forums:
                item["source"] = "公开关注"
            elif "BDUSS" in str(profile.get("follow_error") or ""):
                item["source"] = "profile片段"
            else:
                item["source"] = "隐藏关注"
            item["level_source"] = "官方" if int(item.get("level") or 0) > 0 else "-"
            merged[name] = item
        if not forums:
            for forum in profile.get("like_forum") or []:
                name = forum.get("name")
                if name and name not in merged:
                    merged[name] = {
                        "name": name,
                        "level": 0,
                        "level_name": "",
                        "source": "精选",
                        "level_source": "-",
                    }
        if bar_name and bar_level:
            name = bar_name[:-1] if bar_name.endswith("吧") and len(bar_name) > 2 else bar_name
            level_int = int(bar_level or 0)
            if name and level_int > 0:
                existing = merged.get(name)
                if existing is None:
                    merged[name] = {
                        "name": name,
                        "level": level_int,
                        "level_name": "",
                        "source": "等级",
                        "level_source": "主题",
                    }
                elif not existing.get("level"):
                    existing["level"] = level_int
                    existing["level_source"] = "主题"
        for forum in profile.get("probed_forums") or []:
            name = forum.get("name")
            if not name:
                continue
            existing = merged.get(name)
            if existing is None:
                merged[name] = {
                    "name": name,
                    "level": int(forum.get("level") or 0),
                    "level_name": forum.get("level_name") or "",
                    "source": "候选确认",
                    "level_source": "官方" if int(forum.get("level") or 0) > 0 else "-",
                }
            elif existing.get("source") != "公开关注":
                if int(forum.get("level") or 0) > int(existing.get("level") or 0):
                    existing["level"] = int(forum.get("level") or 0)
                existing["source"] = "候选确认"
        if self.settings.get("love_level_enabled", False):
            for forum in profile.get("level_forums") or []:
                name = str(forum.get("name") or "")
                if not name:
                    continue
                existing = merged.get(name)
                if existing is None:
                    merged[name] = {
                        "name": name,
                        "level": int(forum.get("level") or 0),
                        "level_name": "",
                        "source": "恋文等级",
                        "level_source": "恋文（第三方）",
                    }
                elif not int(existing.get("level") or 0):
                    existing["level"] = int(forum.get("level") or 0)
                    existing["level_source"] = "恋文（第三方）"
        return sorted(merged.values(), key=lambda item: int(item.get("level") or 0), reverse=True)

    def _forum_note(self, profile, bar_name: str = "", bar_level: str = "") -> str:
        forums = profile.get("followed_forums") or []
        hidden = profile.get("hidden_forums") or []
        if forums:
            note = f"共 {len(forums)} 个（完整）"
            levels = profile.get("level_forums") or []
            if self.settings.get("love_level_enabled", False) and levels:
                note += f"；恋文等级另列 {len(levels)} 个（非关注）"
            return note
        if hidden or profile.get("like_forum"):
            count = len(self._merge_forum_lists(profile, bar_name, bar_level))
            if self.mode_var.get() == "guest" and not profile.get("follow_loaded"):
                if self._hidden_follows(profile):
                    return f"对方隐藏了完整关注列表，已从 profile 片段补 {count} 个，不完整"
                return f"未加载公开关注，仅从 profile 片段补 {count} 个，不完整"
            if "BDUSS" in str(profile.get("follow_error") or ""):
                note = f"未登录，仅从 profile 片段补 {count} 个关注，不完整"
            else:
                note = f"对方隐藏了完整关注列表，已从等级头衔/精选恢复 {count} 个，可能不全"
            probed = len(profile.get("probed_forums") or [])
            if probed:
                note += f"；候选确认 {probed} 个"
            if self._hidden_follows(profile):
                note += "（风险提示：关注列表隐藏，未命中警告时标黄）"
            return note
        if profile.get("followed_private"):
            return "对方隐藏了完整关注列表（隐私设置），不可见（风险提示：关注列表隐藏，未命中警告时标黄）"
        return "没有查到关注列表"





    def _post_range_dialog(self, user, start: str, end: str, cache_start: str, cache_end: str):
        win = tk.Toplevel(self.root)
        win.title("导出发言")
        win.configure(bg="#ffffff")
        win.resizable(False, False)
        win.transient(self.root)
        win.grab_set()
        frame = tk.Frame(win, bg="#ffffff", padx=18, pady=14)
        frame.pack(fill="both", expand=True)
        tk.Label(
            frame,
            text="导出发言",
            bg="#ffffff",
            fg="#1f2937",
            font=("Microsoft YaHei UI", 12, "bold"),
        ).pack(anchor="w")
        tk.Label(
            frame,
            text=f"用户：{user.get('name') or user.get('username') or user.get('uid') or ''}",
            bg="#ffffff",
            fg="#475569",
            font=("Microsoft YaHei UI", 9),
        ).pack(anchor="w", pady=(4, 10))
        mode = tk.StringVar(value="current")
        tk.Radiobutton(
            frame,
            text=f"当前查询范围：{start or '不限'} ~ {end or '不限'}",
            value="current",
            variable=mode,
            bg="#ffffff",
            anchor="w",
        ).pack(fill="x", anchor="w")
        tk.Radiobutton(
            frame,
            text=f"已缓存回复范围：{cache_start or '无'} ~ {cache_end or '无'}（连续）",
            value="cache",
            variable=mode,
            bg="#ffffff",
            anchor="w",
        ).pack(fill="x", anchor="w")
        tk.Radiobutton(
            frame, text="自定义", value="custom", variable=mode, bg="#ffffff", anchor="w"
        ).pack(fill="x", anchor="w")
        custom = tk.Frame(frame, bg="#ffffff")
        custom.pack(fill="x", padx=(24, 0), pady=(2, 8))
        tk.Label(custom, text="开始", bg="#ffffff", fg="#475569").pack(side="left")
        start_var = tk.StringVar(value=cache_start or start)
        tk.Entry(custom, textvariable=start_var, width=12).pack(side="left", padx=(4, 10))
        tk.Label(custom, text="结束", bg="#ffffff", fg="#475569").pack(side="left")
        end_var = tk.StringVar(value=cache_end or end)
        tk.Entry(custom, textvariable=end_var, width=12).pack(side="left", padx=(4, 0))
        tk.Label(
            frame,
            text="只导出缓存已有发言，不自动联网补抓；内容为接口片段，最多 120 字。",
            bg="#ffffff",
            fg="#8a94a6",
            font=("Microsoft YaHei UI", 8),
            wraplength=440,
            justify="left",
        ).pack(anchor="w", pady=(0, 10))
        result: Dict[str, Any] = {"value": None}
        buttons = tk.Frame(frame, bg="#ffffff")
        buttons.pack(fill="x")

        def finish(value) -> None:
            result["value"] = value
            win.destroy()

        ttk.Button(buttons, text="取消", command=lambda: finish(None), width=8).pack(side="right")
        ttk.Button(buttons, text="导出", command=lambda: finish(mode.get()), width=8).pack(
            side="right", padx=(0, 8)
        )
        self._center_toplevel(win)
        self.root.wait_window(win)
        if not result["value"]:
            return None
        if result["value"] == "current":
            return start, end
        if result["value"] == "cache":
            if not cache_start and not cache_end:
                messagebox.showinfo("没有缓存", "当前用户还没有可导出的发言缓存。", parent=self.root)
                return None
            return cache_start, cache_end
        chosen_start = start_var.get().strip()
        chosen_end = end_var.get().strip()
        for value in (chosen_start, chosen_end):
            if value and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                messagebox.showerror("日期格式", "日期请用 YYYY-MM-DD。", parent=self.root)
                return None
        return chosen_start, chosen_end


    def _choose_report_path(self, user, profile, mark) -> None:
        initial = report_base(user, profile) + ".xlsx"
        path = filedialog.asksaveasfilename(
            title="导出用户报告",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile=initial,
            parent=self.root,
        )
        if not path:
            return
        base = os.path.splitext(path)[0]
        try:
            xlsx, md = export_user_report(
                base,
                user,
                profile,
                self.last_threads,
                mode="登录" if self.mode_var.get() == "login" else "免登录",
                mark=mark,
            )
        except Exception as error:  # noqa: BLE001
            messagebox.showerror("导出失败", friendly_error(error), parent=self.root)
            return
        self._append_log(f"用户报告已导出：{xlsx}；{md}", "ok")
        messagebox.showinfo(
            "导出完成",
            "已生成：" + chr(10) + xlsx + chr(10) + md,
            parent=self.root,
        )


    def _comment_depth_value(self) -> int:
        text = self.comment_depth_var.get()
        if "10" in text:
            return 10
        if "50" in text:
            return 50
        if "100" in text:
            return 100
        return 0

    def _merge_threads(self, new_threads) -> list:
        merged = {thread.meta.tid: thread for thread in self.last_threads}
        for thread in new_threads:
            merged[thread.meta.tid] = thread
        return list(merged.values())


    def _persist_result(self, threads) -> None:
        if not threads or not self.settings.get("save_result", True):
            return
        try:
            save_threads(
                list(threads),
                VERSION,
                encrypted=bool(self.settings.get("encrypt_result", False)),
            )
        except Exception as error:  # noqa: BLE001
            self._post(("log", f"保存最近结果失败：{error}", "err"))

    def _restore_last_result(self) -> None:
        if not self.settings.get("save_result", True):
            return
        try:
            threads, _meta = load_threads()
        except IncompatibleResult as error:
            archive_incompatible()
            self._append_log(f"上次结果无法恢复：{error}；旧文件已保留。", "err")
            return
        except Exception as error:  # noqa: BLE001
            self._append_log(f"读取上次结果失败：{error}", "err")
            return
        if not threads:
            return
        if self.settings.get("restore_mode", "auto") == "ask":
            if not messagebox.askyesno(
                "恢复上次结果",
                f"发现上次保存的 {len(threads)} 个帖子，恢复吗？",
                parent=self.root,
            ):
                return
        self.last_threads = threads
        self._profile_cache.clear()
        try:
            self._refresh_people()
            self._refresh_grid()
        except Exception:  # noqa: BLE001
            pass
        incomplete = sum(1 for thread in threads if not getattr(thread, "complete", True))
        self.status_var.set(
            f"已恢复 {len(threads)} 个帖子（{len(self._aggregate_users())} 位用户）；"
            + (f"不完整 {incomplete} 帖；" if incomplete else "")
            + "缺失身份可点「补全身份」"
        )


    def _high_risk_items(self) -> List[str]:
        items: List[str] = []
        try:
            if int(self.identity_concurrency_var.get() or 3) >= 5:
                items.append("贴吧ID并发 5（可能触发验证码或临时限流）")
        except (TypeError, ValueError, tk.TclError):
            pass
        try:
            if float(self.delay_var.get() or 0.8) < 0.5:
                items.append("请求间隔小于 0.5 秒")
        except (TypeError, ValueError, tk.TclError):
            pass
        try:
            if int(self.comment_concurrency_var.get() or 1) >= 3:
                items.append("楼中楼并发 3")
        except (TypeError, ValueError, tk.TclError):
            pass
        if self.mode_var.get() == "login":
            items.append("登录抓取")
        if self.settings.get("follow_enhance"):
            items.append("关注检测")
        if getattr(self, "proxy_mode_var", None) is not None and self.proxy_mode_var.get() == "自定义":
            items.append("自定义代理")
        return items



    def _on_escape(self, _event=None):
        grab = self.root.grab_current()
        if grab is not None and isinstance(grab, tk.Toplevel):
            try:
                grab.destroy()
            except tk.TclError:
                pass
            return "break"
        return None

    def _on_close(self) -> None:
        try:
            self.settings["window_geometry"] = self.root.geometry()
            save_settings(self.settings)
        except tk.TclError:
            pass
        active = self._network_task_active()
        if not active:
            self.root.destroy()
            return
        choice = self._ask_choice(
            "任务仍在进行",
            "未导出的已完成部分可能不会保留。",
            [("继续", False), ("停止并退出", True)],
            default=0,
        )
        if not choice:
            return
        self._closing = True
        self.stop_event.set()
        self._supplement_stop.set()
        resume_requests()
        self._batch_cancel = True
        self.status_var.set("正在停止…")
        self._append_log("正在停止任务并保存已完成部分…", "muted")
        self.root.after(100, self._wait_close)

    def _wait_close(self) -> None:
        worker_alive = bool(self.worker and self.worker.is_alive())
        batch_alive = bool(getattr(self, "_batch_running", False))
        supplement_alive = bool(getattr(self, "_supplement_running", False))
        if worker_alive or batch_alive or supplement_alive:
            self.root.after(100, self._wait_close)
            return
        self._persist_result(self.last_threads)
        self.root.destroy()

    # ------------------------------------------------------------- worker

    def _post(self, message: tuple) -> None:
        self.log_queue.put(message)


    def _poll_queue(self) -> None:
        try:
            while True:
                message = self.log_queue.get_nowait()
                self._handle_message(message)
        except queue.Empty:
            pass
        self._notify_cache_eviction()
        self.root.after(120, self._poll_queue)

    def _handle_message(self, message: tuple) -> None:
        kind = message[0]
        if kind == "log":
            self._handle_message_log(message)
        elif kind == "stage":
            self._handle_message_stage(message)
        elif kind == "done":
            self._handle_message_done(message)
        elif kind == "cancelled_partial":
            self._handle_message_cancelled_partial(message)
        elif kind == "probe_done":
            self._handle_message_probe_done(message)
        elif kind == "batch_profile":
            self._handle_message_batch_profile(message)
        elif kind == "batch_done":
            self._handle_message_batch_done(message)
        elif kind == "profile_ready":
            self._handle_message_profile_ready(message)
        elif kind == "enrich_done":
            self._handle_message_enrich_done(message)
        elif kind == "enrich_error":
            self._handle_message_enrich_error(message)
        elif kind == "identity_ask":
            self._handle_message_identity_ask(message)
        elif kind == "identity_supplement_progress":
            self._handle_message_identity_supplement_progress(message)
        elif kind == "identity_supplement_done":
            self._handle_message_identity_supplement_done(message)
        elif kind == "incomplete_retry_done":
            self._handle_message_incomplete_retry_done(message)
        elif kind == "incomplete_retry_stopped":
            self._handle_message_incomplete_retry_stopped(message)
        elif kind == "incomplete_retry_error":
            self._handle_message_incomplete_retry_error(message)
        elif kind == "identity_supplement_stopped":
            self._handle_message_identity_supplement_stopped(message)
        elif kind == "identity_supplement_error":
            self._handle_message_identity_supplement_error(message)
        elif kind == "identity_remember":
            self._handle_message_identity_remember(message)
        elif kind == "levels_filled":
            self._handle_message_levels_filled(message)
        elif kind == "levels_error":
            self._handle_message_levels_error(message)
        elif kind == "lookup_ready":
            self._handle_message_lookup_ready(message)
        elif kind == "levels_ready":
            self._handle_message_levels_ready(message)
        elif kind == "links_ready":
            self._handle_message_links_ready(message)
        elif kind == "love_ready":
            self._handle_message_love_ready(message)
        elif kind == "cancelled":
            self._handle_message_cancelled(message)
        elif kind == "captcha":
            self._handle_captcha_message(message)
        elif kind == "webview_test_done":
            self._handle_webview_test_done(message)
        elif kind == "embedded_done":
            self._handle_embedded_done(message)
        elif kind == "auto_login_done":
            self._handle_auto_login_done(message)
        elif kind == "switch_login_done":
            self._handle_switch_login_done(message)
        elif kind == "manual_cookie_probe":
            self._handle_message_manual_cookie_probe(message)
        elif kind == "rate_limit_ask":
            self._handle_message_rate_limit_ask(message)
        elif kind == "rate_limit":
            self._handle_message_rate_limit(message)
        elif kind == "error":
            self._handle_message_error(message)
        else:
            self._append_log(f"未知消息类型：{kind}", "warn")

    def _network_task_active(self) -> bool:
        return bool(
            self.running
            or getattr(self, "_supplement_running", False)
            or getattr(self, "_batch_running", False)
        )

    def _set_supplement_running(self, running: bool) -> None:
        self._supplement_running = running
        self.start_btn.configure(
            state="disabled" if running else "normal",
            text="补全中，稍后开始" if running else "开始统计",
        )
        if running:
            self.stop_btn.configure(state="normal")
        else:
            self.stop_btn.configure(state="disabled")
            resume_requests()
            set_stop_checker(self.stop_event.is_set)

    def _set_running(self, running: bool) -> None:
        self.running = running
        if running:
            self.start_btn.configure(state="disabled", bg="#9aa7b8")
            self.pause_btn.configure(state="normal", text="暂停")
            self.stop_btn.configure(state="normal")
            self.open_btn.configure(state="disabled")
            self.status_var.set("正在抓取，请稍候…")
            self.progress.pack(side="right")
            try:
                self.progress.start(12)
            except tk.TclError:
                pass
        else:
            self.start_btn.configure(state="normal", bg="#1F4E79")
            self.pause_btn.configure(state="disabled", text="暂停")
            self.stop_btn.configure(state="disabled")
            self._task_paused = False
            resume_requests()
            try:
                self.progress.stop()
            except tk.TclError:
                pass
            self.progress.pack_forget()
            if getattr(self, "_pending_retry", False):
                self.root.after(200, self._retry_incomplete_threads)



def _selftest() -> int:
    """冻结后的环境自检：解析一段 HTML、生成一个 xlsx，结果写到日志文件。"""
    import tempfile

    out = os.environ.get("TIEBA_SELFTEST_OUT") or os.path.join(tempfile.gettempdir(), "tieba_selftest.xlsx")
    log_path = out + ".txt"
    try:
        from tieba_reply_stats.models import Comment, Floor, ThreadData, ThreadMeta
        from tieba_reply_stats.parse import parse_thread_page

        html = (
            '<div class="l_post" data-field=\'{"author":{"user_id":1,"user_name":"测试用户"},'
            '"content":{"post_id":2,"post_no":1,"reply_num":1,"is_lz":1,"date":"2024-01-01 10:00"}}\'>'
            '<div class="d_post_content">你好</div></div>'
        )
        meta, floors = parse_thread_page(html, "https://tieba.baidu.com/p/1", "1")
        if not floors or floors[0].author != "测试用户":
            raise AssertionError(f"解析异常：{floors!r}")

        thread = ThreadData(
            meta=ThreadMeta(tid="1", title="自检", url="https://tieba.baidu.com/p/1"),
            floors=[
                Floor(
                    post_no=1,
                    pid="2",
                    author="测试用户",
                    author_id="1",
                    posted_at=None,
                    content="你好",
                    reply_num=1,
                    is_op=True,
                )
            ],
            comments=[
                Comment(
                    floor_no=1,
                    floor_author="测试用户",
                    pid="9",
                    author="回复者",
                    author_id="2",
                    posted_at=None,
                    content="回复内容",
                )
            ],
        )
        export_workbook(out, [thread])

        from openpyxl import load_workbook

        book = load_workbook(out)
        if "回复统计" not in book.sheetnames:
            raise AssertionError(f"工作表异常：{book.sheetnames}")
        with open(log_path, "w", encoding="utf-8") as handle:
            handle.write("OK\n" + out + "\n")
        return 0
    except Exception as error:  # noqa: BLE001
        with open(log_path, "w", encoding="utf-8") as handle:
            handle.write("FAIL\n" + repr(error) + "\n" + traceback.format_exc())
        return 1


_INSTANCE_SOCKET = None


def _focus_existing_window() -> None:
    if os.name != "nt":
        return
    try:
        import ctypes

        user32 = ctypes.windll.user32
        title = f"{APP_TITLE} v{VERSION}"
        handle = user32.FindWindowW(None, title)
        if handle:
            user32.ShowWindow(handle, 9)
            user32.SetForegroundWindow(handle)
    except Exception:  # noqa: BLE001
        pass


def _acquire_instance():
    global _INSTANCE_SOCKET
    import socket

    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        sock.bind(("127.0.0.1", 47831))
        sock.listen(1)
    except OSError:
        sock.close()
        return None
    _INSTANCE_SOCKET = sock
    return sock


def main() -> int:
    if "--webview2-verify" in sys.argv:
        index = sys.argv.index("--webview2-verify")
        from tieba_reply_stats.webview2_helper import main as verify_main

        sys.argv = [sys.argv[0], *sys.argv[index + 1 :]]
        return verify_main()
    if "--selftest" in sys.argv:
        return _selftest()
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa: BLE001
            pass
    if os.name == "nt":
        try:
            import ctypes

            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except Exception:  # noqa: BLE001
            pass
    if _acquire_instance() is None:
        _focus_existing_window()
        messagebox.showinfo("软件已在运行", "已有一个窗口在运行，请使用已有窗口。")
        return 0
    root = tk.Tk()
    App(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
