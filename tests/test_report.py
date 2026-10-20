from openpyxl import load_workbook

from tieba_reply_stats.models import Floor, ThreadData, ThreadMeta
from tieba_reply_stats.report import export_user_report, safe_name


def make_thread():
    return ThreadData(
        meta=ThreadMeta(tid="1", title="测试帖", forum="测试吧"),
        floors=[
            Floor(post_no=1, pid="1", author="小明", author_id="1", author_name="ming", level="5"),
            Floor(post_no=2, pid="2", author="小红", author_id="2", author_name="hong", level="3"),
        ],
    )


def test_safe_name():
    assert safe_name("a/b:c*?") == "a_b_c"
    assert safe_name("") == "user"


def test_export_report(tmp_path):
    user = {"name": "小明", "username": "ming", "uid": "1", "level": "5"}
    profile = {
        "nickname": "小明",
        "username": "ming",
        "user_id": "1",
        "fans": 10,
        "concern": 20,
        "post_num": 30,
        "followed_forums": [{"name": "测试吧", "level": 5}],
        "recent_posts": [
            {
                "time": "2026-10-19",
                "type": "主题帖",
                "bar": "测试吧",
                "title": "t",
                "content": "c",
                "url": "u",
            }
        ],
    }
    base = str(tmp_path / "report")
    xlsx, md = export_user_report(
        base, user, profile, [make_thread()], mode="login", mark={"level": "红", "bar": "坏吧"}
    )
    book = load_workbook(xlsx)
    assert book.sheetnames == ["基本信息", "关注贴吧", "发言", "本帖参与", "本地标记"]
    text = open(md, encoding="utf-8").read()
    assert "用户报告" in text
    assert "测试吧" in text
    assert "本地标记" in text
