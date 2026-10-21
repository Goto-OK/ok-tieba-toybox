from openpyxl import load_workbook

from tieba_reply_stats.report import export_user_posts


def test_export_posts_sheets_and_links(tmp_path):
    posts = [
        {
            "time": "2025-01-02",
            "type": "回复",
            "bar": "吧",
            "title": "标题",
            "content": "内容",
            "url": "https://tieba.baidu.com/p/1?pid=2",
        },
        {
            "time": "2025-02-03",
            "type": "主题帖",
            "bar": "吧",
            "title": "主题",
            "content": "正文",
            "url": "https://tieba.baidu.com/p/3",
        },
    ]
    path = export_user_posts(
        str(tmp_path / "posts"),
        {"name": "测试", "username": "u", "uid": "1"},
        posts,
        "2025-01-01",
        "2025-12-31",
        {"reply": "2025-01-02 ~ 2025-01-02", "topic": "2025-02-03 ~ 2025-02-03"},
    )
    book = load_workbook(path)
    assert book.sheetnames == ["汇总", "主题帖", "回复"]
    assert book["主题帖"].max_row == 2
    assert book["回复"].max_row == 2
    assert book["主题帖"]["D2"].hyperlink.target == "https://tieba.baidu.com/p/3"
    assert book["回复"]["D2"].hyperlink.target == "https://tieba.baidu.com/p/1?pid=2"
    assert book["汇总"]["B9"].value == "2025-01-02 ~ 2025-01-02"
    assert book["汇总"]["B10"].value == "2025-02-03 ~ 2025-02-03"
